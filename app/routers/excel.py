"""Excel import/export router.

Import understands multiple formats:
  - Las Heras: numeric codes (0.1, 1.1, etc.)
  - Lugones/El Encuentro: date-encoded codes (2025-01-01 = 1.1, day=section, month=item)
  - All share: catalog sheets (00_*), computation sheet (01_C&P), detail sheets

Auto-detects format based on code column content.
"""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from io import BytesIO
from typing import Literal
from urllib.parse import quote
from uuid import UUID

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse

from app.auth import get_current_user, require_editor
from app.budget_prices import budget_config, fetch_all, initial_indirects, today
from app.catalog_prices import fuente_importada, normalize_codigo, price_changed
from app.calculations import (
    calc_budget_summary,
    fraction_to_pct,
    is_section,
    price_item,
    sale_totals,
)
from app.db import get_data_db
from app.routers.catalogs import insert_entries, start_history, update_entries
from app.terrac_export import FMT_PESOS as TERRAC_FMT_PESOS
from app.terrac_export import terrac_bytes
from app.tree import get_parent_candidates, normalize_item_code, safe_float

router = APIRouter()

# Column offsets for 01_C&P (based on Las Heras Excel structure)
# Row 7+ data columns:
COL_ITEM = 0
COL_DESC = 1
COL_UNIDAD = 2
COL_CANTIDAD = 3
# Direct costs (unit prices)
COL_MAT_UNIT = 4
COL_MO_UNIT = 9  # MO total unitario
COL_DIRECTO_UNIT = 10  # General unitario directo
# Direct costs (general/total)
COL_MAT_TOTAL = 11
COL_MO_GENERAL = 12
COL_DIRECTO_GENERAL = 13
# Indirect costs (general)
COL_IND_MAT = 14
COL_IND_MO = 15
COL_IND_GENERAL = 16
# Benefits (general)
COL_BEN_MAT = 17
COL_BEN_MO = 18
COL_BEN_GENERAL = 19
# Net total (unit)
COL_NETO_UNIT_MAT = 20
COL_NETO_UNIT_MO = 21
COL_NETO_UNIT_GEN = 22
# Net total (general)
COL_NETO_MAT = 23
COL_NETO_MO = 24
COL_NETO_GENERAL = 25

# Catalog sheet config: (sheet_name, tipo, header_row)
CATALOG_SHEETS = [
    ("00_Mat", "material", 1),
    ("00_MO", "mano_obra", 2),
    ("00_Eq", "equipo", 2),
    ("00_Sub", "subcontrato", 2),
]


def _cell(df: pd.DataFrame, row: int, col: int) -> float:
    """Safely read a numeric cell from a dataframe."""
    try:
        val = df.iloc[row, col]
        return safe_float(val) or 0.0
    except (IndexError, KeyError):
        return 0.0


def _cell_str(df: pd.DataFrame, row: int, col: int) -> str:
    """Safely read a string cell."""
    try:
        val = df.iloc[row, col]
        if val is None or (isinstance(val, float) and pd.isna(val)):
            return ""
        return str(val).strip()
    except (IndexError, KeyError):
        return ""


def _parse_catalogs(
    df_dict: dict[str, pd.DataFrame], org_id: str, catalog_id: str | None = None,
) -> list[dict]:
    """Parse all catalog sheets into catalog_entries rows."""
    entries: list[dict] = []
    for sheet_name, tipo, header_row in CATALOG_SHEETS:
        if sheet_name not in df_dict:
            continue
        df = df_dict[sheet_name]
        # The actual data starts after the header row
        for i in range(header_row + 1, len(df)):
            code = _cell_str(df, i, 0)
            if not code:
                continue
            desc = _cell_str(df, i, 1)
            if not desc:
                continue
            entries.append({
                "catalog_id": catalog_id,
                "org_id": org_id,
                "tipo": tipo,
                "codigo": code,
                "descripcion": desc,
                "unidad": _cell_str(df, i, 2),
                "precio_con_iva": safe_float(df.iloc[i, 3] if df.shape[1] > 3 else None),
                "precio_sin_iva": safe_float(df.iloc[i, 4] if df.shape[1] > 4 else None) if tipo == "material" else safe_float(df.iloc[i, 3] if df.shape[1] > 3 else None),
                "referencia": _cell_str(df, i, 4) if tipo != "material" else None,
            })
    return entries


def _parse_computation_sheet(
    df: pd.DataFrame, budget_id: str, org_id: str,
) -> list[dict]:
    """Parse the 01_C&P sheet into budget_items rows.

    Data starts at row 7 (0-indexed). Section titles have no CANTIDAD (col 3).
    """
    items: list[dict] = []
    code_to_id_placeholder: dict[str, int] = {}  # code -> index for parent lookup

    date_codes_corrected = 0

    for i in range(7, len(df)):
        # Read RAW value from code column to preserve Timestamps
        code_raw_val = df.iloc[i, COL_ITEM] if df.shape[1] > COL_ITEM else None
        code_raw = _cell_str(df, i, COL_ITEM)
        desc = _cell_str(df, i, COL_DESC)
        cantidad_val = safe_float(df.iloc[i, COL_CANTIDAD] if df.shape[1] > COL_CANTIDAD else None)

        # Track date-code conversions
        if isinstance(code_raw_val, (pd.Timestamp,)):
            date_codes_corrected += 1

        # Use raw value for normalization (preserves Timestamp for date-code detection)
        code_for_normalize = code_raw_val if isinstance(code_raw_val, pd.Timestamp) else code_raw

        # Skip empty rows
        if not code_raw and not desc and not isinstance(code_raw_val, pd.Timestamp):
            continue

        # Section title (no code or no cantidad) — create as parent item
        if cantidad_val is None:
            if code_raw or desc or isinstance(code_raw_val, pd.Timestamp):
                label = f"{code_raw} {desc}".strip() if desc else code_raw
                items.append({
                    "budget_id": budget_id,
                    "org_id": org_id,
                    "parent_id": None,
                    "code": normalize_item_code(code_for_normalize) or None,
                    "description": label,
                    "unidad": None,
                    "cantidad": None,
                    "mat_unitario": 0,
                    "mo_unitario": 0,
                    "mat_total": 0,
                    "mo_total": 0,
                    "directo_total": 0,
                    "indirecto_total": 0,
                    "beneficio_total": 0,
                    "neto_total": 0,
                    "notas": "Seccion",
                    "sort_order": len(items),
                    "_code_norm": normalize_item_code(code_for_normalize),
                })
                norm = normalize_item_code(code_for_normalize)
                if norm:
                    code_to_id_placeholder[norm] = len(items) - 1
            continue

        # Regular item with data
        code_norm = normalize_item_code(code_for_normalize)

        # Find parent by code hierarchy
        parent_idx = None
        for candidate in get_parent_candidates(code_norm):
            if candidate in code_to_id_placeholder:
                parent_idx = code_to_id_placeholder[candidate]
                break

        mat_unit = _cell(df, i, COL_MAT_UNIT)
        mo_unit = _cell(df, i, COL_MO_UNIT)

        items.append({
            "budget_id": budget_id,
            "org_id": org_id,
            "parent_id": None,  # resolved after insert by index
            "code": code_norm or None,
            "description": desc,
            "unidad": _cell_str(df, i, COL_UNIDAD),
            "cantidad": cantidad_val,
            "mat_unitario": mat_unit,
            "mo_unitario": mo_unit,
            "mat_total": _cell(df, i, COL_MAT_TOTAL),
            "mo_total": _cell(df, i, COL_MO_GENERAL),
            "directo_total": _cell(df, i, COL_DIRECTO_GENERAL),
            "indirecto_total": _cell(df, i, COL_IND_GENERAL),
            "beneficio_total": _cell(df, i, COL_BEN_GENERAL),
            "neto_total": _cell(df, i, COL_NETO_GENERAL),
            "notas": "Importado desde Excel",
            "sort_order": len(items),
            "_code_norm": code_norm,
            "_parent_idx": parent_idx,
        })
        if code_norm:
            code_to_id_placeholder[code_norm] = len(items) - 1

    return items, date_codes_corrected


def _section_tipo(header: str) -> str | None:
    """Map a detail-sheet section header (upper-cased) to an item_resources tipo.

    Specific "MANO DE OBRA - X" headers must be checked before the generic
    MATERIALES / MANO DE OBRA ones, otherwise equipos, materiales indirectos and
    subcontratos get misclassified.
    """
    if "EQUIPO" in header:
        return "equipo"
    if "SUBCONTRAT" in header:
        return "subcontrato"
    if "MANO DE OBRA" in header and "MATERIAL" in header:
        return "mo_material"
    if "MANO DE OBRA" in header:
        return "mano_obra"
    if "MATERIALES" in header:
        return "material"
    return None


def _parse_detail_sheets(
    df_dict: dict[str, pd.DataFrame],
    sheet_names: list[str],
    org_id: str,
) -> dict[str, list[dict]]:
    """Parse detail sheets (1.1, 1.2, etc.) into item_resources keyed by item code."""
    resources_by_code: dict[str, list[dict]] = {}

    for sheet_name in sheet_names:
        if sheet_name not in df_dict:
            continue
        df = df_dict[sheet_name]
        code_norm = normalize_item_code(sheet_name)
        if not code_norm:
            continue

        resources: list[dict] = []
        current_tipo: str | None = None

        for i in range(4, len(df)):
            first_cell = _cell_str(df, i, 0)
            upper = first_cell.upper()

            # Detect section headers
            if "TOTAL" in upper:
                current_tipo = None
                continue
            section_tipo = _section_tipo(upper)
            if section_tipo:
                current_tipo = section_tipo
                continue

            if current_tipo is None:
                continue

            # Skip header rows (Codigo, Descripcion, etc.)
            if first_cell.lower() in ("codigo", "código", ""):
                continue

            desc = _cell_str(df, i, 1)
            if not desc:
                continue

            cantidad = safe_float(df.iloc[i, 3] if df.shape[1] > 3 else None)
            dias = safe_float(df.iloc[i, 4] if df.shape[1] > 4 else None)
            desperdicio = fraction_to_pct(safe_float(df.iloc[i, 5] if df.shape[1] > 5 else None))
            cantidad_eff = safe_float(df.iloc[i, 6] if df.shape[1] > 6 else None)
            precio = safe_float(df.iloc[i, 7] if df.shape[1] > 7 else None)
            subtotal = safe_float(df.iloc[i, 8] if df.shape[1] > 8 else None)

            if cantidad is None and cantidad_eff is None:
                continue

            resource = {
                "org_id": org_id,
                "tipo": current_tipo,
                "codigo": first_cell,
                "descripcion": desc,
                "unidad": _cell_str(df, i, 2),
                "cantidad": cantidad,
                "desperdicio_pct": desperdicio,
                "cantidad_efectiva": cantidad_eff or (
                    cantidad * (1 + desperdicio / 100) if cantidad else 0
                ),
                "precio_unitario": precio,
                "subtotal": subtotal or 0,
            }
            if current_tipo == "mano_obra":
                # Excel: Cantidad = trabajadores, Dias, "Desperdicio" = cargas/ineficiencia
                resource["trabajadores"] = cantidad or 0
                resource["dias"] = dias or 0
                resource["cargas_sociales_pct"] = desperdicio
            resources.append(resource)

        if resources:
            resources_by_code[code_norm] = resources

    return resources_by_code


def _insert_entries(db, catalog_id: str, entries: list[dict], filename: str) -> None:  # type: ignore[no-untyped-def]
    """Insert catalog entries in batches, each one starting its price history ("Importado de …")."""
    fuente = fuente_importada(filename)
    for batch_start in range(0, len(entries), 100):
        batch = [{**e, "catalog_id": catalog_id, "fuente": fuente}
                 for e in entries[batch_start:batch_start + 100]]
        insert_entries(db, batch)


def _save_catalog(db, org_id: str, filename: str, entries: list[dict]) -> dict:  # type: ignore[no-untyped-def]
    """Save the price sheets of an imported Excel.

    No price sheets: no list. A list imported before from the same file (same org) is
    updated: changed prices go through the catalog rule (date + history), new codes are
    added, codes the Excel no longer has are left alone. Otherwise a new list is created.
    It is never marked oficial here.
    """
    out = {"catalog_id": None, "catalog_name": None, "catalog_reused": False,
           "precios_actualizados": 0, "precios_nuevos": 0}
    if not entries:
        return out

    previous = (
        db.table("price_catalogs")
        .select("*")
        .eq("org_id", org_id)
        .eq("source_file", filename)
        .order("created_at", desc=True)
        .limit(1)
        .execute()
        .data or []
    )
    if not previous:
        name = f"Catálogo - {filename}"
        created = db.table("price_catalogs").insert({
            "org_id": org_id,
            "name": name,
            "source_file": filename,
        }).execute()
        catalog_id = created.data[0]["id"]
        _insert_entries(db, catalog_id, entries, filename)
        return {**out, "catalog_id": catalog_id, "catalog_name": name, "precios_nuevos": len(entries)}

    catalog_id = str(previous[0]["id"])
    current = fetch_all(
        lambda: db.table("catalog_entries").select("*")
        .eq("catalog_id", catalog_id).eq("org_id", org_id).order("id")
    )
    by_codigo: dict[str, list[dict]] = {}
    for e in current:
        by_codigo.setdefault(normalize_codigo(e.get("codigo")), []).append(e)

    nuevas: list[dict] = []
    changes: list[tuple[dict, dict]] = []
    for entry in entries:
        # Same code and same tipo: a material never takes the price of a mano de obra with
        # its code. Each row of the list is matched once
        candidates = by_codigo.get(normalize_codigo(entry["codigo"]), [])
        match = next((e for e in candidates if e.get("tipo") == entry["tipo"]), None)
        if match is None:
            nuevas.append(entry)
            continue
        candidates.remove(match)
        precio = entry.get("precio_sin_iva")
        # The Excel has no dates: an empty or 0 price is "sin precio" (is_price) and
        # never replaces the price the list already has
        if not precio or not price_changed(match, {"precio_sin_iva": precio}):
            continue
        update = {"precio_sin_iva": precio, "fuente": fuente_importada(filename), "fuente_url": None}
        if entry.get("precio_con_iva") is not None:
            update["precio_con_iva"] = entry["precio_con_iva"]
        changes.append((match, update))

    if changes:
        # Lists saved before every price went to the history: keep their current price there
        start_history(db, org_id, [old for old, _ in changes])
        update_entries(db, org_id, changes)
    _insert_entries(db, catalog_id, nuevas, filename)
    return {**out, "catalog_id": catalog_id, "catalog_name": previous[0].get("name"), "catalog_reused": True,
            "precios_actualizados": len(changes), "precios_nuevos": len(nuevas)}


# ── Endpoints ────────────────────────────────────────────────────────────────


@router.post("/import-excel")
async def import_excel(
    file: UploadFile = File(...),
    user: dict = Depends(require_editor),
):
    """Import a construction budget Excel (Las Heras format)."""
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Subí un archivo de Excel (.xlsx o .xls)")

    contents = await file.read()
    df_dict = pd.read_excel(BytesIO(contents), sheet_name=None, header=None)
    # The app's own simple sheet is not a budget to copy: it is loaded in Cargar obra (it used to create an
    # empty budget and answer "Excel importado")
    if "01_C&P" not in df_dict and "Presupuesto" in df_dict:
        primera = [str(v).strip() for v in df_dict["Presupuesto"].iloc[0].tolist()] if len(df_dict["Presupuesto"]) else []
        if "Codigo" in primera and ("Precio sin IVA" in primera or "Neto Total" in primera):
            raise HTTPException(
                400,
                "Este Excel es la planilla simple que baja la app (Exportar). Para volver a armar el presupuesto "
                "con ella, subila en Cargar obra.",
            )

    db = get_data_db()
    org_id = user["org_id"]

    # 1-2. Price sheets: update the list of this same file, or create one
    entries = _parse_catalogs(df_dict, org_id)
    catalog = _save_catalog(db, org_id, file.filename, entries)
    catalog_id = catalog["catalog_id"]

    # 3. Create budget
    base_name = os.path.splitext(os.path.basename(file.filename))[0].strip()
    budget_name = base_name or f"Presupuesto {datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}"

    budget = db.table("budgets").insert({
        "org_id": org_id,
        "name": budget_name,
        "description": f"Importado desde {file.filename}",
        "source_file": file.filename,
        "status": "draft",
        "indirectos": initial_indirects(db, org_id),
        "precios_al": today().isoformat(),
    }).execute()
    budget_id = budget.data[0]["id"]
    # The Excel brings the direct cost; the price is the cascade of this budget, like everywhere
    config = budget_config(db, org_id, budget.data[0])

    # 4. Parse 01_C&P computation sheet
    items_inserted = 0
    resources_inserted = 0
    neto_excel = 0.0
    neto_app = 0.0

    if "01_C&P" in df_dict:
        parsed_items, date_codes_corrected = _parse_computation_sheet(df_dict["01_C&P"], budget_id, org_id)
        for item in parsed_items:
            if is_section(item):
                continue
            neto_excel += float(item.get("neto_total") or 0)
            price_item(item, config)
            neto_app += float(item.get("neto_total") or 0)

        # Insert items one by one to resolve parent_id references
        idx_to_db_id: dict[int, str] = {}
        for idx, item in enumerate(parsed_items):
            # Resolve parent
            parent_idx = item.pop("_parent_idx", None)
            code_norm = item.pop("_code_norm", "")
            if parent_idx is not None and parent_idx in idx_to_db_id:
                item["parent_id"] = idx_to_db_id[parent_idx]

            result = db.table("budget_items").insert(item).execute()
            if result.data:
                db_id = result.data[0]["id"]
                idx_to_db_id[idx] = db_id
                items_inserted += 1

        # 5. Parse detail sheets and insert item_resources
        system_sheets = {
            "00_Mat", "00_MO", "00_Eq", "00_Sub", "00_JEF + ESTR",
            "01_C&P", "01_VENTA", "01_RESUMEN VENTA",
            "01_RESUMEN MAT.", "01_RESUMEN MAT M.O.",
            "01_RESUMEN SERV. M.O.", "ESTRUCTURA", "TIEMPOS",
        }
        detail_sheet_names = [
            s for s in df_dict.keys()
            if s not in system_sheets and not s.startswith("00_") and not s.startswith("01_")
        ]

        resources_by_code = _parse_detail_sheets(df_dict, detail_sheet_names, org_id)

        # Match resources to inserted items by code
        code_to_db_id: dict[str, str] = {}
        for idx, item in enumerate(parsed_items):
            code = item.get("code")
            if code and idx in idx_to_db_id:
                code_to_db_id[normalize_item_code(code)] = idx_to_db_id[idx]

        for code, resources in resources_by_code.items():
            item_db_id = code_to_db_id.get(code)
            if not item_db_id:
                continue
            for res in resources:
                res["item_id"] = item_db_id
            # Insert in batches
            for batch_start in range(0, len(resources), 50):
                batch = resources[batch_start:batch_start + 50]
                db.table("item_resources").insert(batch).execute()
                resources_inserted += len(batch)

    return {
        "message": "Excel importado",
        "catalog_id": catalog_id,
        "catalog_entries": len(entries),
        "catalog_reused": catalog["catalog_reused"],
        "catalog_name": catalog["catalog_name"],
        "precios_actualizados": catalog["precios_actualizados"],
        "precios_nuevos": catalog["precios_nuevos"],
        "budget_id": budget_id,
        "budget_name": budget_name,
        "items_inserted": items_inserted,
        "resources_inserted": resources_inserted,
        "date_codes_corrected": date_codes_corrected if "01_C&P" in df_dict else 0,
        # "Tu Excel decía $X; con tu Coeficiente de pase da $Y"
        "neto_excel": round(neto_excel, 2),
        "neto_app": round(neto_app, 2),
    }


@router.get("/{budget_id}/export/excel")
async def export_budget_excel(
    budget_id: UUID,
    formato: Literal["simple", "terrac"] = Query("simple"),
    user: dict = Depends(get_current_user),
):
    """Export a budget to Excel.

    ``formato=simple`` (default): one row per work with the whole cascade (planilla simple).
    ``formato=terrac``: Sol's planilla (01_C&P, one sheet per work, Coeficiente de pase),
    which Cargar obra can read back (see app/terrac_export.py).
    """
    db = get_data_db()
    bid = str(budget_id)
    org_id = user["org_id"]

    budget = (
        db.table("budgets")
        .select("*")
        .eq("id", bid)
        .eq("org_id", org_id)
        .single()
        .execute()
    )
    if not budget.data:
        raise HTTPException(404, "Presupuesto no encontrado")

    if formato == "terrac":
        return _terrac_response(db, budget.data, org_id)

    items = (
        db.table("budget_items")
        .select("*")
        .eq("budget_id", bid)
        .eq("org_id", org_id)
        .order("sort_order")
        .execute()
    )
    if not items.data:
        raise HTTPException(404, "El presupuesto no tiene trabajos")

    iva_pct = budget_config(db, org_id, budget.data)["iva_pct"]
    df = pd.DataFrame(excel_rows(items.data, iva_pct))
    total_row = {col: "" for col in df.columns}
    total_row["Codigo"] = "TOTAL"
    for col in EXCEL_TOTALS:
        total_row[col] = pd.to_numeric(df[col], errors="coerce").sum()
    df = pd.concat([df, pd.DataFrame([total_row])], ignore_index=True)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Presupuesto")
        _format_simple_sheet(writer.sheets["Presupuesto"], list(df.columns))
    output.seek(0)

    safe_name = (budget.data["name"] or "presupuesto").replace(" ", "_")
    return _xlsx_response(output, f"{safe_name}_{datetime.now().strftime('%Y%m%d')}.xlsx")


def _xlsx_response(output: BytesIO, filename: str) -> StreamingResponse:
    # Headers are latin-1: an ASCII fallback plus the real name (RFC 5987) for any obra name
    ascii_name = filename.encode("ascii", "replace").decode().replace("?", "_").replace('"', "")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(filename)}'},
    )


# Widths of the planilla simple (the rest: money columns)
_SIMPLE_WIDTHS = {"Codigo": 10, "Descripcion": 60, "Unidad": 8, "Cantidad": 11, "Notas": 50}


def _format_simple_sheet(ws, columns: list[str]) -> None:  # type: ignore[no-untyped-def]
    """Planilla simple: bold header, column widths, pesos format, first row fixed.

    Only formatting: the header row and its names stay as they are (tests and e2e read them).
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="FFE8F5EE")
    money = {c for c in columns if c not in _SIMPLE_WIDTHS}
    for idx, name in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = _SIMPLE_WIDTHS.get(name, 18)
        head = ws.cell(1, idx)
        head.font = Font(bold=True)
        head.fill = head_fill
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, idx)
            if name in money:
                cell.number_format = TERRAC_FMT_PESOS
            elif name == "Cantidad":
                cell.number_format = "#,##0.00"
    for idx in range(1, len(columns) + 1):
        ws.cell(ws.max_row, idx).font = Font(bold=True)  # TOTAL row
    ws.freeze_panes = "A2"


_RESOURCES_CHUNK = 200  # item ids per item_resources request (as in budgets.py)


def _terrac_response(db, budget: dict, org_id: str) -> StreamingResponse:  # type: ignore[no-untyped-def]
    """Planilla Terrac of a budget (formato=terrac)."""
    bid = str(budget["id"])
    items = fetch_all(
        lambda: db.table("budget_items").select("*").eq("budget_id", bid).eq("org_id", org_id).order("id")
    )
    if not items:
        raise HTTPException(404, "El presupuesto no tiene trabajos")
    items.sort(key=lambda i: (i.get("sort_order") is None, i.get("sort_order") or 0))

    ids = [str(i["id"]) for i in items if not is_section(i)]
    resources: dict[str, list[dict]] = {}
    for start in range(0, len(ids), _RESOURCES_CHUNK):
        chunk = ids[start:start + _RESOURCES_CHUNK]
        for r in fetch_all(
            lambda chunk=chunk: db.table("item_resources").select("*")
            .eq("org_id", org_id).in_("item_id", chunk).order("id")
        ):
            resources.setdefault(str(r["item_id"]), []).append(r)

    cfg = _apply_cfg_defaults(budget_config(db, org_id, budget))
    output = terrac_bytes(budget, items, resources, cfg, pdf_totals(items, cfg))
    name = " ".join(str(budget.get("name") or "presupuesto").split())
    return _xlsx_response(output, f"{name} - Planilla Terrac - {today().isoformat()}.xlsx")


# Columns of the exported Excel that get a total (the same sums as calc_budget_summary)
EXCEL_TOTALS = (
    "MAT Total", "MO Total", "Directo Total", "Indirecto Total", "Beneficio Total",
    "Impuestos Total", "Precio sin IVA", "IVA", "Precio con IVA",
)


def excel_rows(items: list[dict], iva_pct: float) -> list[dict]:
    """Rows of the exported Excel: the saved numbers of each work (the whole cascade).

    Section rows carry no amounts, so each totals column adds up to calc_budget_summary.
    """
    rows = []
    for item in items:
        section = is_section(item)
        neto, iva, total = sale_totals(item, iva_pct)

        def amount(value: object, section: bool = section) -> object:
            return "" if section else float(value or 0)

        rows.append({
            "Codigo": item.get("code") or "",
            "Descripcion": item.get("description") or "",
            "Unidad": item.get("unidad") or "",
            "Cantidad": item.get("cantidad") or "",
            "MAT Unitario": amount(item.get("mat_unitario")),
            "MO Unitario": amount(item.get("mo_unitario")),
            "MAT Total": amount(item.get("mat_total")),
            "MO Total": amount(item.get("mo_total")),
            "Directo Total": amount(item.get("directo_total")),
            "Indirecto Total": amount(item.get("indirecto_total")),
            "Beneficio Total": amount(item.get("beneficio_total")),
            "Impuestos Total": amount(item.get("impuestos_total")),
            "Precio sin IVA": amount(neto),
            "IVA": amount(iva),
            "Precio con IVA": amount(total),
            "Notas": item.get("notas") or "",
        })
    return rows


# ── PDF Export ──────────────────────────────────────────────────────────────

# Brand colors
_COLOR_HEADER = "#143D34"   # dark teal — header background
_COLOR_ACCENT = "#2D8D68"   # teal — accent lines and KPI borders
_COLOR_SECTION = "#E8F5EE"  # light green — section row background
_COLOR_SECTION_TEXT = "#143D34"
_COLOR_BORDER = "#D1D5DB"   # light gray borders
_COLOR_ROW_ALT = "#F9FAFB"  # alternate row background
_COLOR_TOTAL_ROW = "#143D34"

# Default indirect config (used if table is empty)
_INDIRECT_DEFAULTS: dict[str, float] = {
    "imprevistos_pct": 3,
    "estructura_pct": 15,
    "jefatura_pct": 8,
    "logistica_pct": 5,
    "herramientas_pct": 3,
    "beneficio_pct": 10,
    "ingresos_brutos_pct": 7,
    "imp_cheque_pct": 1.2,
    "iva_pct": 21,
}


def _fmt_ars(val: float) -> str:
    """Format a number as Argentine peso: $ 1.234.567,89"""
    if val == 0:
        return "$ 0"
    # Use comma as thousands sep, dot as decimal (es-AR style)
    formatted = f"{abs(val):,.2f}"
    # swap separators: 1,234,567.89 -> 1.234.567,89
    formatted = formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"$ {'-' if val < 0 else ''}{formatted}"


def _fmt_pct(val: float) -> str:
    return f"{val:g}%"


def _apply_cfg_defaults(cfg: dict) -> dict:
    result = dict(cfg)
    for k, v in _INDIRECT_DEFAULTS.items():
        if result.get(k) is None:
            result[k] = v
    return result


def _split(total: float, weights: list[float]) -> list[float]:
    """Split a saved amount by weights (the % of each concept), in cents, adding up exactly."""
    cents = round(total * 100)
    peso = sum(weights)
    if not weights:
        return []
    if not peso:
        return [0.0] * (len(weights) - 1) + [cents / 100]
    parts = [round(cents * w / peso) for w in weights]
    last = max(i for i, w in enumerate(weights) if w)
    parts[last] += cents - sum(parts)
    return [p / 100 for p in parts]


def pdf_totals(all_items: list[dict], cfg: dict) -> dict:
    """Totals of the internal PDF: the saved sums (calc_budget_summary), never recalculated.

    The % (labels) come from the budget's config; each amount is the saved one, split
    by those % only where the DB keeps a single amount (the 5 indirects, IIBB + cheque).
    Every subtotal is the sum of the rows above it.
    """
    summary = calc_budget_summary(all_items, cfg["iva_pct"])
    ind_keys = ("imprevistos_pct", "estructura_pct", "jefatura_pct", "logistica_pct", "herramientas_pct")
    ind_pcts = [float(cfg[k]) for k in ind_keys]
    tax_pcts = [float(cfg["ingresos_brutos_pct"]), float(cfg["imp_cheque_pct"])]
    directo = summary["directo_total"]
    subtotal_02 = round(directo + summary["indirecto_total"], 2)
    subtotal_03 = round(subtotal_02 + summary["beneficio_total"], 2)
    return {
        **summary,
        "indirectos": dict(zip(ind_keys, _split(summary["indirecto_total"], ind_pcts))),
        "ind_pcts": dict(zip(ind_keys, ind_pcts)),
        "total_ind_pct": sum(ind_pcts),
        "ben_pct": float(cfg["beneficio_pct"]),
        "iibb_pct": tax_pcts[0],
        "cheque_pct": tax_pcts[1],
        "iibb": _split(summary["impuestos_total"], tax_pcts)[0],
        "cheque": _split(summary["impuestos_total"], tax_pcts)[1],
        "iva_pct": float(cfg["iva_pct"]),
        "subtotal_02": subtotal_02,
        "subtotal_03": subtotal_03,
    }


def _build_page_footer(canvas, doc):
    """Draw page number footer on every page."""
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    canvas.saveState()
    page_num = canvas.getPageNumber()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#9CA3AF"))
    text = f"Página {page_num} — Generado por Presupuestador"
    canvas.drawRightString(doc.pagesize[0] - 1.5 * cm, 0.8 * cm, text)
    # Left: brand line
    canvas.setFillColor(colors.HexColor(_COLOR_ACCENT))
    canvas.drawString(1.5 * cm, 0.8 * cm, "PRESUPUESTADOR — SOLE")
    canvas.restoreState()


def _pdf_header(budget_data: dict, pw: float, styles, with_description: bool = True) -> list:  # type: ignore[no-untyped-def]
    """Header of both PDFs: band with the obra name, date and version, accent bar and description."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    title_style = ParagraphStyle(
        "BTitle", parent=styles["Title"],
        fontSize=22, textColor=colors.white,
        spaceAfter=4, fontName="Helvetica-Bold",
    )
    subtitle_style = ParagraphStyle(
        "BSubtitle", parent=styles["Normal"],
        fontSize=10, textColor=colors.HexColor("#D1FAE5"),
        spaceAfter=0,
    )
    small_gray = ParagraphStyle(
        "SmGray", parent=styles["Normal"],
        fontSize=7, textColor=colors.HexColor("#6B7280"),
    )

    elements: list = []
    budget_name = budget_data.get("name") or "Presupuesto"
    budget_desc = budget_data.get("description") or ""
    created_raw = budget_data.get("created_at", "") or ""
    created_at = created_raw[:10] if created_raw else datetime.now().strftime("%Y-%m-%d")
    version = budget_data.get("version") or "1"

    # ── Header band ───────────────────────────────────────────────────────────
    header_table = Table(
        [[Paragraph(budget_name, title_style), Paragraph(f"Fecha: {created_at}  |  Versión {version}", subtitle_style)]],
        colWidths=[pw * 0.7, pw * 0.3],
    )
    header_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(_COLOR_HEADER)),
        ("TOPPADDING", (0, 0), (-1, -1), 14),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
        ("LEFTPADDING", (0, 0), (0, -1), 16),
        ("RIGHTPADDING", (-1, 0), (-1, -1), 16),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]))
    elements.append(header_table)

    # Accent bar
    accent_bar = Table([["  "]], colWidths=[pw])
    accent_bar.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(_COLOR_ACCENT)),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(accent_bar)
    elements.append(Spacer(1, 0.4 * cm))

    if with_description and budget_desc:
        elements.append(Paragraph(budget_desc, small_gray))
        elements.append(Spacer(1, 0.3 * cm))
    return elements


def pdf_detail_groups(all_items: list[dict]) -> list[tuple[dict | None, list[dict]]]:
    """Groups of the internal PDF's detail table: [(rubro or None, rows)].

    Each top-level section (rubro) gets every row below it at any depth, in tree order:
    the pisos (sections) as titles and their works, so the works listed add up to the
    grand total. A top-level work is its own group (no title), with whatever hangs from
    it. Rows whose parent is missing are treated as top level. Every row appears once.
    """
    ids = {i.get("id") for i in all_items}
    children: dict[object, list[dict]] = {}
    for i in all_items:
        pid = i.get("parent_id")
        children.setdefault(pid if pid in ids else None, []).append(i)

    seen: set = set()

    def below(node: dict) -> list[dict]:
        out: list[dict] = []
        for child in children.get(node.get("id"), []):
            if child.get("id") in seen:
                continue
            seen.add(child.get("id"))
            out.append(child)
            out.extend(below(child))
        return out

    groups: list[tuple[dict | None, list[dict]]] = []
    for top in children.get(None, []):
        if top.get("id") in seen:
            continue
        seen.add(top.get("id"))
        if is_section(top):
            groups.append((top, below(top)))
        else:
            groups.append((None, [top, *below(top)]))
    return groups


@router.get("/{budget_id}/export/pdf")
async def export_budget_pdf(
    budget_id: UUID,
    vista: Literal["interna", "cliente"] = Query("interna"),
    user: dict = Depends(get_current_user),
):
    """Export a budget to PDF with professional layout.

    ``vista=cliente``: the PDF for the client, with the sale price of each work and
    no internal costs (see client_pdf_data). Without it: the internal PDF.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    db = get_data_db()
    bid = str(budget_id)
    org_id = user["org_id"]

    # ── Fetch data ──────────────────────────────────────────────────────────
    budget_result = (
        db.table("budgets")
        .select("*")
        .eq("id", bid)
        .eq("org_id", org_id)
        .single()
        .execute()
    )
    if not budget_result.data:
        raise HTTPException(404, "Presupuesto no encontrado")
    budget_data = budget_result.data

    items_result = (
        db.table("budget_items")
        .select("*")
        .eq("budget_id", bid)
        .eq("org_id", org_id)
        .order("sort_order")
        .execute()
    )
    if not items_result.data:
        raise HTTPException(404, "El presupuesto no tiene trabajos")
    all_items = items_result.data

    # Indirect % of this budget (no error if missing — use defaults)
    cfg = _apply_cfg_defaults(budget_config(db, org_id, budget_data))

    if vista == "cliente":
        return _client_pdf_response(budget_data, all_items, cfg)

    # ── Totals: the saved sums (the same numbers as the editor) ─────────────
    totals = pdf_totals(all_items, cfg)
    mat_total = totals["mat_total"]
    mo_total = totals["mo_total"]
    directo_total = totals["directo_total"]
    indirecto_total = totals["indirecto_total"]
    beneficio_total = totals["beneficio_total"]
    impuestos_total = totals["impuestos_total"]
    neto_display = totals["neto_total"]
    iva_total = totals["iva_total"]
    total_final_display = totals["total_final"]
    items_count = totals["items_count"]

    # ── Build PDF ─────────────────────────────────────────────────────────────
    output = BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=budget_data.get("name") or "Presupuesto",
        author="Presupuestador SOLE",
    )

    # ── Styles ─────────────────────────────────────────────────────────────────
    styles = getSampleStyleSheet()
    C_HEADER = colors.HexColor(_COLOR_HEADER)
    C_ACCENT = colors.HexColor(_COLOR_ACCENT)
    C_SECTION_BG = colors.HexColor(_COLOR_SECTION)
    C_SECTION_TXT = colors.HexColor(_COLOR_SECTION_TEXT)
    C_BORDER = colors.HexColor(_COLOR_BORDER)
    C_ROW_ALT = colors.HexColor(_COLOR_ROW_ALT)
    C_TOTAL = colors.HexColor(_COLOR_TOTAL_ROW)

    section_label_style = ParagraphStyle(
        "SLabel", parent=styles["Normal"],
        fontSize=9, fontName="Helvetica-Bold",
        textColor=C_SECTION_TXT,
    )
    section_heading_style = ParagraphStyle(
        "SHeading", parent=styles["Heading2"],
        fontSize=12, fontName="Helvetica-Bold",
        textColor=C_HEADER, spaceAfter=6, spaceBefore=14,
    )
    cell_style = ParagraphStyle(
        "Cell", parent=styles["Normal"],
        fontSize=7, leading=9,
    )
    cell_bold_style = ParagraphStyle(
        "CellBold", parent=styles["Normal"],
        fontSize=7, leading=9, fontName="Helvetica-Bold",
    )
    small_gray = ParagraphStyle(
        "SmGray", parent=styles["Normal"],
        fontSize=7, textColor=colors.HexColor("#6B7280"),
    )

    # ── Page width for layout ──────────────────────────────────────────────────
    pw = page_size[0] - 3 * cm  # usable page width (landscape A4 ~ 297mm - margins)

    elements: list = []

    # ══════════════════════════════════════════════════════════════════════════
    #  PAGE 1: COVER / SUMMARY
    # ══════════════════════════════════════════════════════════════════════════

    elements += _pdf_header(budget_data, pw, styles)

    # ── KPI boxes (4 columns) ─────────────────────────────────────────────────
    kpi_width = pw / 4

    def kpi_cell(label: str, value: str) -> list:
        return [
            Paragraph(label, ParagraphStyle(
                "KpiLabel", parent=styles["Normal"],
                fontSize=8, textColor=colors.HexColor("#6B7280"), fontName="Helvetica",
            )),
            Paragraph(value, ParagraphStyle(
                "KpiVal", parent=styles["Normal"],
                fontSize=13, fontName="Helvetica-Bold",
                textColor=C_HEADER, spaceAfter=0,
            )),
        ]

    kpi_data = [
        kpi_cell("Trabajos", str(items_count)),
        kpi_cell("Costo Directo", _fmt_ars(directo_total)),
        kpi_cell("Indirectos + Beneficio + Impuestos",
                 _fmt_ars(indirecto_total + beneficio_total + impuestos_total)),
        kpi_cell("PRECIO SIN IVA", _fmt_ars(neto_display)),
    ]

    kpi_rows_mat = []
    for kpi in kpi_data:
        cell_t = Table([[kpi[0]], [kpi[1]]], colWidths=[kpi_width - 0.6 * cm])
        cell_t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BOX", (0, 0), (-1, -1), 1, C_ACCENT),
            ("ROUNDEDCORNERS", [4, 4, 4, 4]),
        ]))
        kpi_rows_mat.append(cell_t)

    kpi_table = Table([kpi_rows_mat], colWidths=[kpi_width] * 4)
    kpi_table.setStyle(TableStyle([
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 0.5 * cm))

    # ── Summary table (cost breakdown) ────────────────────────────────────────
    elements.append(Paragraph("Resumen de Costos", section_heading_style))

    summary_rows = [
        ["Concepto", "Monto"],
        ["Total Materiales", _fmt_ars(mat_total)],
        ["Total Mano de Obra", _fmt_ars(mo_total)],
        ["Costo Directo (01)", _fmt_ars(directo_total)],
        [f"Costos Indirectos ({_fmt_pct(totals['total_ind_pct'])})", _fmt_ars(indirecto_total)],
        [f"Beneficio ({_fmt_pct(totals['ben_pct'])})", _fmt_ars(beneficio_total)],
        ["Impuestos (IIBB + cheque)", _fmt_ars(impuestos_total)],
        ["PRECIO SIN IVA", _fmt_ars(neto_display)],
        [f"IVA ({_fmt_pct(totals['iva_pct'])})", _fmt_ars(iva_total)],
        ["PRECIO CON IVA", _fmt_ars(total_final_display)],
    ]

    col_s1 = 9 * cm
    col_s2 = 5.5 * cm
    sum_table = Table(summary_rows, colWidths=[col_s1, col_s2])
    sum_table.setStyle(TableStyle([
        # Header row
        ("BACKGROUND", (0, 0), (-1, 0), C_HEADER),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        # Alternating rows
        ("BACKGROUND", (0, 2), (-1, 2), C_ROW_ALT),
        ("BACKGROUND", (0, 4), (-1, 4), C_ROW_ALT),
        ("BACKGROUND", (0, 6), (-1, 6), C_ROW_ALT),
        ("BACKGROUND", (0, 8), (-1, 8), C_ROW_ALT),
        # Price without IVA
        ("BACKGROUND", (0, 7), (-1, 7), C_HEADER),
        ("TEXTCOLOR", (0, 7), (-1, 7), colors.white),
        ("FONTNAME", (0, 7), (-1, 7), "Helvetica-Bold"),
        # Last row (price with IVA)
        ("BACKGROUND", (0, -1), (-1, -1), C_HEADER),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        # Alignment
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.5, C_BORDER),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(sum_table)
    elements.append(Spacer(1, 1.2 * cm))

    # ── Signature lines ────────────────────────────────────────────────────────
    sig_style = ParagraphStyle(
        "Sig", parent=styles["Normal"],
        fontSize=9, textColor=colors.HexColor("#374151"),
    )
    sig_line = "_" * 38
    sig_row = [
        [Paragraph(f"{sig_line}<br/>Preparado por", sig_style)],
        [Paragraph(f"{sig_line}<br/>Aprobado por", sig_style)],
    ]
    sig_table = Table(sig_row, colWidths=[pw * 0.45, pw * 0.45])
    sig_table.setStyle(TableStyle([
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    elements.append(sig_table)

    # ══════════════════════════════════════════════════════════════════════════
    #  PAGE 2+: DETAIL TABLE grouped by section
    # ══════════════════════════════════════════════════════════════════════════

    elements.append(PageBreak())
    elements.append(Paragraph("Detalle de trabajos", section_heading_style))

    sections = pdf_detail_groups(all_items)

    # Column layout for detail table
    # Codigo | Descripcion | Unidad | Cantidad | P.Unit MAT | P.Unit MO | Directo | Indirecto | Beneficio
    # | Impuestos | Precio sin IVA  (each row adds up: directo + indirecto + beneficio + impuestos)
    COL_W = [1.6*cm, 5.4*cm, 1.2*cm, 1.7*cm, 2.1*cm, 2.1*cm, 2.4*cm, 2.2*cm, 2.2*cm, 2.2*cm, 2.5*cm]
    detail_header = [
        "Código", "Descripción", "Unid.", "Cantidad",
        "Materiales\npor unidad", "Mano de obra\npor unidad", "Directo", "Indirecto", "Beneficio", "Impuestos", "Precio\nsin IVA",
    ]

    table_rows: list = [detail_header]
    # Track row indices for section rows styling
    section_row_indices: list[int] = []
    piso_row_indices: list[int] = []
    subtotal_row_indices: list[int] = []

    row_idx = 1  # 0 is header

    for sec_item, children in sections:
        if sec_item is not None:
            # Section header row
            sec_code = sec_item.get("code") or ""
            sec_desc = sec_item.get("description") or "—"
            label = f"{sec_code}  {sec_desc}".strip() if sec_code else sec_desc
            table_rows.append([
                Paragraph(label, ParagraphStyle(
                    "SecRow", parent=styles["Normal"],
                    fontSize=8, fontName="Helvetica-Bold", textColor=C_SECTION_TXT,
                )),
                "", "", "", "", "", "", "", "", "", "",
            ])
            section_row_indices.append(row_idx)
            row_idx += 1

        sec_directo = 0.0
        sec_indirecto = 0.0
        sec_beneficio = 0.0
        sec_impuestos = 0.0
        sec_neto = 0.0

        for item in children:
            if is_section(item):
                # A piso (or any section below the rubro): a title row, its works follow
                sub_code = item.get("code") or ""
                sub_desc = item.get("description") or "—"
                table_rows.append([
                    Paragraph(f"{sub_code}  {sub_desc}".strip() if sub_code else sub_desc, cell_bold_style),
                    "", "", "", "", "", "", "", "", "", "",
                ])
                piso_row_indices.append(row_idx)
                row_idx += 1
                continue
            cantidad = item.get("cantidad")
            cantidad_str = _fmt_ars(float(cantidad)).replace("$ ", "").replace(",00", "") if cantidad else "—"
            # Simplified: just show the number without $ for quantity
            try:
                cantidad_str = f"{float(cantidad):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if cantidad else "—"
            except Exception:
                cantidad_str = str(cantidad) if cantidad else "—"

            d = float(item.get("directo_total") or 0)
            ind = float(item.get("indirecto_total") or 0)
            ben = float(item.get("beneficio_total") or 0)
            imp = float(item.get("impuestos_total") or 0)
            neto = float(item.get("neto_total") or 0)

            sec_directo += d
            sec_indirecto += ind
            sec_beneficio += ben
            sec_impuestos += imp
            sec_neto += neto

            desc_para = Paragraph(item.get("description") or "—", cell_style)
            table_rows.append([
                item.get("code") or "",
                desc_para,
                item.get("unidad") or "",
                cantidad_str,
                _fmt_ars(float(item.get("mat_unitario") or 0)),
                _fmt_ars(float(item.get("mo_unitario") or 0)),
                _fmt_ars(d),
                _fmt_ars(ind),
                _fmt_ars(ben),
                _fmt_ars(imp),
                _fmt_ars(neto),
            ])
            row_idx += 1

        # Section subtotal row (only if there was a section header)
        if sec_item is not None and children:
            table_rows.append([
                "",
                Paragraph(f"Subtotal {sec_item.get('description') or ''}", cell_bold_style),
                "", "", "", "",
                _fmt_ars(sec_directo),
                _fmt_ars(sec_indirecto),
                _fmt_ars(sec_beneficio),
                _fmt_ars(sec_impuestos),
                _fmt_ars(sec_neto),
            ])
            subtotal_row_indices.append(row_idx)
            row_idx += 1

    # Grand total row
    table_rows.append([
        "TOTAL", "", "", "", "", "",
        _fmt_ars(directo_total),
        _fmt_ars(indirecto_total),
        _fmt_ars(beneficio_total),
        _fmt_ars(impuestos_total),
        _fmt_ars(neto_display),
    ])
    grand_total_row_idx = row_idx

    detail_table = Table(table_rows, colWidths=COL_W, repeatRows=1)
    # Base style
    ts = [
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), C_HEADER),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 7),
        # Body
        ("FONTSIZE", (0, 1), (-1, -2), 7),
        ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
        # Alignment: numeric columns right
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (2, -1), "LEFT"),
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.3, C_BORDER),
        # Padding
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        # Grand total row
        ("BACKGROUND", (0, grand_total_row_idx), (-1, grand_total_row_idx), C_TOTAL),
        ("TEXTCOLOR", (0, grand_total_row_idx), (-1, grand_total_row_idx), colors.white),
        ("FONTNAME", (0, grand_total_row_idx), (-1, grand_total_row_idx), "Helvetica-Bold"),
        ("FONTSIZE", (0, grand_total_row_idx), (-1, grand_total_row_idx), 8),
    ]
    # Alternating rows (odd body rows)
    for i in range(1, len(table_rows) - 1, 2):
        if i not in section_row_indices and i not in subtotal_row_indices and i not in piso_row_indices:
            ts.append(("BACKGROUND", (0, i), (-1, i), C_ROW_ALT))
    # Piso rows (sections inside a rubro)
    for i in piso_row_indices:
        ts.append(("SPAN", (0, i), (-1, i)))
        ts.append(("LINEBELOW", (0, i), (-1, i), 0.5, C_ACCENT))
    # Section rows
    for i in section_row_indices:
        ts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(_COLOR_SECTION)))
        ts.append(("SPAN", (0, i), (-1, i)))
        ts.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
    # Subtotal rows
    for i in subtotal_row_indices:
        ts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F3F4F6")))
        ts.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
        ts.append(("LINEABOVE", (0, i), (-1, i), 0.5, C_ACCENT))

    detail_table.setStyle(TableStyle(ts))
    elements.append(detail_table)

    # ══════════════════════════════════════════════════════════════════════════
    #  LAST SECTION: CASCADE SUMMARY
    # ══════════════════════════════════════════════════════════════════════════

    elements.append(PageBreak())
    elements.append(Paragraph("Cascada de Costos", section_heading_style))

    t = totals  # saved sums; % from the budget's config
    ind, pcts = t["indirectos"], t["ind_pcts"]
    cascade_rows = [
        ["Concepto", "%", "Monto"],
        ["Subtotal 01 — Costos Directos", "", _fmt_ars(t["directo_total"])],
        ["+ Imprevistos", _fmt_pct(pcts["imprevistos_pct"]), _fmt_ars(ind["imprevistos_pct"])],
        ["+ Estructura", _fmt_pct(pcts["estructura_pct"]), _fmt_ars(ind["estructura_pct"])],
        ["+ Jefatura de Obra", _fmt_pct(pcts["jefatura_pct"]), _fmt_ars(ind["jefatura_pct"])],
        ["+ Logística", _fmt_pct(pcts["logistica_pct"]), _fmt_ars(ind["logistica_pct"])],
        ["+ Herramientas", _fmt_pct(pcts["herramientas_pct"]), _fmt_ars(ind["herramientas_pct"])],
        ["= Subtotal 02 (con Indirectos)", _fmt_pct(t["total_ind_pct"]), _fmt_ars(t["subtotal_02"])],
        ["+ Beneficio", _fmt_pct(t["ben_pct"]), _fmt_ars(t["beneficio_total"])],
        ["= Subtotal 03 (con Beneficio)", "", _fmt_ars(t["subtotal_03"])],
        ["+ Ingresos Brutos", _fmt_pct(t["iibb_pct"]), _fmt_ars(t["iibb"])],
        ["+ Impuesto al Cheque", _fmt_pct(t["cheque_pct"]), _fmt_ars(t["cheque"])],
        ["= PRECIO SIN IVA", "", _fmt_ars(t["neto_total"])],
        ["+ IVA", _fmt_pct(t["iva_pct"]), _fmt_ars(t["iva_total"])],
        ["= PRECIO CON IVA", "", _fmt_ars(t["total_final"])],
    ]

    # Row indices that are subtotals or totals (bold + background)
    cascade_bold_rows = {0, 7, 9, 12, 14}   # header, subtotal 02, 03, neto, total
    cascade_total_rows = {12, 14}             # neto, total final

    col_casc = [10 * cm, 2.5 * cm, 5 * cm]
    casc_table = Table(cascade_rows, colWidths=col_casc)
    ts_c = [
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), C_HEADER),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        # Body
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        # Alignment
        ("ALIGN", (1, 0), (2, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.4, C_BORDER),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        # Subtotal row (subtotal 02 = index 7)
        ("BACKGROUND", (0, 7), (-1, 7), colors.HexColor(_COLOR_SECTION)),
        ("FONTNAME", (0, 7), (-1, 7), "Helvetica-Bold"),
        ("LINEABOVE", (0, 7), (-1, 7), 1, C_ACCENT),
        # Subtotal 03 = index 9
        ("BACKGROUND", (0, 9), (-1, 9), colors.HexColor(_COLOR_SECTION)),
        ("FONTNAME", (0, 9), (-1, 9), "Helvetica-Bold"),
        ("LINEABOVE", (0, 9), (-1, 9), 1, C_ACCENT),
        # NETO = index 12
        ("BACKGROUND", (0, 12), (-1, 12), C_HEADER),
        ("TEXTCOLOR", (0, 12), (-1, 12), colors.white),
        ("FONTNAME", (0, 12), (-1, 12), "Helvetica-Bold"),
        # TOTAL FINAL = index 14
        ("BACKGROUND", (0, 14), (-1, 14), C_HEADER),
        ("TEXTCOLOR", (0, 14), (-1, 14), colors.white),
        ("FONTNAME", (0, 14), (-1, 14), "Helvetica-Bold"),
        ("FONTSIZE", (0, 14), (-1, 14), 11),
        # Alternating plain rows
        ("BACKGROUND", (0, 2), (-1, 2), C_ROW_ALT),
        ("BACKGROUND", (0, 4), (-1, 4), C_ROW_ALT),
        ("BACKGROUND", (0, 6), (-1, 6), C_ROW_ALT),
        ("BACKGROUND", (0, 10), (-1, 10), C_ROW_ALT),
    ]
    casc_table.setStyle(TableStyle(ts_c))
    elements.append(casc_table)
    elements.append(Spacer(1, 0.8 * cm))

    # Nota al pie de la cascada
    elements.append(Paragraph(
        "Nota: Los porcentajes de indirectos y beneficio se aplican en cascada conforme al modelo Excel.",
        small_gray,
    ))

    # ── Build and stream ───────────────────────────────────────────────────────
    doc.build(elements, onFirstPage=_build_page_footer, onLaterPages=_build_page_footer)
    output.seek(0)

    safe_name = (budget_data.get("name") or "presupuesto").replace(" ", "_").replace("/", "-")
    filename = f"{safe_name}_presupuesto.pdf"

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── PDF para el cliente ──────────────────────────────────────────────────────


def _fmt_cantidad(value: object) -> str:
    """1234.5 → '1.234,50' (Argentine format, 2 decimals)."""
    try:
        return f"{float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return str(value or "—")


def _top_section(item: dict, by_id: dict) -> tuple[dict | None, dict | None]:
    """(rubro, piso) of an item: its top-level section and the section right above it, if different."""
    chain: list[dict] = []
    seen: set = set()
    parent = by_id.get(item.get("parent_id"))
    while parent is not None and parent.get("id") not in seen:
        seen.add(parent.get("id"))
        chain.append(parent)
        parent = by_id.get(parent.get("parent_id"))
    sections = [p for p in chain if is_section(p)]
    if not sections:
        return None, None
    rubro = sections[-1]
    piso = sections[0] if sections[0] is not rubro else None
    return rubro, piso


def client_pdf_data(all_items: list[dict], cfg: dict) -> dict:
    """Rows of the client PDF: per rubro, each work with its sale price; then the totals.

    Each work's price is its saved price without IVA (``neto_total``, the same number as
    the editor). Total sin IVA = Σ neto of the works (= calc_budget_summary), IVA = Σ IVA,
    Total con IVA = Σ total_final; works saved without IVA get the IVA step of the cascade
    over their neto (sale_totals). Nothing else of the cascade (direct cost, indirects,
    profit, percentages) is returned.

    The rows are the works (not "Seccion") with a quantity or a price: a work with
    cantidad 0 and price 0 is not shown (it adds nothing). Rubros are the top-level
    "Seccion" rows, with every work below them (also the ones under a piso); works without
    a rubro go in "Otros trabajos" (or in a group without title when there is no rubro at all).
    """
    iva_pct = float(cfg.get("iva_pct") if cfg.get("iva_pct") is not None else 21)
    works = [i for i in all_items if not is_section(i)]
    sales = [sale_totals(i, iva_pct) for i in works]
    leaf_items, cents = [], []
    for item, (neto, _, _) in zip(works, sales):
        if float(item.get("cantidad") or 0) > 0 or round(neto * 100):
            leaf_items.append(item)
            cents.append(round(neto * 100))

    by_id = {i.get("id"): i for i in all_items}
    grupos: dict[object, dict] = {}
    for item, precio_cents in zip(leaf_items, cents):
        rubro, piso = _top_section(item, by_id)
        key = rubro.get("id") if rubro else None
        if key not in grupos:
            titulo = None
            if rubro:
                code, desc = rubro.get("code") or "", rubro.get("description") or "—"
                titulo = f"{code}  {desc}".strip() if code else desc
            grupos[key] = {"titulo": titulo, "filas": [], "_cents": 0}
        grupo = grupos[key]
        cantidad = float(item.get("cantidad") or 0)
        total = precio_cents / 100
        grupo["filas"].append({
            "code": item.get("code") or "",
            "descripcion": item.get("description") or "—",
            "unidad": item.get("unidad") or "",
            "cantidad": cantidad,
            "precio_unitario": round(total / cantidad, 2) if cantidad else total,
            "total": total,
            "piso": piso.get("description") if piso else None,
        })
        grupo["_cents"] += precio_cents

    con_rubros = any(key is not None for key in grupos)
    rubros = []
    for grupo in grupos.values():
        titulo = grupo["titulo"] or ("Otros trabajos" if con_rubros else None)
        rubros.append({"titulo": titulo, "filas": grupo["filas"], "subtotal": grupo.pop("_cents") / 100})

    return {
        "rubros": rubros,
        "total_sin_iva": sum(round(neto * 100) for neto, _, _ in sales) / 100,
        "iva_pct": iva_pct,
        "iva": sum(round(iva * 100) for _, iva, _ in sales) / 100,
        "total_con_iva": sum(round(total * 100) for _, _, total in sales) / 100,
    }


def _client_pdf_response(budget_data: dict, all_items: list[dict], cfg: dict) -> StreamingResponse:
    """Client PDF: same header as the internal one, one table per rubro, totals with IVA."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    data = client_pdf_data(all_items, cfg)

    output = BytesIO()
    page_size = A4
    doc = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=budget_data.get("name") or "Presupuesto",
        author="Presupuestador SOLE",
    )
    styles = getSampleStyleSheet()
    C_HEADER = colors.HexColor(_COLOR_HEADER)
    C_ACCENT = colors.HexColor(_COLOR_ACCENT)
    C_BORDER = colors.HexColor(_COLOR_BORDER)
    C_ROW_ALT = colors.HexColor(_COLOR_ROW_ALT)
    section_heading_style = ParagraphStyle(
        "SHeading", parent=styles["Heading2"],
        fontSize=12, fontName="Helvetica-Bold",
        textColor=C_HEADER, spaceAfter=6, spaceBefore=14,
    )
    rubro_style = ParagraphStyle(
        "Rubro", parent=styles["Heading3"],
        fontSize=10, fontName="Helvetica-Bold",
        textColor=colors.HexColor(_COLOR_SECTION_TEXT), spaceAfter=4, spaceBefore=10,
    )
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)
    piso_style = ParagraphStyle("Piso", parent=styles["Normal"], fontSize=8, leading=10,
                                fontName="Helvetica-Bold", textColor=colors.HexColor(_COLOR_SECTION_TEXT))
    cell_bold_style = ParagraphStyle("CellBold", parent=styles["Normal"], fontSize=8, leading=10,
                                     fontName="Helvetica-Bold")

    pw = page_size[0] - 3 * cm
    elements: list = _pdf_header(budget_data, pw, styles, with_description=False)
    elements.append(Paragraph("Presupuesto", section_heading_style))

    # Ítem | Trabajo | Unidad | Cantidad | Precio unitario | Total
    col_w = [1.5 * cm, pw - 1.5 * cm - 1.4 * cm - 2.0 * cm - 2.8 * cm - 3.0 * cm, 1.4 * cm, 2.0 * cm, 2.8 * cm, 3.0 * cm]
    header = ["Código", "Trabajo", "Unidad", "Cantidad", "Precio unitario", "Total"]

    for rubro in data["rubros"]:
        rows: list = [header]
        piso_rows: list[int] = []
        piso_actual = None
        for fila in rubro["filas"]:
            if fila["piso"] and fila["piso"] != piso_actual:
                piso_actual = fila["piso"]
                rows.append([Paragraph(piso_actual, piso_style), "", "", "", "", ""])
                piso_rows.append(len(rows) - 1)
            rows.append([
                fila["code"],
                Paragraph(fila["descripcion"], cell_style),
                fila["unidad"],
                _fmt_cantidad(fila["cantidad"]),
                _fmt_ars(fila["precio_unitario"]),
                _fmt_ars(fila["total"]),
            ])
        if rubro["titulo"]:
            rows.append(["", Paragraph(f"Subtotal {rubro['titulo']}", cell_bold_style),
                         "", "", "", _fmt_ars(rubro["subtotal"])])
        table = Table(rows, colWidths=col_w, repeatRows=1)
        ts = [
            ("BACKGROUND", (0, 0), (-1, 0), C_HEADER),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (2, -1), "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.3, C_BORDER),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        if rubro["titulo"]:  # subtotal row
            ts += [
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F4F6")),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("LINEABOVE", (0, -1), (-1, -1), 0.5, C_ACCENT),
            ]
        for i in range(1, len(rows) - (1 if rubro["titulo"] else 0), 2):
            if i not in piso_rows:
                ts.append(("BACKGROUND", (0, i), (-1, i), C_ROW_ALT))
        for i in piso_rows:
            ts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(_COLOR_SECTION)))
            ts.append(("SPAN", (0, i), (-1, i)))
        table.setStyle(TableStyle(ts))
        if rubro["titulo"]:
            elements.append(KeepTogether([Paragraph(rubro["titulo"], rubro_style), table]))
        else:
            elements.append(Spacer(1, 0.3 * cm))
            elements.append(table)

    totales = [
        ["Total sin IVA", _fmt_ars(data["total_sin_iva"])],
        [f"IVA ({_fmt_pct(data['iva_pct'])})", _fmt_ars(data["iva"])],
        ["Total con IVA", _fmt_ars(data["total_con_iva"])],
    ]
    tot_table = Table(totales, colWidths=[6 * cm, 4.5 * cm], hAlign="RIGHT")
    tot_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, C_BORDER),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, -1), (-1, -1), C_HEADER),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, -1), (-1, -1), 11),
    ]))
    elements.append(Spacer(1, 0.6 * cm))
    elements.append(KeepTogether([tot_table]))

    doc.build(elements, onFirstPage=_build_page_footer, onLaterPages=_build_page_footer)
    output.seek(0)

    safe_name = (budget_data.get("name") or "presupuesto").replace(" ", "_").replace("/", "-")
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}_presupuesto_cliente.pdf"'},
    )
