"""Planilla Terrac: the budget written as Sol's Excel (PLAN_EXPORTAR 2).

Sheets:
  - ``01_C&P``: her 26 columns (A..Z) with her header (rows 1-7), rubros and pisos as title
    rows and one row per work from row 8, in the layout ``app/obra_import.parse_obra``
    reads, so the file can go back through "Cargar obra" (same works, quantities and, from the
"Coeficiente de pase" sheet, the obra's percentages; works with a formula are priced again).
  - One sheet per work with resources, like her detail sheets ("3.1-1"): a block per
    kind of resource and the work's direct cost at the bottom.
  - ``Coeficiente de pase``: the % of the budget and the cascade of its totals.

Values, not formulas: a file written by the app has no cached results, so anything that
reads values (Cargar obra, pandas, a preview) would see empty cells. Every amount is the
saved one of each work (the same numbers as the editor): N = directo_total and
Z = neto_total to the cent.

Pure functions (no database): the router reads the budget and passes it in.
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.calculations import cascade_factors, is_section, sale_totals
from app.tree import safe_float

SHEET_CP = "01_C&P"
SHEET_COEF = "Coeficiente de pase"
FIRST_ROW = 8  # obra_import.FIRST_ROW
LAST_COL = 26  # Z

# Columns of 01_C&P (1 = A); the ones obra_import reads are A, B, C, D, E, J, N, Z
C_ITEM, C_DESC, C_UNIDAD, C_CANT = 1, 2, 3, 4
C_MAT, C_JORNALES, C_EQUIPOS, C_MO_MAT, C_SUBC, C_MO, C_UNIT = 5, 6, 7, 8, 9, 10, 11  # E..K
C_L, C_M, C_N = 12, 13, 14          # directo general: MAT, M.O., GENERAL
C_O, C_P, C_Q = 15, 16, 17          # gastos indirectos
C_R, C_S, C_T = 18, 19, 20          # beneficio e impuestos
C_U, C_V, C_W = 21, 22, 23          # total neto por unidad
C_X, C_Y, C_Z = 24, 25, 26          # total neto general

# Resource kind → unit column of 01_C&P (material is E; the rest add up to J, "M.O.")
TIPO_COL = {"material": C_MAT, "mano_obra": C_JORNALES, "equipo": C_EQUIPOS,
            "mo_material": C_MO_MAT, "subcontrato": C_SUBC}
# Blocks of a work's sheet, with Sol's titles (Cargar obra recognizes them)
TIPO_TITULO = (
    ("material", "MATERIALES"),
    ("mano_obra", "MANO DE OBRA - PERSONAS"),
    ("equipo", "MANO DE OBRA - EQUIPOS"),
    ("mo_material", "MANO DE OBRA - MATERIALES"),
    ("subcontrato", "MANO DE OBRA - SUBCONTRATOS"),
)

# Sol's peso format (accounting, $ es-AR) and quantities
FMT_PESOS = '_-[$$-2C0A]\\ * #,##0.00_-;\\-[$$-2C0A]\\ * #,##0.00_-;_-[$$-2C0A]\\ * "-"??_-;_-@'
FMT_CANT = "#,##0.00"
FMT_PCT = "0.0#%"

_THIN = Side(style="thin", color="FF808080")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BLACK = PatternFill("solid", fgColor="FF000000")
_RUBRO = PatternFill("solid", fgColor="FF404040")
_PISO = PatternFill("solid", fgColor="FF999999")
_HEAD = PatternFill("solid", fgColor="FFF2F2F2")
_TOTAL = PatternFill("solid", fgColor="FFD9D9D9")
_SOFT = PatternFill("solid", fgColor="FFF3F3F3")
_WHITE_B = Font(bold=True, color="FFFFFFFF")
_BOLD = Font(bold=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_SECTION_TEXT = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*[-.]?\s+(.*)$")  # obra_import._SECTION_RE


def _f(value: object) -> float:
    return safe_float(value) or 0.0


def _cents(value: float) -> int:
    return round(value * 100)


def split_two(total: float, a: float, b: float) -> tuple[float, float]:
    """Split ``total`` between MAT and M.O. in proportion to (a, b), in cents, adding up exactly.

    With no weight (a direct cost of 0) everything goes to M.O.
    """
    cents = _cents(total)
    a, b = max(a, 0.0), max(b, 0.0)
    if a + b <= 0:
        return 0.0, cents / 100
    first = round(cents * a / (a + b))
    return first / 100, (cents - first) / 100


def _split_cents(total: float, weights: list[float]) -> list[float]:
    """``total`` split by weights, in cents, adding up exactly (all to the first without weights)."""
    cents = _cents(total)
    peso = sum(w for w in weights if w > 0)
    if peso <= 0:
        return [cents / 100] + [0.0] * (len(weights) - 1)
    parts = [round(cents * max(w, 0) / peso) for w in weights]
    biggest = max(range(len(weights)), key=lambda i: weights[i])
    parts[biggest] += cents - sum(parts)
    return [p / 100 for p in parts]


# ── Values of one work in 01_C&P ────────────────────────────────────────────


def work_values(item: dict, resources: list[dict]) -> dict[int, float]:
    """Columns E..Z of one work (1 = A), from its saved numbers.

    E = mat_unitario; F..I = mo_unitario split by the kinds of its resources (jornales,
    equipos, materiales de M.O., subcontratos; what the client buys does not count);
    without resources everything goes to F. L = mat_total, N = directo_total and
    M = N − L, so N is the saved direct cost to the cent. Indirects (O..Q),
    beneficio + impuestos (R..T) and the net (X..Z, Z = neto_total) split between MAT
    and M.O. in proportion to L and M; U..W are X..Z per unit.
    """
    cantidad = _f(item.get("cantidad"))
    mat_unit = round(_f(item.get("mat_unitario")), 2)
    mo_unit = round(_f(item.get("mo_unitario")), 2)

    weights = {tipo: 0.0 for tipo in ("mano_obra", "equipo", "mo_material", "subcontrato")}
    for r in resources:
        if r.get("tipo") in weights and not r.get("lo_compra_cliente"):
            weights[r["tipo"]] += _f(r.get("subtotal"))
    f, g, h, i = _split_cents(mo_unit, list(weights.values()))

    directo = item.get("directo_total")
    mat_total = item.get("mat_total")
    l_ = round(_f(mat_total), 2) if mat_total is not None else round(cantidad * mat_unit, 2)
    if directo is None:
        mo_total = item.get("mo_total")
        m_ = round(_f(mo_total), 2) if mo_total is not None else round(cantidad * mo_unit, 2)
        n_ = round(l_ + m_, 2)
    else:
        n_ = round(_f(directo), 2)
        m_ = round(n_ - l_, 2)

    indirecto = round(_f(item.get("indirecto_total")), 2)
    ben_imp = round(_f(item.get("beneficio_total")) + _f(item.get("impuestos_total")), 2)
    neto = round(sale_totals(item)[0], 2)
    o_, p_ = split_two(indirecto, l_, m_)
    r_, s_ = split_two(ben_imp, l_, m_)
    x_, y_ = split_two(neto, l_, m_)
    if cantidad:
        w_ = round(neto / cantidad, 2)
        u_ = round(x_ / cantidad, 2)
        v_ = round(w_ - u_, 2)
    else:
        u_ = v_ = w_ = 0.0

    return {
        C_MAT: mat_unit, C_JORNALES: f, C_EQUIPOS: g, C_MO_MAT: h, C_SUBC: i,
        C_MO: mo_unit, C_UNIT: round(mat_unit + mo_unit, 2),
        C_L: l_, C_M: m_, C_N: n_,
        C_O: o_, C_P: p_, C_Q: indirecto,
        C_R: r_, C_S: s_, C_T: ben_imp,
        C_U: u_, C_V: v_, C_W: w_,
        C_X: x_, C_Y: y_, C_Z: neto,
    }


# ── Tree → rows of 01_C&P ────────────────────────────────────────────────────


def _section_name(item: dict, code: str) -> str:
    """Name of a rubro / piso without a code it may already carry ("3.1- FUNDACIONES")."""
    name = " ".join(str(item.get("description") or "").split())
    m = _SECTION_TEXT.match(name)
    if m and m.group(1).rstrip(".") == code:
        name = m.group(2).strip()
    return name or "Sin nombre"


def cp_rows(items: list[dict]) -> list[dict]:
    """Rows of 01_C&P in tree order: {"kind": rubro/piso/work, "item", "code"}.

    Top-level sections are rubros ("1- TAREAS PRELIMINARES"); every section below them is
    a piso ("3.1- FUNDACIONES"), the two levels Cargar obra reads. Section codes are the
    saved ones when they fit that reading (rubro: "3", piso: "3.1" inside rubro 3);
    otherwise they are taken from their works' codes or numbered in order.
    Inside each section (and at the top) its own works go before its sub-sections:
    Cargar obra hangs a work from the last title above it. Every item appears once.
    """
    ids = {i.get("id") for i in items}
    children: dict[object, list[dict]] = {}
    for i in items:
        pid = i.get("parent_id")
        children.setdefault(pid if pid in ids and pid != i.get("id") else None, []).append(i)

    seen: set = set()
    rows: list[dict] = []
    rubro_codes: set[str] = set()
    piso_codes: set[str] = set()

    def fresh(node_id: object) -> list[dict]:
        out = []
        for child in children.get(node_id, []):
            if child.get("id") not in seen:
                seen.add(child.get("id"))
                out.append(child)
        return out

    def work_codes(section: dict, deep: bool = False) -> list[str]:
        """Codes of the works below a section (directly, or at any depth)."""
        out: list[str] = []
        stack, visited = list(children.get(section.get("id"), [])), set()
        while stack:
            node = stack.pop(0)
            if node.get("id") in visited:
                continue
            visited.add(node.get("id"))
            if not is_section(node):
                out.append(str(node.get("code") or "").strip())
            if deep or not is_section(node):
                stack.extend(children.get(node.get("id"), []))
        return out

    def pick_code(own: str, pattern: str, used: set[str], hints: list[str]) -> str | None:
        for candidate in [own, *hints]:
            if candidate and re.fullmatch(pattern, candidate) and candidate not in used:
                return candidate
        return None

    def section_row(s: dict, rubro_code: str | None) -> None:
        own = str(s.get("code") or "").strip().rstrip(".")
        if rubro_code is None:
            code = pick_code(own, r"\d+", rubro_codes, [c.split(".")[0] for c in work_codes(s, deep=True)])
            n = 1
            while code is None:
                code = str(n) if str(n) not in rubro_codes else None
                n += 1
            rubro_codes.add(code)
            rows.append({"kind": "rubro", "item": s, "code": code})
            emit(s.get("id"), code)
        else:
            hints = [".".join(c.split(".")[:2]) for c in work_codes(s) if c.count(".") >= 2]
            code = pick_code(own, re.escape(rubro_code) + r"\.\d+", piso_codes, hints)
            n = 1
            while code is None:
                code = f"{rubro_code}.{n}" if f"{rubro_code}.{n}" not in piso_codes else None
                n += 1
            piso_codes.add(code)
            rows.append({"kind": "piso", "item": s, "code": code})
            emit(s.get("id"), rubro_code)

    def emit(node_id: object, rubro_code: str | None) -> None:
        works: list[dict] = []
        sections: list[dict] = []

        def collect(nid: object) -> None:
            # Works hanging from a work (an imported 1.1 → 1.1.1) follow it
            for kid in fresh(nid):
                if is_section(kid):
                    sections.append(kid)
                else:
                    works.append(kid)
                    collect(kid.get("id"))

        collect(node_id)
        rows.extend({"kind": "work", "item": w} for w in works)
        for s in sections:
            section_row(s, rubro_code)

    emit(None, None)
    # Anything left (a cycle of parent ids) goes at the end
    for i in items:
        if i.get("id") not in seen:
            seen.add(i.get("id"))
            if is_section(i):
                section_row(i, None)
            else:
                rows.append({"kind": "work", "item": i})
    return rows


# ── Sheet names ─────────────────────────────────────────────────────────────


def sheet_name(code: object, used: set[str], n: int) -> str:
    """A valid, unique sheet name for a work: its code (≤ 31 characters, without []:*?/\\).

    A repeated name gets " (2)", " (3)"…: not "-2", because "3.1-2" is how Sol names the
    sheet of work 3.1.2 and Cargar obra would read it so. ``used`` holds lower-cased names.
    """
    base = _INVALID_SHEET_CHARS.sub("_", " ".join(str(code or "").split())).strip("'").strip()
    base = base[:31].strip() or f"Trabajo {n}"
    name, k = base, 2
    while name.lower() in used:
        suffix = f" ({k})"
        name = base[: 31 - len(suffix)].rstrip() + suffix
        k += 1
    used.add(name.lower())
    return name


# ── Writing ─────────────────────────────────────────────────────────────────


def _set(ws: Worksheet, row: int, col: int, value: object, fmt: str | None = None,
         font: Font | None = None, fill: PatternFill | None = None) -> None:
    cell = ws.cell(row, col, value)
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font
    if fill:
        cell.fill = fill


def _cp_header(ws: Worksheet, obra: str) -> None:
    ws["A1"] = obra
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "PLANILLA COMPUTO Y PRESUPUESTO"
    for col in range(1, LAST_COL + 1):
        ws.cell(3, col).fill = _BLACK
    ws["A3"].font = Font(bold=True, size=22, color="FFFFFFFF")

    top = {"A5": "ITEM", "B5": "DESCRIPCIÓN", "C5": "UNIDAD", "D5": "CANTIDAD",
           "E5": "SUBTOTAL 01: GASTOS DIRECTOS", "O5": "SUBTOTAL 02: GASTOS INDIRECTOS",
           "R5": "SUBTOTAL 03: BENEFICIO E IMPUESTOS", "U5": "TOTAL (NETO)",
           "E6": "V.UNITARIO", "L6": "V.GENERAL", "O6": "V. GENERAL", "R6": "V. GENERAL",
           "U6": "V. UNITARIO", "X6": "V.GENERAL"}
    sub = ["MAT.", "M.O.: JORNALES", "M.O.: EQUIPOS", "M.O.: MATERIALES", "M.O.: SUBCONTRATOS", "M.O.",
           "GENERAL", "MAT.", "M.O", "GENERAL", "MAT.", "M.O.", "GENERAL", "MAT.", "M.O.", "GENERAL",
           "MAT.", "M.O", "GENERAL", "MAT.", "M.O", "GENERAL"]
    for ref, text in top.items():
        ws[ref] = text
    for offset, text in enumerate(sub):
        ws.cell(7, C_MAT + offset, text)
    for rng in ("A5:A7", "B5:B7", "C5:C7", "D5:D7", "E5:N5", "O5:Q5", "R5:T5", "U5:Z5",
                "E6:K6", "L6:N6", "O6:Q6", "R6:T6", "U6:W6", "X6:Z6"):
        ws.merge_cells(rng)
    for r in (5, 6, 7):
        for col in range(1, LAST_COL + 1):
            cell = ws.cell(r, col)
            cell.border = _BOX
            cell.alignment = _CENTER
            cell.fill = _HEAD
            cell.font = _BOLD
    for r, h in ((1, 24), (2, 9), (3, 30.75), (5, 30), (7, 30)):
        ws.row_dimensions[r].height = h
    widths = {"A": 10, "B": 70, "C": 8, "D": 12, "E": 20.7}
    for col in range(1, LAST_COL + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = widths.get(letter, 16)
    ws.freeze_panes = "C8"


def _write_cp(ws: Worksheet, obra: str, rows: list[dict], resources: dict[str, list[dict]]) -> dict:
    """Fill 01_C&P; returns the totals row values {col: amount} and its row number."""
    _cp_header(ws, obra)
    sums = {col: 0 for col in (C_L, C_M, C_N, C_O, C_P, C_Q, C_R, C_S, C_T, C_X, C_Y, C_Z)}
    r = FIRST_ROW
    for row in rows:
        item = row["item"]
        if row["kind"] in ("rubro", "piso"):
            fill, font = (_RUBRO, _WHITE_B) if row["kind"] == "rubro" else (_PISO, _BOLD)
            ws.cell(r, C_DESC, f"{row['code']}- {_section_name(item, row['code'])}")
            for col in range(1, LAST_COL + 1):
                ws.cell(r, col).fill = fill
                ws.cell(r, col).font = font
            r += 1
            continue
        _set(ws, r, C_ITEM, str(item.get("code") or "").strip() or None, "@")
        ws.cell(r, C_DESC, " ".join(str(item.get("description") or "").split()) or "(sin descripción)")
        ws.cell(r, C_UNIDAD, item.get("unidad") or None)
        _set(ws, r, C_CANT, _f(item.get("cantidad")), FMT_CANT)
        values = work_values(item, resources.get(str(item.get("id")), []))
        for col, value in values.items():
            _set(ws, r, col, value, FMT_PESOS)
            if col in sums:
                sums[col] += _cents(value)
        r += 1

    total_row = r + 1
    _set(ws, total_row, C_ITEM, "TOTAL DEL PRESUPUESTO", font=_BOLD)
    totals = {col: cents / 100 for col, cents in sums.items()}
    for col in range(1, LAST_COL + 1):
        ws.cell(total_row, col).fill = _TOTAL
        ws.cell(total_row, col).border = Border(top=_THIN, bottom=_THIN)
    for col, value in totals.items():
        _set(ws, total_row, col, value, FMT_PESOS, font=_BOLD)
    return {"row": total_row, "values": totals}


def _write_work_sheet(ws: Worksheet, obra: str, item: dict, resources: list[dict]) -> None:
    """One work, like Sol's detail sheets: a block per kind of resource, then its direct cost."""
    cantidad = _f(item.get("cantidad"))
    unidad = item.get("unidad") or ""
    ws["A1"] = obra
    ws["A1"].font = _BOLD
    for col, text in ((1, "Código"), (2, "Descripción"), (8, "Cantidad"), (9, "Unidad")):
        _set(ws, 2, col, text, font=Font(italic=True, color="FF595959"))
    _set(ws, 3, 1, str(item.get("code") or ""), "@")
    ws.cell(3, 2, " ".join(str(item.get("description") or "").split()))
    _set(ws, 3, 8, cantidad, FMT_CANT)
    ws.cell(3, 9, unidad)
    ws.merge_cells("B3:G3")
    for col in range(1, 10):
        ws.cell(3, col).fill = _BLACK
        ws.cell(3, col).font = _WHITE_B
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 30

    r = 5
    for tipo, titulo in TIPO_TITULO:
        block = [res for res in resources if res.get("tipo") == tipo]
        if not block:
            continue
        mo = tipo == "mano_obra"
        _set(ws, r, 1, titulo, font=_WHITE_B, fill=_BLACK)
        for col in range(2, 10):
            ws.cell(r, col).fill = _BLACK
        r += 1
        heads = ["Código", "Descripción", "Unidad", "Cantidad", "Días" if mo else "",
                 "Cargas sociales" if mo else "Desperdicio",
                 "Jornales (con cargas)" if mo else "Cantidad + Desperdicio",
                 "Precio Unitario", "Subtotal", "Nota"]
        for col, text in enumerate(heads, start=1):
            _set(ws, r, col, text or None, font=_BOLD, fill=_HEAD)
            ws.cell(r, col).border = _BOX
            ws.cell(r, col).alignment = _CENTER
        r += 1
        first = r
        total = 0
        for res in block:
            cliente = bool(res.get("lo_compra_cliente"))
            if mo:
                qty, dias = _f(res.get("trabajadores")), _f(res.get("dias"))
                pct = _f(res.get("cargas_sociales_pct"))
                default_eff = qty * dias * (1 + pct / 100)
            else:
                qty, dias = _f(res.get("cantidad")), None
                pct = _f(res.get("desperdicio_pct"))
                default_eff = qty * (1 + pct / 100)
            efectiva = safe_float(res.get("cantidad_efectiva"))
            subtotal = 0.0 if cliente else round(_f(res.get("subtotal")), 2)
            total += _cents(subtotal)
            _set(ws, r, 1, str(res.get("codigo") or "") or None, "@")
            ws.cell(r, 2, res.get("descripcion") or None)
            ws.cell(r, 3, res.get("unidad") or None)
            _set(ws, r, 4, qty, FMT_CANT)
            if mo:
                _set(ws, r, 5, dias, FMT_CANT)
            _set(ws, r, 6, pct / 100, FMT_PCT)
            _set(ws, r, 7, round(efectiva if efectiva is not None else default_eff, 2), FMT_CANT)
            _set(ws, r, 8, round(_f(res.get("precio_unitario")), 2), FMT_PESOS)
            _set(ws, r, 9, subtotal, FMT_PESOS)
            if cliente:
                _set(ws, r, 10, "Lo compra el cliente", font=Font(italic=True, color="FF9C5700"),
                     fill=PatternFill("solid", fgColor="FFFFEB9C"))
            r += 1
        for rr in range(first, r):
            for col in range(1, 10):
                ws.cell(rr, col).border = _BOX
        _set(ws, r, 1, f"TOTAL {titulo}", font=_BOLD, fill=_TOTAL)
        for col in range(2, 10):
            ws.cell(r, col).fill = _TOTAL
        _set(ws, r, 9, total / 100, FMT_PESOS, font=_BOLD)
        r += 1
        _set(ws, r, 1, f"TOTAL x {unidad}".strip(), font=_BOLD, fill=_SOFT)
        for col in range(2, 10):
            ws.cell(r, col).fill = _SOFT
        _set(ws, r, 9, round(total / 100 / cantidad, 2) if cantidad else 0.0, FMT_PESOS, font=_BOLD)
        r += 3

    directo = work_values(item, resources)[C_N]
    _set(ws, r, 1, "TOTAL DEL TRABAJO (COSTO DIRECTO)", font=_WHITE_B, fill=_BLACK)
    for col in range(2, 10):
        ws.cell(r, col).fill = _BLACK
    _set(ws, r, 9, directo, FMT_PESOS, font=_WHITE_B)
    r += 1
    unit = round(_f(item.get("mat_unitario")) + _f(item.get("mo_unitario")), 2)
    _set(ws, r, 1, f"TOTAL x {unidad}".strip(), font=_BOLD, fill=_SOFT)
    for col in range(2, 10):
        ws.cell(r, col).fill = _SOFT
    _set(ws, r, 9, unit, FMT_PESOS, font=_BOLD)

    for letter, width in zip("ABCDEFGHIJ", (12, 40, 9, 11, 9, 13, 14, 18, 18, 20)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A4"


# Rows of the "Coeficiente de pase" sheet that carry a % (label in A, fraction in B), so a
# Planilla Terrac uploaded again in Cargar obra keeps the obra's percentages
COEF_TITLE = "COEFICIENTE DE PASE"
COEF_LABELS = {
    "Imprevistos": "imprevistos_pct",
    "Gastos de estructura": "estructura_pct",
    "Jefatura de obra": "jefatura_pct",
    "Logística": "logistica_pct",
    "Herramientas": "herramientas_pct",
    "Beneficio": "beneficio_pct",
    "Ingresos brutos": "ingresos_brutos_pct",
    "Impuesto al cheque": "imp_cheque_pct",
    "IVA": "iva_pct",
}


def read_coeficiente(wb) -> dict[str, float] | None:  # type: ignore[no-untyped-def]
    """The 9 % (as 15 = 15%) of a Planilla Terrac made by the app, or None.

    None unless the workbook has the app's "Coeficiente de pase" sheet with all 9 values
    between 0 and 100 %: an Excel made by hand never changes the obra's percentages.
    """
    if SHEET_COEF not in wb.sheetnames:
        return None
    ws = wb[SHEET_COEF]
    if str(ws["A3"].value or "").strip() != COEF_TITLE:
        return None
    found: dict[str, float] = {}
    for row in ws.iter_rows(min_row=6, max_row=40, max_col=2, values_only=True):
        label, value = str(row[0] or "").strip(), row[1]
        key = COEF_LABELS.get(label)
        if key is None or isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        if not 0 <= value <= 1:
            return None
        found[key] = round(float(value) * 100, 4)
    return found if len(found) == len(COEF_LABELS) else None


def _write_coef(ws: Worksheet, obra: str, cfg: dict, totals: dict) -> None:
    """The 9 % of the budget, the cascade of its saved totals and the coeficiente de pase."""
    ws["A1"] = obra
    ws["A1"].font = _BOLD
    ws["A3"] = COEF_TITLE
    for col in range(1, 5):
        ws.cell(3, col).fill = _BLACK
    ws["A3"].font = Font(bold=True, size=16, color="FFFFFFFF")
    ws.row_dimensions[3].height = 24
    for col, text in enumerate(("Concepto", "%", "Sobre", "En esta obra"), start=1):
        _set(ws, 5, col, text, font=_BOLD, fill=_HEAD)
        ws.cell(5, col).border = _BOX

    ind = totals["indirectos"]
    pct = {k: float(cfg[k]) / 100 for k in (
        "imprevistos_pct", "estructura_pct", "jefatura_pct", "logistica_pct", "herramientas_pct",
        "beneficio_pct", "ingresos_brutos_pct", "imp_cheque_pct", "iva_pct")}
    lines: list[tuple[str, float | None, str, float, bool]] = [
        ("SUBTOTAL 01 - COSTOS DIRECTOS", None, "", totals["directo_total"], True),
        ("Imprevistos", pct["imprevistos_pct"], "Costo directo", ind["imprevistos_pct"], False),
        ("Gastos de estructura", pct["estructura_pct"], "Costo directo", ind["estructura_pct"], False),
        ("Jefatura de obra", pct["jefatura_pct"], "Costo directo", ind["jefatura_pct"], False),
        ("Logística", pct["logistica_pct"], "Costo directo", ind["logistica_pct"], False),
        ("Herramientas", pct["herramientas_pct"], "Costo directo", ind["herramientas_pct"], False),
        ("SUBTOTAL 02 - CON GASTOS INDIRECTOS", None, "", totals["subtotal_02"], True),
        ("Beneficio", pct["beneficio_pct"], "Subtotal 02", totals["beneficio_total"], False),
        ("SUBTOTAL 03 - CON BENEFICIO", None, "", totals["subtotal_03"], True),
        ("Ingresos brutos", pct["ingresos_brutos_pct"], "Subtotal 03", totals["iibb"], False),
        ("Impuesto al cheque", pct["imp_cheque_pct"], "Subtotal 03", totals["cheque"], False),
        ("PRECIO SIN IVA", None, "", totals["neto_total"], True),
        ("IVA", pct["iva_pct"], "Precio sin IVA", totals["iva_total"], False),
        ("PRECIO CON IVA", None, "", totals["total_final"], True),
    ]
    r = 6
    for label, p, sobre, amount, strong in lines:
        font = _BOLD if strong else None
        fill = _TOTAL if strong else None
        _set(ws, r, 1, label, font=font, fill=fill)
        _set(ws, r, 2, p, FMT_PCT if p is not None else None, fill=fill)
        _set(ws, r, 3, sobre or None, fill=fill)
        _set(ws, r, 4, amount, FMT_PESOS, font=font, fill=fill)
        for col in range(1, 5):
            ws.cell(r, col).border = _BOX
        r += 1

    coef = cascade_factors(cfg)["coeficiente"]
    iva = 1 + pct["iva_pct"]
    r += 1
    _set(ws, r, 1, "Coeficiente de pase (sin IVA)", font=_BOLD)
    _set(ws, r, 4, coef, "0.0000", font=_BOLD)
    r += 1
    sin_iva = round(100 * coef, 2)
    con_iva = round(100 * coef * iva, 2)
    ws.cell(r, 1, f"Por cada $100 de costo directo, ${_ars(sin_iva)} sin IVA (${_ars(con_iva)} con IVA).")
    r += 2
    ws.cell(r, 1, "Los montos son los guardados de cada trabajo: los mismos del editor.").font = Font(
        italic=True, color="FF595959")
    for letter, width in zip("ABCD", (40, 10, 16, 22)):
        ws.column_dimensions[letter].width = width


def _ars(value: float) -> str:
    """1234.5 → '1.234,50'."""
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def terrac_workbook(budget: dict, items: list[dict], resources: dict[str, list[dict]],
                    cfg: dict, totals: dict) -> Workbook:
    """The Planilla Terrac of a budget.

    ``items``: every budget_item, in sort order. ``resources``: item id → its
    item_resources. ``cfg``: the budget's cascade config with defaults.
    ``totals``: the router's pdf_totals (the saved sums, split by the %).
    """
    obra = " ".join(str(budget.get("name") or "Presupuesto").split())
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CP
    rows = cp_rows(items)
    _write_cp(ws, obra, rows, resources)

    used = {SHEET_CP.lower(), SHEET_COEF.lower()}
    n = 0
    for row in rows:
        item = row["item"]
        if row["kind"] != "work" or is_section(item):
            continue
        n += 1
        own = resources.get(str(item.get("id"))) or []
        if not own:
            continue
        name = sheet_name(item.get("code"), used, n)
        _write_work_sheet(wb.create_sheet(name), obra, item, own)

    _write_coef(wb.create_sheet(SHEET_COEF), obra, cfg, totals)
    return wb


def terrac_bytes(budget: dict, items: list[dict], resources: dict[str, list[dict]],
                 cfg: dict, totals: dict) -> BytesIO:
    """terrac_workbook saved as .xlsx, ready to stream."""
    output = BytesIO()
    terrac_workbook(budget, items, resources, cfg, totals).save(output)
    output.seek(0)
    return output
