"""Arma ginkgo_sin_precios.xlsx: el Excel de Ginkgo solo con cantidades (para scripts/e2e_sin_precios.cjs).

Uso (desde la raíz del repo):
  EXCEL_DIR=<carpeta con ginkgo.xlsx> python3 scripts/make_ginkgo_sin_precios.py
  python3 scripts/make_ginkgo_sin_precios.py <carpeta con ginkgo.xlsx>
Deja solo la hoja 01_C&P y vacía sus columnas E, J, N y Z (costos) en las filas de datos. Lo demás
queda como valores: las cantidades de Ginkgo son fórmulas que apuntan a otras hojas, y esas hojas se borran.
Guarda ginkgo_sin_precios.xlsx en la misma carpeta.
"""
from __future__ import annotations

import os
import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from app.obra_import import COL_DIR_TOTAL, COL_MAT_UNIT, COL_MO_UNIT, COL_NETO_TOTAL, FIRST_ROW, SHEET  # noqa: E402

COLUMNAS = (COL_MAT_UNIT, COL_MO_UNIT, COL_DIR_TOTAL, COL_NETO_TOTAL)  # E, J, N, Z


def main() -> None:
    carpeta = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("EXCEL_DIR")
    if not carpeta:
        sys.exit("Decime la carpeta de ginkgo.xlsx: EXCEL_DIR=<carpeta> o como argumento.")
    origen = Path(carpeta) / "ginkgo.xlsx"
    if not origen.exists():
        sys.exit(f"No encuentro {origen}")
    destino = origen.with_name("ginkgo_sin_precios.xlsx")

    # Values, not formulas: the quantities point to sheets that are about to be deleted
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(origen, data_only=True)
    for nombre in list(wb.sheetnames):
        if nombre != SHEET:
            del wb[nombre]
    ws = wb[SHEET]

    vaciadas = 0
    for r in range(FIRST_ROW, ws.max_row + 1):
        celdas = [ws.cell(r, c) for c in COLUMNAS]
        if any(c.value not in (None, "") for c in celdas):
            vaciadas += 1
        for c in celdas:
            c.value = None

    wb.save(destino)
    print(f"Listo: {destino} (hoja {SHEET}; {vaciadas} filas con costos vaciadas en E, J, N y Z).")


if __name__ == "__main__":
    main()
