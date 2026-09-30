"""
Importa las 4 listas de precios del Excel "Maestro" de TERRAC
(solapas 00_Mat, 00_MO, 00_Eq y 00_Sub) como catálogos con fecha y proveedor.

Uso:
  # 1. Solo el informe (no toca la base):
  python3 import_maestro.py "MODELO DE PRESUPUESTACION RESUMEN.xlsx"

  # 2. Importar en Supabase (crea 4 catálogos nuevos):
  python3 import_maestro.py "MODELO DE PRESUPUESTACION RESUMEN.xlsx" --apply

Para --apply hacen falta las variables de entorno:
  DATA_SUPABASE_URL  o  SUPABASE_URL
  DATA_SUPABASE_KEY  o  SUPABASE_KEY
  ORG_ID

Antes de --apply hay que correr migrations/004_catalog_price_dates.sql.
Siempre se escribe el informe de códigos duplicados y precios sin fecha
(por defecto en output/informe_maestro_precios.md).
"""

from __future__ import annotations

import argparse
import os
import sys
import warnings
from pathlib import Path

import openpyxl

from app.catalog_prices import history_row
from app.maestro_import import TIPO_LABELS, parse_workbook, report_markdown

CHUNK = 200


def load_workbook(path: str):  # type: ignore[no-untyped-def]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=True)


def apply_to_supabase(entries_by_tipo: dict[str, list[dict]], source_file: str, prefix: str) -> None:
    from supabase import create_client

    url = os.environ.get("DATA_SUPABASE_URL") or os.environ.get("SUPABASE_URL")
    key = os.environ.get("DATA_SUPABASE_KEY") or os.environ.get("SUPABASE_KEY")
    org_id = os.environ.get("ORG_ID")
    if not url or not key or not org_id:
        sys.exit("ERROR: faltan DATA_SUPABASE_URL, DATA_SUPABASE_KEY u ORG_ID")

    db = create_client(url, key)
    for tipo, entries in entries_by_tipo.items():
        if not entries:
            continue
        name = f"{prefix} - {TIPO_LABELS[tipo]}"
        catalog = db.table("price_catalogs").insert({
            "org_id": org_id, "name": name, "source_file": source_file,
        }).execute()
        catalog_id = catalog.data[0]["id"]

        rows = [{**e, "catalog_id": catalog_id, "org_id": org_id} for e in entries]
        inserted: list[dict] = []
        for i in range(0, len(rows), CHUNK):
            inserted += db.table("catalog_entries").insert(rows[i:i + CHUNK]).execute().data or []

        history = [history_row(e) for e in inserted if e.get("precio_sin_iva") is not None]
        for i in range(0, len(history), CHUNK):
            db.table("catalog_price_history").insert(history[i:i + CHUNK]).execute()

        print(f"  {name}: {len(inserted)} entradas, {len(history)} precios en el historial")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Importa el Maestro de precios de TERRAC")
    parser.add_argument("excel", help="Ruta al .xlsx del Maestro")
    parser.add_argument("--apply", action="store_true", help="Escribir en Supabase (si no, solo el informe)")
    parser.add_argument("--report", default="output/informe_maestro_precios.md", help="Ruta del informe")
    parser.add_argument("--prefix", default="Maestro TERRAC", help="Prefijo del nombre de los catálogos")
    args = parser.parse_args(argv)

    source = Path(args.excel).name
    entries_by_tipo, report = parse_workbook(load_workbook(args.excel))

    Path(args.report).parent.mkdir(parents=True, exist_ok=True)
    Path(args.report).write_text(report_markdown(report, source), encoding="utf-8")

    for sheet, info in report["hojas"].items():
        print(
            f"{sheet}: {info['entradas']} entradas, {len(info.get('sin_fecha', []))} sin fecha, "
            f"{len(info.get('duplicados', []))} códigos duplicados"
        )
    if report["faltantes"]:
        print(f"Hojas que no se encontraron: {', '.join(report['faltantes'])}")
    print(f"Informe: {args.report}")

    if args.apply:
        print("Importando en Supabase...")
        apply_to_supabase(entries_by_tipo, source, args.prefix)
    else:
        print("No se escribió nada en la base (usar --apply para importar).")


if __name__ == "__main__":
    main()
