"""Cargar una obra desde la app: /obras/analizar y /obras/cargar."""

from __future__ import annotations

import copy
import json
import os
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.budget_prices import today
from app.main import create_app
from tests.test_obra_import import TEMPLATES, _workbook
from tests.test_recipes_api import MOCK_USER, ORG, FakeDB

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel(wb=None) -> bytes:
    buf = BytesIO()
    (wb or _workbook()).save(buf)
    return buf.getvalue()


def _tables(**over) -> dict:
    templates = [
        {**copy.deepcopy(t), "id": f"tmpl-{c}", "org_id": ORG, "nombre": f"Receta {c}", "desperdicio_pct": None}
        for c, t in TEMPLATES.items()
    ]
    entries = [
        ("LH18", "material", 500), ("HADN6", "material", 100), ("H30", "material", 1000),
        ("MO-OF", "mano_obra", 200), ("EPS-500", "material", 0), ("ES", "material", 30),
        ("RP-KP", "material", 40), ("RP-PORC", "material", 0),
    ]
    tables = {
        "budgets": [],
        "budget_items": [],
        "item_resources": [],
        "item_templates": templates,
        "indirect_config": [{"id": "cfg", "org_id": ORG, "desperdicio_pct": 5}],
        "price_catalogs": [{"id": "cat1", "org_id": ORG, "name": "Maestro TERRAC - Materiales"}],
        "catalog_entries": [
            # EPS-500 is an old Maestro zero: undated, so "sin precio" (a dated 0 would be a price)
            {"id": f"e-{c}", "catalog_id": "cat1", "org_id": ORG, "codigo": c, "tipo": t,
             "precio_sin_iva": p, "fecha_precio": None if c == "EPS-500" else "2026-01-01"} for c, t, p in entries
        ],
        "catalog_price_history": [],
    }
    tables.update(over)
    return tables


@pytest.fixture
def client():
    app = create_app()
    from app.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: MOCK_USER
    return TestClient(app)


@pytest.fixture
def db():
    fake = FakeDB(_tables())
    with patch("app.routers.obras.get_data_db", return_value=fake), \
         patch("app.routers.analysis.get_data_db", return_value=fake):
        yield fake


def analizar(client, asignaciones=None, wb=None):
    data = {"asignaciones": json.dumps(asignaciones)} if asignaciones else {}
    return client.post("/obras/analizar", files={"file": ("obra.xlsx", _excel(wb), XLSX)}, data=data)


def cargar(client, nombre="OBRA PRUEBA", permitir=False, asignaciones=None, wb=None):
    data = {"nombre": nombre, "permitir_sin_precio": str(permitir).lower()}
    if asignaciones:
        data["asignaciones"] = json.dumps(asignaciones)
    return client.post("/obras/cargar", files={"file": ("obra.xlsx", _excel(wb), XLSX)}, data=data)


def _tareas(body) -> dict:
    return {t["clave"]: t for t in body["tareas"]}


ARISTAS = "ARISTAS DE YESO EN PAREDES | m"
TELGOPOR = "TELGOPOR 50 MM + CONTRAPISO EN AZOTEA ACCESIBLE E: 8CM | m2"
TENSORES = "TENSORES 20 CM X 40 CM. | m"
OBRADOR = "OBRADOR | gl"
MURO_CARGA = "MURO DE CARGA EN LADRILLO HUECO DEL 18 | m2"


HALL = "CONTRAPISO EN HALL + RAMPAS. ESP.=10CM. | m2"
CONTRAPISO = "CONTRAPISO | m2"
BALCONES = "TELGOPOR 50 MM + CONTRAPISO/CARPETA EN AZOTEA INACCESIBLE (BALCONES) E=4CM | m2"


def _workbook_contrapisos():
    """The obra plus three contrapisos (Ginkgo's names): thickness in the name, or not."""
    wb = _workbook()
    ws = wb["01_C&P"]
    for r, (code, desc) in enumerate([
        ("4.3-6", "CONTRAPISO EN HALL + RAMPAS. ESP.=10cm."),
        ("4.3-7", "CONTRAPISO"),
        ("4.3-8", "TELGOPOR 50 mm + CONTRAPISO/CARPETA EN AZOTEA INACCESIBLE (BALCONES) e=4cm"),
    ], start=ws.max_row + 1):
        for col, value in ((1, code), (2, desc), (3, "m²"), (4, 10), (5, 1), (10, 1), (14, 20), (26, 30)):
            ws.cell(r, col, value)
    return wb


def _fix_eps(db):
    next(e for e in db.tables["catalog_entries"] if e["codigo"] == "EPS-500")["precio_sin_iva"] = 10


class TestAnalizar:
    def test_contract(self, client, db):
        res = analizar(client)
        assert res.status_code == 200, res.text
        body = res.json()
        assert set(body) == {"archivo", "titulo_dudoso", "catalogo_oficial", "excel_con_precios", "titulo",
                             "fecha_precios", "resumen", "tareas", "precios", "recetas", "correcciones_excel",
                             "listo"}
        assert body["catalogo_oficial"] is False
        assert body["excel_con_precios"] is True
        assert body["titulo_dudoso"] is False  # "obra.xlsx": nothing to compare the title with
        assert body["resumen"] == {
            "rubros": 3, "pisos": 3, "trabajos": 7, "grupos": 6,
            "verdes": 3, "amarillos": 2, "rojos": 1, "total_excel": 13850.0,
        }
        assert [t["estado"] for t in body["tareas"]] == ["rojo", "amarillo", "amarillo", "verde", "verde", "verde"]
        for t in body["tareas"]:
            assert set(t) == {"clave", "descripcion", "unidad", "veces", "cantidad_total", "total_excel", "codigos",
                              "estado", "motivo_rojo", "receta", "sugerencias", "pregunta", "avisos",
                              "precios_faltantes"}
        tareas = _tareas(body)

        telgopor = tareas[TELGOPOR]
        assert (telgopor["estado"], telgopor["motivo_rojo"]) == ("rojo", "precio")
        assert telgopor["precios_faltantes"] == ["EPS-500"]
        assert telgopor["receta"]["origen"] == "regla"
        assert telgopor["receta"]["porque"] == "Coincide la descripción"
        assert telgopor["receta"]["nombre"] == "Receta 8.3 + Receta 5.2.3"
        assert [p["factor"] for p in telgopor["receta"]["partes"]] == [1.0, 0.08]
        assert telgopor["pregunta"]["receta"] == "5.2.3"
        assert telgopor["pregunta"]["valor"] == 0.08
        assert set(telgopor["pregunta"]) == {"tipo", "texto", "receta", "unidad_receta", "unidad_obra", "valor",
                                             "dato", "origen_valor"}
        # The thickness is in the name ("e: 8cm"): shown as a fact, saying where it came from
        assert telgopor["pregunta"]["dato"] == 'Cada m² lleva 0,08 m³ de "Receta 5.2.3" (por los 8 cm del nombre)'
        assert telgopor["pregunta"]["origen_valor"] == "nombre"
        assert telgopor["avisos"] == ["Compuesto: placas EPS + contrapiso de cascote."]  # sin códigos

        tensores = tareas[TENSORES]
        assert tensores["estado"] == "verde"
        assert tensores["pregunta"]["texto"] == '¿Cuántos m³ de "Receta 4.1.7" lleva cada ml de este trabajo?'
        assert (tensores["pregunta"]["unidad_receta"], tensores["pregunta"]["unidad_obra"]) == ("m3", "ml")
        assert tensores["pregunta"]["dato"] == 'Cada ml lleva 0,08 m³ de "Receta 4.1.7"'
        assert tensores["pregunta"]["origen_valor"] == "regla"

        assert len(body["precios"]) == 1
        for p in body["precios"]:
            assert set(p) == {"codigo", "descripcion", "unidad", "tipo", "problema", "motivo", "recursos", "items",
                              "entradas", "referencias", "propuesta", "catalogo_destino"}
            # Without an oficial catalog: no references, no destination; no Excel price to propose
            assert (p["referencias"], p["catalogo_destino"], p["propuesta"]) == ([], None, None)

        piso = next(t for t in body["tareas"] if t["descripcion"].startswith("COLOCACION DE REVESTIMIENTOS"))
        assert (piso["veces"], piso["cantidad_total"], piso["codigos"]) == (2, 40.0, ["4.2.4", "4.3.5"])
        assert piso["pregunta"] is None
        assert piso["avisos"] == ["No incluye el porcelanato: queda como material que compra el cliente."]

        obrador = tareas[OBRADOR]
        assert (obrador["estado"], obrador["receta"], obrador["motivo_rojo"]) == ("amarillo", None, None)

        assert {r["codigo"] for r in body["recetas"]} == set(TEMPLATES)
        assert body["recetas"][0] == {"codigo": "4.1.7", "nombre": "Receta 4.1.7", "unidad": "m3", "categoria": None}
        assert body["listo"] is False

    def test_summary_and_price_problems(self, client, db):
        body = analizar(client).json()
        precios = {p["codigo"]: p for p in body["precios"]}
        assert precios["EPS-500"]["problema"] == "sin_precio"
        assert precios["EPS-500"]["entradas"][0]["id"] == "e-EPS-500"
        assert "RP-PORC" not in precios  # lo compra el cliente
        assert db.tables["budgets"] == []  # analizar no escribe

        _fix_eps(db)
        body = analizar(client).json()
        assert body["resumen"]["rojos"] == 0
        assert body["listo"] is True

    def test_question_needs_the_value(self, client, db):
        body = analizar(client, {ARISTAS: {"plantillas": [["5.1.4", None]]}}).json()
        aristas = _tareas(body)[ARISTAS]
        assert (aristas["estado"], aristas["motivo_rojo"]) == ("rojo", "pregunta")
        assert aristas["pregunta"]["valor"] is None
        assert aristas["pregunta"]["dato"] is None
        assert aristas["pregunta"]["origen_valor"] is None
        assert aristas["pregunta"]["texto"] == '¿Cuántos m² de "Receta 5.1.4" lleva cada m de este trabajo?'
        assert aristas["receta"]["origen"] == "manual"

        body = analizar(client, {ARISTAS: {"plantillas": [["5.1.4", 2.5]]}}).json()
        aristas = _tareas(body)[ARISTAS]
        assert aristas["estado"] == "verde"
        assert aristas["pregunta"]["valor"] == 2.5
        assert aristas["pregunta"]["dato"] == 'Cada m lleva 2,5 m² de "Receta 5.1.4"'
        assert aristas["pregunta"]["origen_valor"] == "mano"
        assert aristas["receta"]["partes"][0]["factor"] == 2.5
        assert aristas["receta"]["porque"] == "La elegiste vos"

    def test_confirmations(self, client, db):
        body = analizar(client, {
            OBRADOR: {"plantillas": [], "confirmada": True},
            # Confirming the rule keeps its conversion when the screen sends none
            TENSORES: {"plantillas": [["4.1.7", None]], "confirmada": True},
        }).json()
        tareas = _tareas(body)
        assert (tareas[OBRADOR]["estado"], tareas[OBRADOR]["receta"]) == ("verde", None)
        assert tareas[TENSORES]["estado"] == "verde"
        assert tareas[TENSORES]["receta"]["origen"] == "regla"
        assert tareas[TENSORES]["pregunta"]["valor"] == 0.08

    def test_suggestion_until_confirmed(self, client, db):
        """A suggestion by similar words is never the recipe, however good: the card stays
        yellow without recipe (the Excel price) and offers it in "sugerencias"."""
        db.tables["item_templates"].append({"id": "tmpl-6.9", "org_id": ORG, "codigo": "6.9", "unidad": "m",
                                            "nombre": "ARISTAS DE YESO", "parametros": [], "recursos": []})
        aristas = _tareas(analizar(client).json())[ARISTAS]
        assert (aristas["estado"], aristas["receta"], aristas["motivo_rojo"]) == ("amarillo", None, None)
        assert aristas["sugerencias"][0] == {"codigo": "6.9", "nombre": "ARISTAS DE YESO", "unidad": "m",
                                             "porque": "Se parece por 'aristas', 'yeso'"}

        # Sol picks it: a recipe chosen by hand
        body = analizar(client, {ARISTAS: {"plantillas": [["6.9", 1]], "confirmada": True}}).json()
        aristas = _tareas(body)[ARISTAS]
        assert (aristas["estado"], aristas["receta"]["origen"]) == ("verde", "manual")
        assert aristas["receta"]["porque"] == "La elegiste vos"
        assert aristas["sugerencias"] == []

    def test_weak_suggestion_does_not_bring_its_prices(self, client, db):
        """Ginkgo: "membrana líquida" looked like "pintura en paredes" and the card went red
        with the paint's missing prices. Now it stays yellow, with the Excel price."""
        db.tables["item_templates"].append({"id": "tmpl-7.4.2", "org_id": ORG, "codigo": "7.4.2", "unidad": "m",
                                            "nombre": "ARISTAS PINTADAS", "parametros": [],
                                            "recursos": [{"tipo": "material", "codigo": "NO-ESTA", "formula": "Q"}]})
        body = analizar(client).json()
        aristas = _tareas(body)[ARISTAS]
        assert (aristas["estado"], aristas["receta"], aristas["precios_faltantes"]) == ("amarillo", None, [])
        assert aristas["sugerencias"][0]["codigo"] == "7.4.2"
        assert all(p["codigo"] != "NO-ESTA" for p in body["precios"])

    def test_memory(self, client, db):
        db.tables["obra_recetas_memoria"] = [
            {"id": "m1", "org_id": ORG, "clave": ARISTAS, "plantillas": [["5.1.4", 2.0]], "veces": 1},
            {"id": "m2", "org_id": "otra-org", "clave": OBRADOR, "plantillas": [["5.1.4", 1.0]], "veces": 1},
        ]
        tareas = _tareas(analizar(client).json())
        assert tareas[ARISTAS]["estado"] == "verde"
        assert tareas[ARISTAS]["receta"]["origen"] == "memoria"
        assert tareas[ARISTAS]["receta"]["porque"] == "Ya la usaste así en otra obra"
        assert tareas[ARISTAS]["pregunta"]["valor"] == 2.0
        assert tareas[OBRADOR]["receta"] is None  # memory of another org

        # The screen's choice wins over memory
        tareas = _tareas(analizar(client, {ARISTAS: {"plantillas": []}}).json())
        assert tareas[ARISTAS]["receta"] is None

    def test_recipe_not_in_the_app(self, client, db):
        body = analizar(client, {ARISTAS: {"plantillas": [["9.9", 1], ["5.1.4", None]]}}).json()
        aristas = _tareas(body)[ARISTAS]
        assert (aristas["estado"], aristas["motivo_rojo"]) == ("rojo", "receta_inexistente")
        assert aristas["receta"]["partes"][0]["nombre"] is None
        assert body["listo"] is False

    def test_bad_choices(self, client, db):
        assert analizar(client, {ARISTAS: {"plantillas": "5.1.4"}}).status_code == 400
        res = client.post("/obras/analizar", files={"file": ("obra.xlsx", _excel(), XLSX)},
                          data={"asignaciones": "{no"})
        assert res.status_code == 400

    def test_duplicated_code(self, client, db):
        # LH18 exists dated (2026-01-01) in the fixture; a second one, undated, must not block
        db.tables["catalog_entries"].append(
            {"id": "e-dup", "catalog_id": "cat1", "org_id": ORG, "codigo": "lh18", "tipo": "material",
             "precio_sin_iva": 9, "fecha_precio": None})
        precios = {p["codigo"]: p for p in analizar(client).json()["precios"]}
        assert "LH18" not in precios
        # Two undated copies of the same code: nobody can tell which one is right
        for e in db.tables["catalog_entries"]:
            if e["codigo"] in ("LH18", "lh18"):
                e["fecha_precio"] = None
        precios = {p["codigo"]: p for p in analizar(client).json()["precios"]}
        assert precios["LH18"]["problema"] == "duplicado"
        assert len(precios["LH18"]["entradas"]) == 2

    def test_future_price_in_another_catalog_does_not_hide_the_one_in_force(self, client, db):
        """Codex: LH18 at 500 (2026-01-01, in force) and a copy at 900 dated in the future.
        The analysis must price LH18 at 500 today, not pick the future entry and go red."""
        db.tables["price_catalogs"].append({"id": "cat2", "org_id": ORG, "name": "Lista futura"})
        db.tables["catalog_entries"].append(
            {"id": "e-fut", "catalog_id": "cat2", "org_id": ORG, "codigo": "LH18", "tipo": "material",
             "precio_sin_iva": 900, "fecha_precio": "2099-01-01"})
        _fix_eps(db)
        body = analizar(client).json()
        assert all(p["codigo"] != "LH18" for p in body["precios"])
        assert cargar(client).status_code == 200
        lh18 = next(r for r in db.tables["item_resources"] if r["codigo"] == "LH18")
        assert lh18["precio_unitario"] == 500
        assert lh18["catalog_entry_id"] == "e-LH18"

    def test_thickness_from_the_name(self, client, db):
        """The m³ of contrapiso per m² comes from the thickness in the name; without one, the
        rule's 10 cm, saying it was assumed."""
        _fix_eps(db)
        tareas = _tareas(analizar(client, wb=_workbook_contrapisos()).json())
        hall, sin, balcon = tareas[HALL], tareas[CONTRAPISO], tareas[BALCONES]
        assert [p["factor"] for p in hall["receta"]["partes"]] == [0.1]
        assert hall["pregunta"]["dato"] == 'Cada m² lleva 0,1 m³ de "Receta 5.2.3" (por los 10 cm del nombre)'
        assert hall["pregunta"]["origen_valor"] == "nombre"
        assert sin["pregunta"]["dato"] == 'Cada m² lleva 0,1 m³ de "Receta 5.2.3" (supuse 10 cm)'
        assert sin["pregunta"]["origen_valor"] == "regla"
        # "TELGOPOR 50 mm": the 50 mm is the board, the contrapiso is the e=4cm
        assert [p["factor"] for p in balcon["receta"]["partes"]] == [1.0, 0.04]
        assert balcon["pregunta"]["dato"] == 'Cada m² lleva 0,04 m³ de "Receta 5.2.3" (por los 4 cm del nombre)'
        assert all(t["estado"] == "verde" for t in (hall, sin, balcon))

    def test_thickness_origin_after_confirming_or_changing(self, client, db):
        wb = _workbook_contrapisos()
        tareas = _tareas(analizar(client, {
            # Confirming sends back the value shown: it is still the rule's assumption
            CONTRAPISO: {"plantillas": [["5.2.3", 0.1]], "confirmada": True},
            # Sol writes another thickness: hers
            HALL: {"plantillas": [["5.2.3", 0.12]], "confirmada": True},
            BALCONES: {"plantillas": [["8.3", None], ["5.2.3", None]], "confirmada": True},
        }, wb=wb).json())
        assert tareas[CONTRAPISO]["pregunta"]["dato"].endswith("(supuse 10 cm)")
        assert tareas[CONTRAPISO]["receta"]["origen"] == "regla"
        assert (tareas[HALL]["pregunta"]["dato"], tareas[HALL]["pregunta"]["origen_valor"]) == (
            'Cada m² lleva 0,12 m³ de "Receta 5.2.3"', "mano")
        assert tareas[BALCONES]["pregunta"]["origen_valor"] == "nombre"
        assert tareas[BALCONES]["pregunta"]["valor"] == 0.04

        db.tables["obra_recetas_memoria"] = [
            {"id": "m1", "org_id": ORG, "clave": CONTRAPISO, "plantillas": [["5.2.3", 0.07]], "veces": 1}]
        sin = _tareas(analizar(client, wb=wb).json())[CONTRAPISO]
        assert (sin["pregunta"]["dato"], sin["pregunta"]["origen_valor"]) == (
            'Cada m² lleva 0,07 m³ de "Receta 5.2.3"', "memoria")

    def test_title_that_does_not_match_the_file(self, client, db):
        wb = _workbook()
        wb["01_C&P"]["A1"] = "EDIFICIO LAS HERAS"
        res = client.post("/obras/analizar",
                          files={"file": ("EDIFICIO GINKGO_Computo y Presupuesto_V2.xlsx", _excel(wb), XLSX)})
        assert res.json()["titulo_dudoso"] is True
        res = client.post("/obras/analizar",
                          files={"file": ("Edificio Las Heras - cómputo.xlsx", _excel(wb), XLSX)})
        assert res.json()["titulo_dudoso"] is False

    def test_the_apps_own_simple_export_says_what_it_is(self, client, db):
        """Carlos, 06/10: the old "Excel Formato Terrac" downloaded the simple sheet; uploading it
        must say what it is and what to upload instead, not just "no tiene la hoja"."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Presupuesto"
        ws.append(["Codigo", "Descripcion", "Unidad", "Cantidad", "Directo Total", "Neto Total"])
        res = client.post("/obras/analizar", files={"file": ("Ginkgo_terrac.xlsx", _excel(wb), XLSX)})
        assert res.status_code == 400
        assert "planilla simple que baja la app" in res.json()["detail"]
        assert "Planilla Terrac" in res.json()["detail"]

    def test_missing_sheet_lists_the_sheets_it_has(self, client, db):
        wb = openpyxl.Workbook()
        wb.active.title = "Hoja1"
        wb.create_sheet("Cómputo")
        res = client.post("/obras/analizar", files={"file": ("otra.xlsx", _excel(wb), XLSX)})
        assert res.status_code == 400
        detail = res.json()["detail"]
        assert "no tiene la hoja 01_C&P" in detail and "Hoja1, Cómputo" in detail

    def test_rejects_other_files(self, client, db):
        res = client.post("/obras/analizar", files={"file": ("obra.csv", b"a,b", "text/csv")})
        assert res.status_code == 400


class TestVaEnCero:
    def test_dated_zero_is_a_price(self, client, db):
        """Sol: "Va en $0" on EPS-500 saves 0 dated today; the card stops being red and the
        resource loads at 0, linked to its catalog entry."""
        cat, eid = "00000000-0000-0000-0000-00000000ca71", "00000000-0000-0000-0000-0000000000e5"
        db.tables["price_catalogs"][0]["id"] = cat
        for e in db.tables["catalog_entries"]:
            e["catalog_id"] = cat
            if e["codigo"] == "EPS-500":
                e["id"] = eid
        body = analizar(client).json()
        assert "EPS-500" in [p["codigo"] for p in body["precios"]]  # an undated 0: no price
        assert _tareas(body)[TELGOPOR]["estado"] == "rojo"

        with patch("app.routers.catalogs.get_data_db", return_value=db):
            res = client.patch(f"/catalogs/{cat}/entries/{eid}", json={"precio_sin_iva": 0})
        assert res.status_code == 200, res.text
        eps = next(e for e in db.tables["catalog_entries"] if e["id"] == eid)
        assert (eps["precio_sin_iva"], eps["fecha_precio"]) == (0, today().isoformat())

        body = analizar(client).json()
        assert body["precios"] == []
        telgopor = _tareas(body)[TELGOPOR]
        assert (telgopor["estado"], telgopor["precios_faltantes"]) == ("verde", [])
        assert body["listo"] is True

        assert cargar(client).status_code == 200
        r = next(r for r in db.tables["item_resources"] if r["codigo"] == "EPS-500")
        assert (r["precio_unitario"], r["catalog_entry_id"], r["precio_fecha"]) == (0, eid, today().isoformat())


class TestCargar:
    def test_blocked_while_prices_missing(self, client, db):
        res = cargar(client)
        assert res.status_code == 409
        detail = res.json()["detail"]
        assert detail["rojos"] == [TELGOPOR]
        assert "EPS-500" in json.dumps(detail["precios"])
        assert detail["mensaje"] == "Falta 1 precio: corregilo antes de cargar, o cargá igual con ese precio en $0"
        assert db.tables["budgets"] == []

    def test_blocked_without_conversion(self, client, db):
        _fix_eps(db)
        res = cargar(client, permitir=True, asignaciones={ARISTAS: {"plantillas": [["5.1.4", None]]}})
        assert res.status_code == 409
        assert res.json()["detail"]["rojos"] == [ARISTAS]
        assert res.json()["detail"]["mensaje"] == "Hay 1 trabajo en rojo: resolvelo antes de cargar"
        assert db.tables["budgets"] == []

    def test_blocked_with_missing_recipe(self, client, db):
        res = cargar(client, permitir=True, asignaciones={ARISTAS: {"plantillas": [["9.9", 1]]}})
        assert res.status_code == 409
        assert res.json()["detail"]["rojos"] == [ARISTAS]

    def test_saves_memory(self, client, db):
        _fix_eps(db)
        elegidas = {ARISTAS: {"plantillas": [["5.1.4", 2.5]], "confirmada": True},
                    OBRADOR: {"plantillas": [], "confirmada": True}}
        res = cargar(client, asignaciones=elegidas)
        assert res.status_code == 200, res.text
        assert res.json()["memoria_guardada"] == 2
        memoria = {m["clave"]: m for m in db.tables["obra_recetas_memoria"]}
        assert memoria[ARISTAS]["plantillas"] == [["5.1.4", 2.5]]
        assert memoria[ARISTAS]["org_id"] == ORG
        assert (memoria[OBRADOR]["plantillas"], memoria[OBRADOR]["veces"]) == ([], 1)
        aristas = next(i for i in db.tables["budget_items"] if i["code"] == "4.2.2")
        assert aristas["template_id"] == "tmpl-5.1.4"

        # Same task again: one row per (org, clave), updated
        elegidas[ARISTAS]["plantillas"] = [["5.1.4", 3]]
        assert cargar(client, nombre="OTRA", asignaciones=elegidas).json()["memoria_guardada"] == 2
        memoria = [m for m in db.tables["obra_recetas_memoria"] if m["clave"] == ARISTAS]
        assert len(memoria) == 1
        assert (memoria[0]["plantillas"], memoria[0]["veces"]) == ([["5.1.4", 3.0]], 2)

        # Next obra: proposed from memory
        aristas = _tareas(analizar(client).json())[ARISTAS]
        assert (aristas["receta"]["origen"], aristas["estado"]) == ("memoria", "verde")

    def test_combined_recipes_keep_their_waste(self, client, db):
        """Codex: in a combined item each resource inherits the waste of ITS recipe, also
        after the cascade; and a waste set later on the budget wins for all of them."""
        _fix_eps(db)
        waste = {"tmpl-8.3": 10, "tmpl-5.2.3": 25}
        for t in db.tables["item_templates"]:
            t["desperdicio_pct"] = waste.get(t["id"], t["desperdicio_pct"])
        assert cargar(client).status_code == 200
        item = next(i for i in db.tables["budget_items"] if i["code"] == "4.2.3")
        assert item["template_id"] == "tmpl-8.3"

        def res():
            return {r["codigo"]: r for r in db.tables["item_resources"] if r["item_id"] == item["id"]}

        assert (res()["EPS-500"]["template_id"], res()["ES"]["template_id"]) == ("tmpl-8.3", "tmpl-5.2.3")
        assert (res()["EPS-500"]["desperdicio_pct"], res()["EPS-500"]["desperdicio_origen"]) == (10, "plantilla")
        assert (res()["ES"]["desperdicio_pct"], res()["ES"]["desperdicio_origen"]) == (25, "plantilla")
        assert res()["ES"]["cantidad_efectiva"] == pytest.approx(res()["ES"]["cantidad"] * 1.25, rel=1e-3)

        # "Recalcular obra" keeps each recipe's %
        budget_id = db.tables["budgets"][0]["id"]
        assert client.post(f"/budgets/{budget_id}/cascade-recalculate").status_code == 200
        assert (res()["EPS-500"]["desperdicio_pct"], res()["ES"]["desperdicio_pct"]) == (10, 25)

        # A waste set on the budget wins over both recipes (presupuesto > receta > organización)
        db.tables["budgets"][0]["desperdicio_pct"] = 7
        assert client.post(f"/budgets/{budget_id}/cascade-recalculate").status_code == 200
        assert (res()["EPS-500"]["desperdicio_pct"], res()["EPS-500"]["desperdicio_origen"]) == (7, "presupuesto")
        assert (res()["ES"]["desperdicio_pct"], res()["ES"]["desperdicio_origen"]) == (7, "presupuesto")

    def test_loads_after_fixing_prices(self, client, db):
        eps = next(e for e in db.tables["catalog_entries"] if e["codigo"] == "EPS-500")
        eps["precio_sin_iva"] = 10
        res = cargar(client)
        assert res.status_code == 200, res.text
        body = res.json()
        budget = db.tables["budgets"][0]
        assert budget["name"] == "OBRA PRUEBA"
        assert budget["precios_al"]

        items = {i["code"]: i for i in db.tables["budget_items"]}
        assert items["4.2.1"]["parent_id"] == items["4.2"]["id"]
        assert items["4.2"]["parent_id"] == items["4"]["id"]
        assert items["4.2.1"]["template_id"] == "tmpl-5.1.4"
        # Sin receta: precio del Excel
        assert items["4.2.2"]["mat_unitario"] == 5
        # Con receta: calculado desde los recursos por la cascada de la app
        lh18 = next(r for r in db.tables["item_resources"] if r["codigo"] == "LH18")
        assert lh18["precio_unitario"] == 500
        assert lh18["desperdicio_pct"] == 15.0
        hadn6 = next(r for r in db.tables["item_resources"] if r["codigo"] == "HADN6")
        assert (hadn6["desperdicio_pct"], hadn6["desperdicio_origen"]) == (5, "organizacion")
        assert items["4.2.1"]["mat_unitario"] > 0
        assert items["4.2.1"]["neto_total"] > items["4.2.1"]["directo_total"]  # indirectos aplicados
        assert body["con_receta"] == 5

    def test_forced_load_and_name_taken(self, client, db):
        assert cargar(client, permitir=True).status_code == 200
        eps = next(r for r in db.tables["item_resources"] if r["codigo"] == "EPS-500")
        assert eps["precio_unitario"] == 0
        res = cargar(client, permitir=True)
        assert res.status_code == 409
        assert "Ya existe" in res.text

    def test_unconfirmed_tasks_and_times(self, client, db):
        """Sol loads with yellows: the answer says how many went in unconfirmed (with the proposed
        recipe or the Excel price) and their notes end in "Para confirmar."; the greens do not."""
        _fix_eps(db)
        wb = _workbook()
        ws = wb["01_C&P"]
        # A rule with a question nobody answered: yellow with its proposed recipe
        row = ws.max_row + 1
        for col, value in ((1, "4.3-6"), (2, "MURO DE CARGA EN LADRILLO HUECO DEL 18"), (3, "m²"), (4, 10),
                           (5, 1), (10, 1), (14, 20), (26, 30)):
            ws.cell(row, col, value)
        tareas = _tareas(analizar(client, wb=wb).json())
        assert tareas[MURO_CARGA]["estado"] == "amarillo"
        assert tareas[MURO_CARGA]["receta"]["codigo"] == "5.1.4"

        res = cargar(client, wb=wb)
        assert res.status_code == 200, res.text
        body = res.json()
        # By the Excel's total, biggest first: OBRADOR 1500, ARISTAS 150, MURO DE CARGA 30
        assert body["sin_confirmar"] == {"total": 3, "con_receta": 1, "sin_receta": 2,
                                         "claves": [OBRADOR, ARISTAS, MURO_CARGA]}
        items = {i["code"]: i for i in db.tables["budget_items"]}
        for code in ("1.1", "4.2.2", "4.3.6"):
            assert items[code]["notas"].endswith(" Para confirmar."), code
        assert items["4.3.6"]["template_id"] == "tmpl-5.1.4"
        for code in ("3.1.3", "4.2.1", "4.2.3", "4.2.4", "4.3.5", "4", "4.2"):
            assert "Para confirmar" not in items[code]["notas"], code

        assert set(body["tiempos"]) == {"analisis_s", "items_s", "recursos_s", "cascada_s", "total_s"}
        assert all(isinstance(v, (int, float)) and v >= 0 for v in body["tiempos"].values())

        # Confirmed on the screen: nothing left to confirm
        res = cargar(client, nombre="OTRA", wb=wb, asignaciones={
            OBRADOR: {"plantillas": [], "confirmada": True},
            ARISTAS: {"plantillas": [], "confirmada": True},
            MURO_CARGA: {"plantillas": [["5.1.4", 1]], "confirmada": True},
        })
        assert res.json()["sin_confirmar"] == {"total": 0, "con_receta": 0, "sin_receta": 0, "claves": []}
        budget_id = res.json()["budget_id"]
        assert not any("Para confirmar" in i["notas"] for i in db.tables["budget_items"]
                       if i["budget_id"] == budget_id)

    def test_failure_leaves_nothing(self, client, db):
        with patch("app.routers.analysis._run_cascade", side_effect=RuntimeError("se cortó")):
            res = cargar(client, permitir=True)
        assert res.status_code == 500
        assert "no quedó nada a medias" in res.text
        assert db.tables["budgets"] == []


# ── Precios que el Excel ya trae y catálogo oficial ──────────────────────────


def _excel_with_prices():
    """The obra's Excel plus a detail sheet where Sol typed the prices the catalog lacks."""
    from tests.test_obra_import import _detail_sheet

    wb = _workbook()
    _detail_sheet(wb, "4.9-2", "4.9-2", "TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm", [
        ("MATERIALES", [("EPS-500", "Placa EPS 30 mm", "u", 10000), ("H30", "Hormigón H30", "m3", 90000)]),
    ])
    mat = wb.create_sheet("00_Mat")
    mat.append(["CODIGO", "MATERIALES", "UNIDAD", "PRECIO CON IVA", "PRECIO SIN IVA"])
    mat.append(["EPS-500", "Placa EPS", "u", None, 8000, "ML", datetime(2026, 9, 17)])
    return wb


def _drop_entry(db, codigo):
    db.tables["catalog_entries"] = [e for e in db.tables["catalog_entries"] if e["codigo"] != codigo]


def _consulta(db, rows, catalog=("cat-lh", "Las Heras", "2025-01-01")):
    """A reference catalog (old obra) with [(id, codigo, tipo, precio, fecha)]."""
    cid, name, created = catalog
    if not any(c["id"] == cid for c in db.tables["price_catalogs"]):
        db.tables["price_catalogs"].append({"id": cid, "org_id": ORG, "name": name, "created_at": created,
                                            "oficial": False})
    for eid, codigo, tipo, precio, fecha in rows:
        db.tables["catalog_entries"].append({"id": eid, "catalog_id": cid, "org_id": ORG, "codigo": codigo,
                                             "descripcion": f"{codigo} viejo", "unidad": "u", "tipo": tipo,
                                             "precio_sin_iva": precio, "fecha_precio": fecha})


def _oficial(db, cid="cat1"):
    for c in db.tables["price_catalogs"]:
        c.setdefault("oficial", False)
        if c["id"] == cid:
            c["oficial"] = True


class TestPreciosDelExcel:
    def test_proposal_for_missing_and_priceless_codes(self, client, db):
        _drop_entry(db, "H30")
        precios = {p["codigo"]: p for p in analizar(client, wb=_excel_with_prices()).json()["precios"]}
        eps, h30 = precios["EPS-500"], precios["H30"]
        assert (eps["problema"], h30["problema"]) == ("sin_precio", "no_esta")
        assert h30["motivo"] == "No está en el catálogo"
        # The detail sheet wins over the 00_Mat list, which stays as "otro"
        assert eps["propuesta"] == {
            "codigo": "EPS-500", "descripcion": "Placa EPS 30 mm", "unidad": "u", "tipo": "material",
            "precio": 10000.0, "fecha": None, "proveedor": None, "nota": None, "origen": "detalle", "hoja": "4.9-2",
            "trabajo": "TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm",
            "otros": [{"precio": 8000.0, "hoja": "00_Mat"}],
        }
        assert h30["propuesta"]["precio"] == 90000.0
        assert eps["referencias"] == [] and h30["catalogo_destino"] is None  # no oficial catalog

    def test_no_proposal_for_a_repeated_code(self, client, db):
        for e in db.tables["catalog_entries"]:
            if e["codigo"] == "EPS-500":
                e["fecha_precio"] = None
        db.tables["catalog_entries"].append({"id": "e-dup", "catalog_id": "cat1", "org_id": ORG,
                                             "codigo": "EPS-500", "tipo": "material", "precio_sin_iva": 0,
                                             "fecha_precio": None})
        precios = {p["codigo"]: p for p in analizar(client, wb=_excel_with_prices()).json()["precios"]}
        assert precios["EPS-500"]["problema"] == "duplicado"
        assert precios["EPS-500"]["propuesta"] is None

    def test_blocked_load_also_proposes(self, client, db):
        res = cargar(client, wb=_excel_with_prices())
        assert res.status_code == 409
        assert res.json()["detail"]["precios"][0]["propuesta"]["precio"] == 10000.0


class TestCatalogoOficial:
    def test_references_only_with_an_oficial_catalog(self, client, db):
        _consulta(db, [("lh-eps", "EPS-500", "material", 6592, "2026-03-12")])
        body = analizar(client).json()
        assert body["catalogo_oficial"] is False
        # Without oficial every catalog prices: the Las Heras copy has a price and is used
        assert body["precios"] == []

        _oficial(db)
        body = analizar(client).json()
        assert body["catalogo_oficial"] is True
        eps = {p["codigo"]: p for p in body["precios"]}["EPS-500"]
        assert eps["problema"] == "sin_precio"
        assert [e["id"] for e in eps["entradas"]] == ["e-EPS-500"]  # only the oficial entry
        assert eps["referencias"] == [{
            "id": "lh-eps", "catalog_id": "cat-lh", "catalogo": "Las Heras", "codigo": "EPS-500",
            "descripcion": "EPS-500 viejo", "unidad": "u", "tipo": "material", "precio_sin_iva": 6592,
            "fecha_precio": "2026-03-12",
        }]

    def test_references_newest_first_at_most_five(self, client, db):
        _oficial(db)
        _consulta(db, [
            ("r-old", "EPS-500", "material", 1, "2024-01-01"),
            ("r-none", "EPS-500", "material", 2, None),
            ("r-new", "EPS-500", "material", 3, "2026-05-01"),
            ("r-zero", "EPS-500", "material", 0, "2026-09-01"),  # no price: not a reference
            ("r-mid", "EPS-500", "material", 4, "2025-06-01"),
            ("r-mid2", "eps-500", "material", 5, "2025-01-01"),
            ("r-mid3", "EPS-500", "material", 6, "2024-06-01"),
        ])
        eps = {p["codigo"]: p for p in analizar(client).json()["precios"]}["EPS-500"]
        assert [r["id"] for r in eps["referencias"]] == ["r-new", "r-mid", "r-mid2", "r-mid3", "r-old"]

    def test_code_only_in_a_reference_catalog_is_missing(self, client, db):
        """With an oficial catalog, the old catalogs do not price, even when the code is only there."""
        _drop_entry(db, "H30")
        _consulta(db, [("lh-h30", "H30", "material", 95000, "2026-02-01")])
        _oficial(db)
        db.tables["price_catalogs"].append({"id": "cat-mo", "org_id": ORG, "name": "Maestro - Mano de obra",
                                            "created_at": "2026-09-01", "oficial": True})
        db.tables["catalog_entries"].append({"id": "e-mo2", "catalog_id": "cat-mo", "org_id": ORG,
                                             "codigo": "MO-AY", "tipo": "mano_obra", "precio_sin_iva": 80,
                                             "fecha_precio": "2026-01-01"})
        h30 = {p["codigo"]: p for p in analizar(client, wb=_excel_with_prices()).json()["precios"]}["H30"]
        assert (h30["problema"], h30["motivo"]) == ("no_esta", "No está en el catálogo oficial")
        assert h30["entradas"] == []
        assert [r["id"] for r in h30["referencias"]] == ["lh-h30"]
        assert h30["propuesta"]["precio"] == 90000.0
        # Created where most materials are (cat1), not in the newer oficial catalog of labor
        assert h30["catalogo_destino"] == {"id": "cat1", "name": "Maestro TERRAC - Materiales"}

    def test_destination_tie_goes_to_the_newest_oficial(self, client, db):
        _drop_entry(db, "H30")
        _oficial(db)
        db.tables["price_catalogs"][0]["created_at"] = "2026-01-01"
        db.tables["price_catalogs"].append({"id": "cat-new", "org_id": ORG, "name": "Maestro nuevo",
                                            "created_at": "2026-09-01", "oficial": True})
        for e in db.tables["catalog_entries"]:
            if e["tipo"] == "material":
                e["catalog_id"] = "cat-new" if e["codigo"] in ("LH18", "HADN6", "ES", "RP-KP") else "cat1"
        # 4 materials each: cat-new LH18, HADN6, ES, RP-KP; cat1 EPS-500, RP-PORC and these two
        db.tables["catalog_entries"] += [
            {"id": f"x{i}", "catalog_id": "cat1", "org_id": ORG, "codigo": f"X{i}", "tipo": "material",
             "precio_sin_iva": 1, "fecha_precio": "2026-01-01"} for i in range(2)
        ]
        h30 = {p["codigo"]: p for p in analizar(client).json()["precios"]}["H30"]
        assert h30["catalogo_destino"] == {"id": "cat-new", "name": "Maestro nuevo"}

    def test_load_uses_only_oficial_prices(self, client, db):
        _fix_eps(db)
        _consulta(db, [("lh-lh18", "LH18", "material", 900, "2026-09-01")])  # newer, but only a reference
        _oficial(db)
        assert cargar(client).status_code == 200
        lh18 = next(r for r in db.tables["item_resources"] if r["codigo"] == "LH18")
        assert (lh18["precio_unitario"], lh18["catalog_entry_id"]) == (500, "e-LH18")

    def test_update_prices_relinks_a_reference_entry_by_code(self, client, db):
        """A resource linked to an entry of a reference catalog takes the oficial one by code."""
        _fix_eps(db)
        assert cargar(client).status_code == 200
        _consulta(db, [("lh-lh18", "LH18", "material", 900, "2026-01-01")])
        lh18 = next(r for r in db.tables["item_resources"] if r["codigo"] == "LH18")
        lh18["catalog_entry_id"], lh18["precio_unitario"] = "lh-lh18", 900
        _oficial(db)
        budget_id = db.tables["budgets"][0]["id"]
        res = client.post(f"/budgets/{budget_id}/actualizar-precios")
        assert res.status_code == 200, res.text
        lh18 = next(r for r in db.tables["item_resources"] if r["id"] == lh18["id"])
        assert (lh18["precio_unitario"], lh18["catalog_entry_id"]) == (500, "e-LH18")


# ── Diferencias con el Excel ─────────────────────────────────────────────────

BUDGET_ID = "00000000-0000-0000-0000-0000000000d1"


def _item(id_, code, desc, unidad, cantidad, parent=None, neto=0.0, excel=None, directo=0.0, excel_dir=None,
          template=None, seccion=False, sort=0):
    return {"id": id_, "budget_id": BUDGET_ID, "org_id": ORG, "code": code, "description": desc,
            "unidad": unidad, "cantidad": cantidad, "parent_id": parent, "neto_total": neto,
            "excel_neto": excel, "directo_total": directo, "excel_directo": excel_dir, "template_id": template,
            "notas": "Seccion" if seccion else "Receta del Maestro", "sort_order": sort}


def _budget_with_excel(db):
    db.tables["budgets"].append({"id": BUDGET_ID, "org_id": ORG, "name": "EDIFICIO GINKGO",
                                 "precios_al": "2026-10-03", "source_file": "ginkgo.xlsx"})
    db.tables["budget_items"] += [
        _item("s4", "4", "ALBAÑILERIA", None, None, seccion=True, sort=0),
        _item("i-obr", "1.1", "OBRADOR", "gl", 1, parent="s4", neto=1000, excel=1500, directo=800,
              excel_dir=1200, sort=1),
        _item("p1", "4.2", "PRIMER PISO", None, None, parent="s4", seccion=True, sort=2),
        # Listed out of order on purpose: items come by sort_order
        _item("i-m2", "4.3.1", "MURO HUECO 18", "m2", 50, parent="p2", neto=600, excel=500, directo=450,
              excel_dir=400, template="tmpl-5.1.4", sort=6),
        _item("i-m1", "4.2.1", "MURO HUECO 18", "m²", 100, parent="p1", neto=1200, excel=1000, directo=900,
              excel_dir=800, template="tmpl-5.1.4", sort=3),
        _item("i-lim", "4.2.2", "LIMPIEZA", "gl", 1, parent="p1", neto=1030, excel=1000, sort=4),
        _item("p2", "4.3", "SEGUNDO PISO", None, None, parent="s4", seccion=True, sort=5),
        _item("i-ay", "4.3.2", "AYUDA DE GREMIOS", "gl", 1, parent="p2", neto=50, excel=0, sort=7),
        _item("i-cero", "4.3.3", "CERO", "gl", 0, parent="p2", neto=0, excel=0, sort=8),
        _item("i-viejo", "4.3.4", "AGREGADO A MANO", "gl", 1, parent="p2", neto=999, excel=None, sort=9),
    ]


def diferencias(client, budget_id=BUDGET_ID):
    return client.get(f"/obras/{budget_id}/diferencias")


class TestDiferencias:
    def test_saves_the_excel_totals_when_loading(self, client, db):
        _fix_eps(db)
        assert cargar(client).status_code == 200
        items = {i["code"]: i for i in db.tables["budget_items"]}
        assert (items["4.2.1"]["excel_directo"], items["4.2.1"]["excel_neto"]) == (9000, 12000)
        assert (items["1.1"]["excel_directo"], items["1.1"]["excel_neto"]) == (1000, 1500)
        assert (items["4"]["excel_neto"], items["4.2"]["excel_directo"]) == (None, None)  # rubros y pisos

        body = diferencias(client, items["4.2.1"]["budget_id"]).json()
        assert body["total"]["excel_neto"] == 13850.0
        piso = next(t for t in body["trabajos"] if t["descripcion"].startswith("COLOCACION DE REVESTIMIENTOS"))
        assert [(i["code"], i["piso"]) for i in piso["items"]] == [("4.2.4", "PRIMER PISO"), ("4.3.5", "SEGUNDO PISO")]

    def test_compares_task_by_task(self, client, db):
        _budget_with_excel(db)
        res = diferencias(client)
        assert res.status_code == 200, res.text
        body = res.json()
        assert set(body) == {"budget_id", "nombre", "precios_al", "source_file", "nivel_excel", "total", "resumen",
                             "trabajos"}
        assert (body["budget_id"], body["nombre"], body["precios_al"], body["source_file"]) == (
            BUDGET_ID, "EDIFICIO GINKGO", "2026-10-03", "ginkgo.xlsx")
        # The hand-added item without Excel totals is left out
        assert body["total"] == {"app_neto": 3880.0, "excel_neto": 4000.0, "diferencia": -120.0,
                                 "diferencia_pct": -3.0, "app_directo": 2150.0, "excel_directo": 2400.0,
                                 "diferencia_directo": -250.0, "diferencia_directo_pct": -10.4,
                                 "margen_app_pct": 80.5, "margen_excel_pct": 66.7,
                                 # This Excel adds 66.7 %: closest to the app's final price (59.5 %)
                                 "app_nivel": 3880.0, "diferencia_nivel": -120.0, "diferencia_nivel_pct": -3.0}
        assert body["nivel_excel"]["nivel"] == "neto"
        assert body["resumen"] == {"trabajos": 5, "mas_caros": 2, "mas_baratos": 1, "parecidos": 2,
                                   "sin_receta": 4,
                                   "directo": {"mas_caros": 1, "mas_baratos": 1, "parecidos": 3},
                                   "nivel": {"mas_caros": 2, "mas_baratos": 1, "parecidos": 2},
                                   "app_neto": 3880.0}
        # Biggest difference first (in absolute value)
        assert [t["descripcion"] for t in body["trabajos"]] == [
            "OBRADOR", "MURO HUECO 18", "AYUDA DE GREMIOS", "LIMPIEZA", "CERO"]

        muro = body["trabajos"][1]
        assert set(muro) == {"clave", "descripcion", "unidad", "veces", "cantidad_total", "receta", "sin_receta",
                             "app_neto", "excel_neto", "diferencia", "diferencia_pct", "app_directo",
                             "excel_directo", "app_unitario", "excel_unitario", "items",
                             "diferencia_directo", "diferencia_directo_pct", "margen_app_pct", "margen_excel_pct",
                             "app_unitario_directo", "excel_unitario_directo",
                             "app_nivel", "diferencia_nivel", "diferencia_nivel_pct", "app_unitario_nivel"}
        assert muro["clave"] == "MURO HUECO 18 | m2"  # m² and m2 are the same task
        assert (muro["veces"], muro["cantidad_total"]) == (2, 150.0)
        assert (muro["receta"], muro["sin_receta"]) == ({"codigo": "5.1.4", "nombre": "Receta 5.1.4"}, False)
        assert (muro["app_neto"], muro["excel_neto"], muro["diferencia"], muro["diferencia_pct"]) == (
            1800.0, 1500.0, 300.0, 20.0)
        assert (muro["app_directo"], muro["excel_directo"]) == (1350.0, 1200.0)
        assert (muro["app_unitario"], muro["excel_unitario"]) == (12.0, 10.0)
        assert muro["items"] == [
            {"id": "i-m1", "code": "4.2.1", "piso": "PRIMER PISO", "cantidad": 100.0, "app_neto": 1200.0,
             "excel_neto": 1000.0, "diferencia": 200.0, "app_directo": 900.0, "excel_directo": 800.0,
             "diferencia_directo": 100.0, "app_nivel": 1200.0, "diferencia_nivel": 200.0},
            {"id": "i-m2", "code": "4.3.1", "piso": "SEGUNDO PISO", "cantidad": 50.0, "app_neto": 600.0,
             "excel_neto": 500.0, "diferencia": 100.0, "app_directo": 450.0, "excel_directo": 400.0,
             "diferencia_directo": 50.0, "app_nivel": 600.0, "diferencia_nivel": 100.0},
        ]

        obrador, ayuda, cero = body["trabajos"][0], body["trabajos"][2], body["trabajos"][4]
        assert (obrador["receta"], obrador["sin_receta"], obrador["diferencia_pct"]) == (None, True, -33.3)
        assert obrador["items"][0]["piso"] == "ALBAÑILERIA"  # its direct parent
        assert (ayuda["diferencia"], ayuda["diferencia_pct"]) == (50.0, None)  # the Excel said 0
        assert (cero["app_unitario"], cero["excel_unitario"], cero["diferencia_pct"]) == (None, None, None)

    def test_direct_cost_and_margin(self, client, db):
        """Sol puts a different markup on each task: comparing the direct cost shows the recipes,
        and the margin says how much each one adds on top."""
        _budget_with_excel(db)
        trabajos = {t["descripcion"]: t for t in diferencias(client).json()["trabajos"]}
        muro = trabajos["MURO HUECO 18"]
        assert (muro["diferencia_directo"], muro["diferencia_directo_pct"]) == (150.0, 12.5)
        assert (muro["margen_app_pct"], muro["margen_excel_pct"]) == (33.3, 25.0)
        assert (muro["app_unitario_directo"], muro["excel_unitario_directo"]) == (9.0, 8.0)
        obrador = trabajos["OBRADOR"]
        assert (obrador["diferencia_directo"], obrador["diferencia_directo_pct"]) == (-400.0, -33.3)
        assert (obrador["margen_app_pct"], obrador["margen_excel_pct"]) == (25.0, 25.0)
        assert obrador["items"][0]["diferencia_directo"] == -400.0
        # Without a direct cost there is no margin; nor a % over an Excel that says 0
        limpieza = trabajos["LIMPIEZA"]
        assert (limpieza["margen_app_pct"], limpieza["margen_excel_pct"]) == (None, None)
        assert (limpieza["diferencia_directo"], limpieza["diferencia_directo_pct"]) == (0.0, None)
        cero = trabajos["CERO"]
        assert (cero["app_unitario_directo"], cero["excel_unitario_directo"]) == (None, None)

        # The 5 % rule by direct cost: 4 % is "parecido", 6 % is not; an Excel at 0 with cost is "más caro"
        items = {i["id"]: i for i in db.tables["budget_items"]}
        items["i-lim"].update(directo_total=1040, excel_directo=1000)
        items["i-ay"].update(directo_total=50, excel_directo=0)
        body = diferencias(client).json()
        assert body["resumen"]["directo"] == {"mas_caros": 2, "mas_baratos": 1, "parecidos": 2}
        items["i-lim"]["directo_total"] = 1060
        body = diferencias(client).json()
        assert body["resumen"]["directo"] == {"mas_caros": 3, "mas_baratos": 1, "parecidos": 1}
        # The final price summary and the order do not follow the direct cost
        assert (body["resumen"]["mas_caros"], body["resumen"]["mas_baratos"], body["resumen"]["parecidos"]) == (
            2, 1, 2)
        assert [t["descripcion"] for t in body["trabajos"]] == [
            "OBRADOR", "MURO HUECO 18", "AYUDA DE GREMIOS", "LIMPIEZA", "CERO"]

    def test_budget_loaded_before_saving_the_excel_totals(self, client, db):
        _budget_with_excel(db)
        for i in db.tables["budget_items"]:
            i.pop("excel_neto")
        res = diferencias(client)
        assert res.status_code == 409
        assert res.json()["detail"] == (
            'Este presupuesto no tiene guardados los totales del Excel. Cargá la obra de nuevo desde '
            '"Cargar obra" para poder compararla.')

    def test_budget_of_another_org(self, client, db):
        _budget_with_excel(db)
        db.tables["budgets"][-1]["org_id"] = "otra-org"
        assert diferencias(client).status_code == 404
        assert diferencias(client, "00000000-0000-0000-0000-00000000dead").status_code == 404


# ── Diferencias al nivel del Excel ───────────────────────────────────────────


def _cascada(directo):
    """(indirecto, beneficio, neto) with the default config: 34 % indirect, 10 % benefit, 8.2 % taxes."""
    indirecto = round(directo * 0.34, 2)
    beneficio = round((directo + indirecto) * 0.10, 2)
    return indirecto, beneficio, round((directo + indirecto + beneficio) * 1.082, 2)


def _budget_al_nivel(db, excel, **budget):
    """MURO (directo 1000) and REVOQUE (directo 500) calculated by the app's cascade;
    ``excel`` gives (excel_directo, excel_neto) for each one."""
    db.tables["budgets"].append({"id": BUDGET_ID, "org_id": ORG, "name": "OBRA", "precios_al": None,
                                 "source_file": "obra.xlsx", **budget})
    for (id_, desc, cantidad, directo, sort), (excel_dir, excel_neto) in zip(
            (("i-muro", "MURO", 100, 1000, 1), ("i-rev", "REVOQUE", 10, 500, 2)), excel):
        indirecto, beneficio, neto = _cascada(directo)
        db.tables["budget_items"].append({
            **_item(id_, f"1.{sort}", desc, "m2", cantidad, neto=neto, excel=excel_neto, directo=directo,
                    excel_dir=excel_dir, sort=sort),
            "indirecto_total": indirecto, "beneficio_total": beneficio})


class TestDiferenciasAlNivel:
    def test_excel_up_to_the_indirects(self, client, db):
        """Sol's Excel adds 34 % (Ginkgo): the app is compared up to its indirects, not its final price."""
        _budget_al_nivel(db, [(900, 1206), (500, 670)])
        body = diferencias(client).json()
        assert body["nivel_excel"] == {
            "nivel": "indirectos", "factor_excel": 1.34, "factor_app": 1.34,
            "texto": "Tu Excel le suma 34% al costo directo: llega hasta los indirectos. Comparo la app hasta ahí."}
        muro, revoque = body["trabajos"]
        # app_nivel = directo + indirecto, against the Excel's final price
        assert (muro["app_nivel"], muro["excel_neto"], muro["diferencia_nivel"], muro["diferencia_nivel_pct"]) == (
            1340.0, 1206.0, 134.0, 11.1)
        assert (muro["app_unitario_nivel"], muro["excel_unitario"]) == (13.4, 12.06)
        assert (revoque["app_nivel"], revoque["diferencia_nivel"], revoque["diferencia_nivel_pct"]) == (
            670.0, 0.0, 0.0)
        assert muro["items"][0]["app_nivel"] == 1340.0
        assert muro["items"][0]["diferencia_nivel"] == 134.0
        assert (body["total"]["app_nivel"], body["total"]["diferencia_nivel"],
                body["total"]["diferencia_nivel_pct"]) == (2010.0, 134.0, 7.1)
        assert body["resumen"]["nivel"] == {"mas_caros": 1, "mas_baratos": 0, "parecidos": 1}
        # The old fields do not change: final price and direct cost as always
        assert (muro["app_neto"], muro["diferencia"], muro["diferencia_pct"]) == (1594.87, 388.87, 32.2)
        assert (muro["app_directo"], muro["diferencia_directo"], muro["diferencia_directo_pct"]) == (
            1000.0, 100.0, 11.1)
        assert (muro["app_unitario"], muro["excel_unitario_directo"]) == (15.95, 9.0)
        assert body["total"]["app_neto"] == body["resumen"]["app_neto"] == 2392.30
        assert (body["resumen"]["mas_caros"], body["resumen"]["parecidos"]) == (2, 0)
        assert body["resumen"]["directo"] == {"mas_caros": 1, "mas_baratos": 0, "parecidos": 1}

    def test_excel_up_to_the_final_price(self, client, db):
        _budget_al_nivel(db, [(1000, 1594.87), (500, 797.43)])
        body = diferencias(client).json()
        assert body["nivel_excel"] == {
            "nivel": "neto", "factor_excel": 1.5949, "factor_app": 1.5949,
            "texto": "Tu Excel le suma 59,5% al costo directo: llega hasta los impuestos. Comparo la app hasta ahí."}
        for t in body["trabajos"]:
            assert (t["app_nivel"], t["diferencia_nivel"]) == (t["app_neto"], 0.0)
            assert t["app_unitario_nivel"] == t["app_unitario"]
        assert body["resumen"]["nivel"] == {"mas_caros": 0, "mas_baratos": 0, "parecidos": 2}

    def test_excel_up_to_the_benefit(self, client, db):
        _budget_al_nivel(db, [(1000, 1480), (500, 730)])  # 47.3 %: 34 % + 10 % on top
        body = diferencias(client).json()
        assert body["nivel_excel"]["nivel"] == "beneficio"
        assert body["nivel_excel"]["factor_app"] == 1.474
        assert body["nivel_excel"]["texto"] == (
            "Tu Excel le suma 47,3% al costo directo: llega hasta el beneficio. Comparo la app hasta ahí.")
        muro = body["trabajos"][0]
        assert (muro["app_nivel"], muro["diferencia_nivel"]) == (1474.0, -6.0)

    def test_excel_without_margin(self, client, db):
        _budget_al_nivel(db, [(900, 900), (500, 500)])
        body = diferencias(client).json()
        assert body["nivel_excel"] == {
            "nivel": "directo", "factor_excel": 1.0, "factor_app": 1.0,
            "texto": "Tu Excel no le suma nada al costo directo: comparo costo directo."}
        muro = body["trabajos"][0]
        assert (muro["app_nivel"], muro["diferencia_nivel"], muro["diferencia_nivel_pct"]) == (1000.0, 100.0, 11.1)
        assert (muro["app_nivel"], muro["diferencia_nivel"]) == (muro["app_directo"], muro["diferencia_directo"])

    def test_old_items_without_indirect_or_benefit(self, client, db):
        _budget_al_nivel(db, [(900, 1206), (500, 670)])
        muro = next(i for i in db.tables["budget_items"] if i["id"] == "i-muro")
        muro["indirecto_total"] = None
        muro.pop("beneficio_total")
        body = diferencias(client).json()
        assert body["nivel_excel"]["nivel"] == "indirectos"
        assert body["trabajos"][0]["app_nivel"] == 1000.0  # null counts as 0

    def test_factors_follow_the_budget_config(self, client, db):
        """The same % as the cascade: the obra's own values over the organization's, over the defaults."""
        db.tables["indirect_config"][0]["estructura_pct"] = 0  # the organization: 19 % of indirects
        _budget_al_nivel(db, [(1000, 1190), (500, 595)])
        nivel = diferencias(client).json()["nivel_excel"]
        assert (nivel["nivel"], nivel["factor_excel"], nivel["factor_app"]) == ("indirectos", 1.19, 1.19)
        assert nivel["texto"].startswith("Tu Excel le suma 19% al costo directo: llega hasta los indirectos.")
        db.tables["budgets"][0]["indirectos"] = {"estructura_pct": 15}  # this obra: 34 % again
        nivel = diferencias(client).json()["nivel_excel"]
        assert (nivel["nivel"], nivel["factor_excel"], nivel["factor_app"]) == ("indirectos", 1.19, 1.34)
        db.tables["budgets"][0]["indirectos"] = {"estructura_pct": 0, "beneficio_pct": 0, "ingresos_brutos_pct": 0,
                                                 "imp_cheque_pct": 0, "imprevistos_pct": 0, "jefatura_pct": 0,
                                                 "logistica_pct": 0, "herramientas_pct": 19}
        assert diferencias(client).json()["nivel_excel"]["factor_app"] == 1.19  # an explicit 0 is respected

    def test_order_by_mode(self, client, db):
        _budget_al_nivel(db, [(900, 1206), (500, 670)])
        db.tables["budget_items"].append({  # its final price is close to the Excel, its indirects are not
            **_item("i-ay", "1.3", "AYUDA", "gl", 1, neto=300, excel=134, directo=100, excel_dir=100, sort=3),
            "indirecto_total": 300, "beneficio_total": 0})
        orden = {modo: [t["descripcion"] for t in client.get(
            f"/obras/{BUDGET_ID}/diferencias", params={"modo": modo} if modo else None).json()["trabajos"]]
            for modo in (None, "neto", "directo", "nivel")}
        assert orden[None] == orden["neto"] == ["MURO", "AYUDA", "REVOQUE"]
        assert orden["directo"][0] == "MURO"
        assert orden["nivel"] == ["AYUDA", "MURO", "REVOQUE"]
        assert client.get(f"/obras/{BUDGET_ID}/diferencias", params={"modo": "otro"}).status_code == 422

    def test_excel_without_direct_cost(self, client, db):
        _budget_al_nivel(db, [(0, 1206), (0, 670)])
        nivel = diferencias(client).json()["nivel_excel"]
        assert (nivel["nivel"], nivel["factor_excel"]) == ("neto", None)
        assert nivel["texto"] == "Tu Excel no trae el costo directo: comparo el precio final."


# ── Excel sin precios ────────────────────────────────────────────────────────

LIMPIEZA = "LIMPIEZA FINAL DE OBRA | gl"
SIN_RECETA_Y_SIN_PRECIO = "Hay {n} trabajos sin fórmula y sin precio en el Excel: elegí una fórmula antes de cargar"


def _workbook_sin_precios():
    """The test obra with quantities only: columns E, J, N and Z empty (no detail sheets)."""
    wb = _workbook()
    ws = wb["01_C&P"]
    for r in range(8, ws.max_row + 1):
        for col in (5, 10, 14, 26):
            ws.cell(r, col).value = None
    return wb


def _aristas_recipe(db):
    """A recipe that looks like ARISTAS, so the card has something to suggest ("Quizás sea")."""
    db.tables["item_templates"].append({"id": "tmpl-6.9", "org_id": ORG, "codigo": "6.9", "unidad": "m",
                                        "nombre": "ARISTAS DE YESO", "parametros": [], "recursos": []})


class TestExcelSinPrecios:
    def test_tasks_without_recipe_are_red(self, client, db):
        _aristas_recipe(db)
        res = analizar(client, wb=_workbook_sin_precios())
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["excel_con_precios"] is False
        assert body["resumen"]["total_excel"] == 0
        tareas = _tareas(body)

        # With a recipe: as always (green, or red for a missing catalog price)
        assert (tareas[TENSORES]["estado"], tareas[TENSORES]["motivo_rojo"]) == ("verde", None)
        assert (tareas[TELGOPOR]["estado"], tareas[TELGOPOR]["motivo_rojo"]) == ("rojo", "precio")
        assert tareas[TELGOPOR]["precios_faltantes"] == ["EPS-500"]

        # Without a recipe: red, the Excel has no price to use; "Quizás sea" still comes
        for clave in (OBRADOR, ARISTAS):
            t = tareas[clave]
            assert (t["estado"], t["motivo_rojo"], t["receta"], t["total_excel"]) == ("rojo", "sin_receta", None, 0)
        assert tareas[ARISTAS]["sugerencias"][0]["codigo"] == "6.9"
        assert body["listo"] is False

        # Choosing "the Excel price" on the screen does not help: it is $0
        body = analizar(client, {OBRADOR: {"plantillas": [], "confirmada": True}}, wb=_workbook_sin_precios()).json()
        assert _tareas(body)[OBRADOR]["motivo_rojo"] == "sin_receta"

    def test_any_cost_counts_as_prices(self, client, db):
        wb = _workbook_sin_precios()
        wb["01_C&P"].cell(9, 10).value = 1000  # OBRADOR's M.O. per unit (J), nothing else
        assert analizar(client, wb=wb).json()["excel_con_precios"] is True

    def test_load_blocked_until_a_recipe_is_chosen(self, client, db):
        _fix_eps(db)
        _aristas_recipe(db)
        res = cargar(client, wb=_workbook_sin_precios())
        assert res.status_code == 409
        detail = res.json()["detail"]
        assert detail["mensaje"] == SIN_RECETA_Y_SIN_PRECIO.format(n=2)
        assert sorted(detail["rojos"]) == sorted([OBRADOR, ARISTAS])
        assert db.tables["budgets"] == []

        res = cargar(client, wb=_workbook_sin_precios(), asignaciones={ARISTAS: {"plantillas": [["6.9", 1]]}})
        assert res.status_code == 409
        assert res.json()["detail"] == {
            "mensaje": "Hay 1 trabajo sin fórmula y sin precio en el Excel: elegí una fórmula antes de cargar",
            "rojos": [OBRADOR]}

        elegidas = {ARISTAS: {"plantillas": [["6.9", 1]]}, OBRADOR: {"plantillas": [["5.1.4", 1]]}}
        res = cargar(client, wb=_workbook_sin_precios(), asignaciones=elegidas)
        assert res.status_code == 200, res.text
        assert res.json()["total_excel"] == 0
        items = {i["code"]: i for i in db.tables["budget_items"]}
        assert (items["1.1"]["template_id"], items["4.2.2"]["template_id"]) == ("tmpl-5.1.4", "tmpl-6.9")
        assert (items["4.2.1"]["excel_neto"], items["4.2.1"]["excel_directo"]) == (0, 0)

        # Nothing to compare against
        res = diferencias(client, items["1.1"]["budget_id"])
        assert res.status_code == 409
        assert res.json()["detail"] == "Este Excel no traía precios: no hay con qué comparar."

    def test_excel_with_prices_and_a_task_at_zero(self, client, db):
        """The usual Excel: nothing changes; a task without recipe whose row says $0 stays yellow."""
        wb = _workbook()
        ws = wb["01_C&P"]
        row = ws.max_row + 1
        for col, value in ((1, "4.3-6"), (2, "LIMPIEZA FINAL DE OBRA"), (3, "gl"), (4, 1), (5, 0), (10, 0),
                           (14, 0), (26, 0)):
            ws.cell(row, col, value)
        body = analizar(client, wb=wb).json()
        assert body["excel_con_precios"] is True
        tareas = _tareas(body)
        # Sol priced that row at $0 on purpose: yellow (confirmable), not red
        assert (tareas[LIMPIEZA]["estado"], tareas[LIMPIEZA]["motivo_rojo"]) == ("amarillo", None)
        assert tareas[LIMPIEZA]["total_excel"] == 0
        # The rest as in test_contract
        assert (tareas[OBRADOR]["estado"], tareas[OBRADOR]["motivo_rojo"]) == ("amarillo", None)
        assert (tareas[ARISTAS]["estado"], tareas[ARISTAS]["motivo_rojo"]) == ("amarillo", None)
        assert (tareas[TELGOPOR]["estado"], tareas[TELGOPOR]["motivo_rojo"]) == ("rojo", "precio")
        assert tareas[TENSORES]["estado"] == "verde"  # $0 in the Excel, but it has a recipe
        assert body["resumen"]["total_excel"] == 13850.0

        # It loads like any other yellow: with the Excel price, which is $0 (Sol's decision)
        _fix_eps(db)
        res = cargar(client, wb=wb)
        assert res.status_code == 200
        limpieza = next(i for i in db.tables["budget_items"] if i["description"] == "LIMPIEZA FINAL DE OBRA")
        assert (limpieza["neto_total"], limpieza["excel_neto"], limpieza["template_id"]) == (0, 0, None)
        assert limpieza["notas"].endswith("Para confirmar.")

    def test_differences_with_some_excel_prices_still_compare(self, client, db):
        _budget_with_excel(db)
        for i in db.tables["budget_items"]:
            if i["id"] != "i-lim" and i.get("excel_neto") is not None:
                i.update(excel_neto=0, excel_directo=0)
        assert diferencias(client).status_code == 200
        next(i for i in db.tables["budget_items"] if i["id"] == "i-lim")["excel_neto"] = 0
        res = diferencias(client)
        assert res.status_code == 409
        assert res.json()["detail"] == "Este Excel no traía precios: no hay con qué comparar."
