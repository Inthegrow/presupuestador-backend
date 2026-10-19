"""Lo que encontró Emilia el 19/10 (PLAN_PRUEBAS_EMILIA, entrega 11), del lado del servidor.

Planillas chicas armadas acá con openpyxl: dónde está el cómputo (find_computo), títulos de rubro y piso,
códigos, notas, la regla «código del modelo», trabajos sin fórmula y sin precio ($0), Importar Excel que
manda a Cargar obra, borrar presupuestos, el modelo de Terrac (Trabajos típicos) y el precio viejo.
Al final, las mismas pruebas con los archivos de Emilia si están en EXCEL_DIR (si no, se saltean).
"""

from __future__ import annotations

import copy
import inspect
import io
import json
import os
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.budget_prices import PRECIO_VIEJO_DIAS, es_precio_viejo, precio_dias, precio_viejo_dias
from app.main import create_app
from app.obra_import import (
    NOTA_SIN_PRECIO,
    code_recipe,
    find_computo,
    parse_any,
    parse_obra,
    section_title,
)
from tests.test_obras_api import _tables as obra_tables
from tests.test_obras_api import _tareas, analizar, cargar
from tests.test_recipes_api import (
    BUDGET,
    ITEM,
    MOCK_USER,
    ORG,
    RESOURCE_DEFAULTS,
    TEMPLATE,
    FakeDB,
    base_tables,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LEADER = {**MOCK_USER, "role": "leader", "orgs": [{**MOCK_USER["orgs"][0], "role": "leader"}]}


def _client(user=None) -> TestClient:
    app = create_app()
    from app.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: user or MOCK_USER
    return TestClient(app)


@pytest.fixture
def client():
    return _client()


def _bytes(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write(ws, rows, start: int) -> None:
    """rows: (A, B, C, D, E, J, N, Z), from row ``start``; shorter tuples leave the rest empty."""
    for r, row in enumerate(rows, start=start):
        row = tuple(row) + (None,) * (8 - len(row))
        for col, value in zip((1, 2, 3, 4, 5, 10, 14, 26), row):
            if value is not None:
                ws.cell(r, col, value)


# ── Planillas como las de Emilia ─────────────────────────────────────────────


def hoja1_workbook(header=("CODIGO", "DESCRIPCIÓN", "UNIDAD", "CANTIDAD")) -> openpyxl.Workbook:
    """Su primera prueba: 4 columnas, solapa «Hoja1», títulos en la fila 1, códigos-fecha, sin precios."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(list(header))
    for row in [
        (None, "1- TAREAS PRELIMINARES", None, None),
        (datetime(2026, 1, 1), "OBRADOR", "gl", 1),
        (datetime(2026, 2, 1), "BANOS QUIMICOS (3)", "mes", 16),
        (None, "2- EXCAVACIONES", None, None),
        (datetime(2026, 1, 2), "PERFILADO DE EXCAVACIÓN DE BASES", "unidad", 35),
        (None, "3- ESTRUCTURA DE HORMIGÓN ARMADO", None, None),
        (None, "3.1- FUNDACIONES", None, None),
        ("3.1-1", "BASES AISLADAS 2,5m x 2,5m.", "unid", 35),
        ("3.1-3", "TENSORES 20 cm x 40 cm.", "ml", 100),
        (None, "3.2- SOBRE PLANTA BAJA", None, None),
        ("3.2-1", "ESTRUCTURA EN HORMIGON ARMADO. LOSA, VIGAS Y COLUMNAS", "m²", 414),
    ]:
        ws.append(list(row))
    return wb


def prueba3_workbook() -> openpyxl.Workbook:
    """Su «Prueba 3»: solapa 01_C&P, títulos de columna en la fila 8 y trabajos desde la 9, sin precios."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_C&P"
    _write(ws, [
        ("CODIGO", "DESCRIPCIÓN", "UNIDAD", "CANTIDAD"),
        (None, "1- TAREAS PRELIMINARES"),
        (date(2026, 1, 1), "OBRADOR", "gl", 1),
        (date(2026, 2, 1), "BANOS QUIMICOS (3)", "mes", 16),
        (None, "2- EXCAVACIONES"),
        (date(2026, 1, 2), "PERFILADO DE EXCAVACIÓN", "unidad", 35),
        (None, "3- ESTRUCTURA DE HOMIGÓN ARMADO"),
        (None, "3.1- FUNDACIONES"),
        ("3.1-1", "BASES AISLADAS 2,5m x 2,5m.", "unid", 35),
        (None, "3.2- SOBRE PLANTA BAJA"),
        ("3.2-1", "ESTRUCTURA EN HORMIGON ARMADO. LOSA, VIGAS Y COLUMNAS", "m²", 414),
        (None, "7 PINTURA"),
        ("7.1", "PINTURA EN PAREDES", "m²", 10),
    ], start=8)
    return wb


def _detail_sheet(wb, name: str, codigo: str, nombre: str) -> None:
    """A detail sheet of one work (Sol's layout): its materials and its labour, with prices."""
    ws = wb.create_sheet(name)
    ws["A3"], ws["B3"] = codigo, nombre
    header = ["Código", "Descripción", "Unidad", "Cantidad", "Dias", "Desperdicio", "Cantidad + Desperdicio",
              "Precio Unitario", "Subtotal"]
    for r, row in enumerate([
        ["MATERIALES"], header,
        ["H30", "Hormigon H30", "m3", 108, None, 0.1, 119, 169000, 20111000],
        ["TOTAL MATERIALES"],
        ["MANO DE OBRA - PERSONAS"], header,
        ["MO-OF", "Oficial", "u", 3, 7, 0.25, 27, 90000, 2430000],
        ["TOTAL"],
    ], start=5):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)


def modelo_workbook(sheet: str = "CYP") -> openpyxl.Workbook:
    """El modelo de Terrac: cómputo en «CYP», títulos de rubro en la columna A, una hoja por trabajo y 00_Mat."""
    wb = openpyxl.Workbook()
    wb.active.title = "Condiciones"
    mat = wb.create_sheet("00_Mat")
    mat.append(["MATERIALES"])
    mat.append(["Codigo", "Descripcion", "Unidad", "Con IVA", "Sin IVA"])
    mat.append(["H30", "Hormigon H30", "m3", 204490, 169000])
    ws = wb.create_sheet(sheet)
    ws["A1"] = "EDIFICIO LAS HERAS"
    for col, text in enumerate(("ITEM", "DESCRIPCIÓN ", "UNIDAD", "CANTIDAD", "SUBTOTAL 01: GASTOS DIRECTOS"), 1):
        ws.cell(5, col, text)
    _write(ws, [
        ("1- TAREAS PRELIMINARES",),
        ("1.1", "OBRADOR", "mes", 1, 0, 0, 0, 0),
        ("2- MOVIMIENTO DE SUELO ", None, None, 1.0),  # the title carries a 1 in CANTIDAD
        ("2.1", "EXCAVACIÓN MECÁNICA (SOLO EQUIPOS)", None, 1.0),
        ("4 - ESTRUCTURA",),
        ("4 -1 FUNDACIONES",),
        ("4.1.3", "Bases aisladas ", "m³", 108, 274466.12, 83922.74, 38705996.8, 43641655.99),
        ("4-3 SOBRE SUBSUELO",),
        (None, "Replicar puntos del subsuelo", None, 1.0),
        ("7 TERMINACIONES",),
        ("7-1 REVESTIMIENTO ",),
        ("7.1.1", "REVESTIMIENTO HORIZANTAL", "m2", 1, 206700, 345000, 551700, 739571.48),
        ("7-5 OTRAS TERMINACIONES",),
        (None, "Dejar un item libre para terminaciones adicionales", None, 1.0),
    ], start=8)
    _detail_sheet(wb, "4.1.3", "4.1.3", "BASES AISLADAS")
    return wb


def _filas(parsed: dict, nivel: str) -> list[tuple]:
    return [(f["codigo"], f["descripcion"]) for f in parsed["filas"] if f["nivel"] == nivel]


# ── Dónde está el cómputo (E3, E6) ──────────────────────────────────────────


class TestFindComputo:
    def test_cyp_with_titles_in_row_5_and_terrac_prices(self):
        info = find_computo(modelo_workbook())
        assert (info["hoja"], info["fila_titulos"], info["primera_fila"]) == ("CYP", 5, 6)
        assert info["columnas"] == {"codigo": 1, "descripcion": 2, "unidad": 3, "cantidad": 4}
        assert info["precios"] == {"mat_unit": 5, "mo_unit": 10, "directo": 14, "neto": 26}  # E/J/N/Z

    def test_01_cp_without_titles_row_keeps_the_terrac_layout(self):
        from tests.test_obra_import import _workbook
        info = find_computo(_workbook())
        assert (info["hoja"], info["fila_titulos"], info["primera_fila"]) == ("01_C&P", None, 8)
        assert info["precios"] is not None

    def test_01_cp_with_titles_in_row_8(self):
        info = find_computo(prueba3_workbook())
        assert (info["hoja"], info["fila_titulos"], info["primera_fila"]) == ("01_C&P", 8, 9)

    @pytest.mark.parametrize("name", ["C&P", "CYP E1", "cyp", "01_C&P (2)"])
    def test_other_names_of_the_cómputo(self, name):
        assert find_computo(modelo_workbook(sheet=name))["hoja"] == name

    def test_order_01_cp_then_cyp_then_cp(self):
        wb = modelo_workbook(sheet="C&P")
        wb.create_sheet("CYP")
        assert find_computo(wb)["hoja"] == "CYP"
        wb.create_sheet("01_C&P")
        assert find_computo(wb)["hoja"] == "01_C&P"

    @pytest.mark.parametrize("header", [
        ("CODIGO", "DESCRIPCIÓN", "UNIDAD", "CANTIDAD"),
        ("código", "descripcion", "unidad", "cantidad"),
        ("Ítem", "Descripción:", "Unid.", "Cantidad"),
        (None, "DESCRIPCION", None, "CANTIDAD"),  # Código and Unidad are optional
    ])
    def test_any_sheet_with_descripcion_and_cantidad(self, header):
        info = find_computo(hoja1_workbook(header))
        assert (info["hoja"], info["fila_titulos"], info["primera_fila"]) == ("Hoja1", 1, 2)
        assert info["precios"] is None  # not the TERRAC layout: no price columns

    def test_titles_row_within_the_first_15_rows(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Obra"
        ws.cell(15, 2, "Descripción")
        ws.cell(15, 3, "Cantidad")
        assert find_computo(wb)["primera_fila"] == 16
        wb2 = openpyxl.Workbook()
        wb2.active.cell(16, 2, "Descripción")
        wb2.active.cell(16, 3, "Cantidad")
        assert find_computo(wb2) is None

    def test_price_lists_and_detail_sheets_are_never_the_cómputo(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Portada"
        _detail_sheet(wb, "4.1.3", "4.1.3", "BASES")
        _detail_sheet(wb, "3.1-1", "3.1.1", "BASES")
        lista = wb.create_sheet("00_Mat")
        lista.append(["Codigo", "Descripcion", "Unidad", "Cantidad"])
        assert find_computo(wb) is None


# ── Títulos, códigos y notas (E4, E6) ───────────────────────────────────────


class TestTitlesAndCodes:
    @pytest.mark.parametrize("text,expected", [
        ("1- TAREAS PRELIMINARES", ("1", "TAREAS PRELIMINARES")),
        ("1 - TAREAS PRELIMINARES", ("1", "TAREAS PRELIMINARES")),
        ("4 -1 FUNDACIONES", ("4.1", "FUNDACIONES")),
        ("4-2 SUBSUELO", ("4.2", "SUBSUELO")),
        ("3.1- FUNDACIONES", ("3.1", "FUNDACIONES")),
        ("7 TERMINACIONES", ("7", "TERMINACIONES")),
        ("2- MOVIMIENTO DE SUELO ", ("2", "MOVIMIENTO DE SUELO")),
        ("4. ALBAÑILERIA", ("4", "ALBAÑILERIA")),
        ("5.  TERMINACIONES", ("5", "TERMINACIONES")),
        ("3.10- SOBRE SALA DE MAQUINAS", ("3.10", "SOBRE SALA DE MAQUINAS")),
        ("Replicar puntos del subsuelo", None),
        ("TOTAL DEL PRESUPUESTO", None),
        ("4.1.1", None),
        ("12", None),
    ])
    def test_section_title(self, text, expected):
        assert section_title(text) == expected

    def test_titles_in_column_a_of_the_model(self):
        parsed = parse_obra(modelo_workbook())
        assert parsed["hoja"] == "CYP" and parsed["titulo"] == "EDIFICIO LAS HERAS"
        assert _filas(parsed, "rubro") == [("1", "TAREAS PRELIMINARES"), ("2", "MOVIMIENTO DE SUELO"),
                                           ("4", "ESTRUCTURA"), ("7", "TERMINACIONES")]
        assert _filas(parsed, "subrubro") == [("4.1", "FUNDACIONES"), ("4.3", "SOBRE SUBSUELO"),
                                              ("7.1", "REVESTIMIENTO"), ("7.5", "OTRAS TERMINACIONES")]
        by_code = {f["codigo"]: f for f in parsed["filas"]}
        assert by_code["4.1.3"]["parent"] == by_code["4.1"]["orden"]
        assert by_code["4.1"]["parent"] == by_code["4"]["orden"]
        assert by_code["2.1"]["parent"] == by_code["2"]["orden"]
        assert by_code["4.1.3"]["excel"]["neto"] == 43641655.99

    def test_notes_are_not_works(self):
        parsed = parse_obra(modelo_workbook())
        descripciones = [f["descripcion"] for f in parsed["filas"]]
        assert not any(d.startswith(("Replicar", "Dejar un item")) for d in descripciones)
        avisos = [p for p in parsed["problemas"] if "es una nota, no un trabajo" in p]
        assert len(avisos) == 2
        assert "Replicar puntos del subsuelo" in avisos[0] and "Dejar un item libre" in avisos[1]
        assert len(_filas(parsed, "item")) == 4

    def test_dates_column_titles_and_dashes(self):
        parsed = parse_obra(prueba3_workbook())
        items = _filas(parsed, "item")
        # The column titles row is never a work; dates are día.mes; 3.1-1 is 3.1.1
        assert ("CODIGO", "DESCRIPCIÓN") not in items and all(d != "DESCRIPCIÓN" for _, d in items)
        assert [c for c, _ in items] == ["1.1", "1.2", "2.1", "3.1.1", "3.2.1", "7.1"]
        assert _filas(parsed, "subrubro") == [("3.1", "FUNDACIONES"), ("3.2", "SOBRE PLANTA BAJA")]
        assert ("7", "PINTURA") in _filas(parsed, "rubro")
        by_code = {f["codigo"]: f for f in parsed["filas"]}
        # E12: each piso has its works
        assert by_code["3.1.1"]["parent"] == by_code["3.1"]["orden"]
        assert by_code["3.2.1"]["parent"] == by_code["3.2"]["orden"]
        assert parsed["titulo"] == ""

    def test_date_codes(self):
        from app.obra_import import _code
        assert _code(date(2026, 2, 1)) == "1.2"
        assert _code(date(2026, 1, 2)) == "2.1"
        assert _code(datetime(2026, 1, 1)) == "1.1"
        assert _code("3.1-1") == "3.1.1"

    def test_hoja1_reads_like_prueba3(self):
        a, b = parse_any(hoja1_workbook()), parse_obra(prueba3_workbook())
        assert a["hoja"] == "Hoja1" and a["titulo"] == ""
        assert [c for c, _ in _filas(a, "item")] == ["1.1", "1.2", "2.1", "3.1.1", "3.1.3", "3.2.1"]
        assert _filas(a, "subrubro") == _filas(b, "subrubro")
        assert all(not any(f["excel"].values()) for f in a["filas"] if f["nivel"] == "item")

    def test_code_with_name_and_nothing_to_price_is_a_title(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "01_C&P"
        _write(ws, [("1", "ALBAÑILERÍA"), ("1.1", "MURO", "m2", 10, 100, 50, 1500, 1980)], start=8)
        parsed = parse_obra(wb)
        assert _filas(parsed, "rubro") == [("1", "ALBAÑILERÍA")]
        assert _filas(parsed, "item") == [("1.1", "MURO")]

    def test_ginkgo_fixture_unchanged(self):
        """Regression (E16): the obra of tests/test_obra_import reads as always."""
        from tests.test_obra_import import _workbook
        parsed = parse_obra(_workbook())
        assert [(f["nivel"], f["codigo"]) for f in parsed["filas"]] == [
            ("rubro", "1"), ("item", "1.1"), ("rubro", "3"), ("subrubro", "3.1"), ("item", "3.1.3"),
            ("rubro", "4"), ("subrubro", "4.2"), ("item", "4.2.1"), ("item", "4.2.2"), ("item", "4.2.3"),
            ("item", "4.2.4"), ("subrubro", "4.3"), ("item", "4.3.5"),
        ]
        assert parsed["titulo"] == "EDIFICIO DE PRUEBA" and parsed["hoja"] == "01_C&P"
        assert parsed["problemas"] == ["Fila 20: código 4.2-5 dentro de 4.3 → 4.3.5.", "Fila 18: sin unidad, se tomó m²."]


# ── Regla «código del modelo» (E6) ──────────────────────────────────────────

BASES = {"codigo": "4.1.3", "nombre": "BASES AISLADAS", "unidad": "m3", "recursos": [], "parametros": []}


class TestCodeRecipe:
    def test_same_code_and_similar_description(self):
        fila = {"codigo": "4.1.3", "descripcion": "Bases aisladas", "unidad": "m³"}
        assert code_recipe(fila, {"4.1.3": BASES}) == "4.1.3"

    def test_same_code_other_work(self):
        fila = {"codigo": "4.1.3", "descripcion": "MURO DE MAMPOSTERIA HUECO DEL 12", "unidad": "m²"}
        assert code_recipe(fila, {"4.1.3": BASES}) is None

    def test_same_code_other_unit(self):
        fila = {"codigo": "4.1.3", "descripcion": "Bases aisladas", "unidad": "u"}
        assert code_recipe(fila, {"4.1.3": BASES}) is None

    def test_the_code_as_the_excel_says_it(self):
        fila = {"codigo": "4.2.3", "codigo_excel": "4.1.3", "descripcion": "BASES AISLADAS", "unidad": None}
        assert code_recipe(fila, {"4.1.3": BASES}) == "4.1.3"


@pytest.fixture
def obra_db():
    tables = obra_tables()
    tables["item_templates"].append({**BASES, "id": "tmpl-4.1.3", "org_id": ORG, "desperdicio_pct": None,
                                     "recursos": [{"tipo": "material", "codigo": "H30", "formula": "Q"}]})
    fake = FakeDB(tables)
    with patch("app.routers.obras.get_data_db", return_value=fake), \
         patch("app.routers.analysis.get_data_db", return_value=fake):
        yield fake


BASES_CLAVE = "BASES AISLADAS | m3"
OBRADOR = "OBRADOR | gl"


class TestAnalizarModelo:
    def test_model_reads_cyp_and_uses_the_formula_of_its_code(self, client, obra_db):
        res = analizar(client, wb=modelo_workbook())
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["hoja"] == "CYP" and body["planilla_simple"] is False
        bases = _tareas(body)[BASES_CLAVE]
        assert bases["receta"]["codigo"] == "4.1.3" and bases["receta"]["origen"] == "codigo"
        assert bases["receta"]["porque"] == "Tiene el código de la fórmula y se parece la descripción"
        assert bases["estado"] == "verde" and bases["avisos"] == []
        assert any("Replicar puntos" in c for c in body["correcciones_excel"])
        # The model's own prices (00_Mat) are offered, as always
        assert body["excel_con_precios"] is True

    def test_load_the_model_with_its_formula(self, client, obra_db):
        res = cargar(client, wb=modelo_workbook(), permitir=True)
        assert res.status_code == 200, res.text
        bases = next(i for i in obra_db.tables["budget_items"] if i["code"] == "4.1.3")
        assert bases["template_id"] == "tmpl-4.1.3"
        assert "mismo código" in bases["notas"]
        piso = next(i for i in obra_db.tables["budget_items"] if i["id"] == bases["parent_id"])
        assert (piso["code"], piso["description"]) == ("4.1", "FUNDACIONES")


# ── Sin fórmula y sin precio: se cargan en $0 (E5) ──────────────────────────


class TestSinPrecio:
    def test_red_until_chosen_then_loads_at_zero(self, client, obra_db):
        res = analizar(client, wb=hoja1_workbook())
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["hoja"] == "Hoja1" and body["excel_con_precios"] is False
        rojos = [t for t in body["tareas"] if t["motivo_rojo"] == "sin_receta"]
        assert body["sin_formula_sin_precio"] == len(rojos) > 0
        assert _tareas(body)[OBRADOR]["sin_precio"] is False

        # Blocked while they wait for a fórmula or for "Cargarlo sin precio ($0)"
        res = cargar(client, wb=hoja1_workbook(), permitir=True)
        assert res.status_code == 409
        assert "cargalos sin precio ($0)" in res.json()["detail"]["mensaje"]
        assert obra_db.tables["budgets"] == []

        todos = {t["clave"]: {"sin_precio": True} for t in rojos}
        body = analizar(client, todos, wb=hoja1_workbook()).json()
        assert body["sin_formula_sin_precio"] == 0
        obrador = _tareas(body)[OBRADOR]
        assert (obrador["estado"], obrador["motivo_rojo"], obrador["sin_precio"]) == ("verde", None, True)

        res = cargar(client, wb=hoja1_workbook(), permitir=True, asignaciones=todos)
        assert res.status_code == 200, res.text
        obrador = next(i for i in obra_db.tables["budget_items"] if i["description"] == "OBRADOR")
        assert (obrador["template_id"], obrador["mat_unitario"], obrador["mo_unitario"]) == (None, 0, 0)
        assert obrador["neto_total"] == 0 and obrador["cantidad"] == 1
        assert obrador["notas"].startswith(NOTA_SIN_PRECIO) and "Para confirmar" not in obrador["notas"]
        # "Sin precio" is for this obra only: nothing goes to the memory of fórmulas
        assert obra_db.tables.get("obra_recetas_memoria", []) == []

    def test_sin_precio_in_an_excel_with_prices_loads_at_zero_too(self, client, obra_db):
        from tests.test_obras_api import ARISTAS, _fix_eps
        _fix_eps(obra_db)
        res = cargar(client, asignaciones={ARISTAS: {"plantillas": [], "sin_precio": True}})
        assert res.status_code == 200, res.text
        aristas = next(i for i in obra_db.tables["budget_items"] if i["description"] == "ARISTAS DE YESO EN PAREDES")
        assert (aristas["mat_unitario"], aristas["mo_unitario"], aristas["neto_total"]) == (0, 0, 0)
        assert aristas["excel_neto"] == 150  # what the Excel said is kept, for the comparison

    def test_bad_assignment_is_rejected(self, client, obra_db):
        res = analizar(client, {OBRADOR: {"sin_precio": "si"}}, wb=hoja1_workbook())
        assert res.status_code == 400


# ── Importar Excel (E1, E2, E7, E9, E16) ────────────────────────────────────


@pytest.fixture
def excel_db():
    fake = FakeDB({"indirect_config": [], "price_catalogs": [], "catalog_entries": [],
                   "catalog_price_history": [], "budgets": [], "budget_items": [], "item_resources": []})
    with patch("app.routers.excel.get_data_db", return_value=fake):
        yield fake


def importar(client, wb, filename="obra.xlsx"):
    return client.post("/budgets/import-excel", files={"file": (filename, _bytes(wb), XLSX)})


def _nothing_created(db) -> bool:
    return all(db.tables[t] == [] for t in ("budgets", "budget_items", "item_resources", "price_catalogs",
                                             "catalog_entries"))


class TestImportExcel:
    @pytest.mark.parametrize("make", [hoja1_workbook, prueba3_workbook])
    def test_computo_without_prices_goes_to_cargar_obra(self, client, excel_db, make):
        wb = make()
        lista = wb.create_sheet("00_Mat")  # a price list too: it is not created either
        lista.append(["MATERIALES"])
        lista.append(["Codigo", "Descripcion", "Unidad", "Con IVA", "Sin IVA"])
        lista.append(["H30", "Hormigon", "m3", 121, 100])
        res = importar(client, wb)
        assert res.status_code == 422
        assert res.json()["detail"] == {
            "codigo": "CARGAR_OBRA", "motivo": "sin_precios",
            "mensaje": "Este Excel trae el cómputo sin precios. Para que la app lo calcule con las fórmulas, "
                       "cargalo en Cargar obra.",
        }
        assert _nothing_created(excel_db)

    def test_without_computo_goes_to_cargar_obra(self, client, excel_db):
        wb = openpyxl.Workbook()
        wb.active.title = "Portada"
        lista = wb.create_sheet("00_Mat")
        lista.append(["MATERIALES"])
        lista.append(["Codigo", "Descripcion", "Unidad", "Con IVA", "Sin IVA"])
        lista.append(["H30", "Hormigon", "m3", 121, 100])
        res = importar(client, wb)
        assert res.status_code == 422
        detail = res.json()["detail"]
        assert (detail["codigo"], detail["motivo"]) == ("CARGAR_OBRA", "sin_computo")
        assert "No encontré el cómputo" in detail["mensaje"] and "Cargar obra" in detail["mensaje"]
        assert _nothing_created(excel_db)

    def test_model_with_titles_in_column_a_and_detail_sheets(self, client, excel_db):
        res = importar(client, modelo_workbook(), "modelo.xlsx")
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["hoja"] == "CYP" and body["catalog_entries"] == 1
        items = {i["code"]: i for i in excel_db.tables["budget_items"]}
        secciones = [(i["code"], i["description"]) for i in excel_db.tables["budget_items"] if i["notas"] == "Seccion"]
        assert ("4.1", "FUNDACIONES") in secciones and ("2", "MOVIMIENTO DE SUELO") in secciones
        # The column titles are never a section, nor the notes a work
        assert not any("DESCRIPCI" in (d or "") or "Replicar" in (d or "") for _, d in secciones)
        assert body["items_inserted"] == len(excel_db.tables["budget_items"]) == 12
        assert items["4.1.3"]["parent_id"] == items["4.1"]["id"] and items["4.1"]["parent_id"] == items["4"]["id"]
        # Its resources come from its detail sheet (named with its code)
        recursos = [r for r in excel_db.tables["item_resources"] if r["item_id"] == items["4.1.3"]["id"]]
        assert sorted(r["codigo"] for r in recursos) == ["H30", "MO-OF"] and body["resources_inserted"] == 2
        # A work without detail and without price still goes in
        assert items["2.1"]["neto_total"] == 0 and items["1.1"]["cantidad"] == 1
        assert body["neto_excel"] == pytest.approx(43641655.99 + 739571.48)
        assert len(body["avisos"]) == 2

    def test_works_in_batches_with_their_parents(self, client, excel_db):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "01_C&P"
        filas = [(None, "1- ALBAÑILERÍA"), (None, "1.1- PLANTA BAJA")]
        filas += [(f"1.1.{n}", f"MURO {n}", "m2", 10, 100, 50, 1500, 1980) for n in range(1, 251)]
        _write(ws, filas, start=8)
        calls: list[int] = []
        original = excel_db.table

        def table(name):
            q = original(name)
            if name == "budget_items":
                insert = q.insert
                q.insert = lambda data, **k: (calls.append(len(data) if isinstance(data, list) else 1),
                                              insert(data, **k))[1]
            return q

        excel_db.table = table
        res = importar(client, wb)
        assert res.status_code == 200, res.text
        assert res.json()["items_inserted"] == 252
        assert calls == [1, 1, 100, 100, 50]
        piso = next(i for i in excel_db.tables["budget_items"] if i["code"] == "1.1")
        assert all(i["parent_id"] == piso["id"] for i in excel_db.tables["budget_items"] if i["code"].count(".") == 2)

    def test_runs_in_a_thread(self):
        from app.routers.excel import import_excel
        assert not inspect.iscoroutinefunction(import_excel)

    def test_a_failure_leaves_nothing_half_done(self, client, excel_db):
        original = excel_db.table

        def table(name):
            if name == "item_resources":
                raise RuntimeError("se cortó")
            return original(name)

        excel_db.table = table
        res = importar(client, modelo_workbook())
        assert res.status_code == 500
        assert "no quedó nada a medias" in res.json()["detail"]
        assert excel_db.tables["budgets"] == []


# ── Borrar un presupuesto (E11) ─────────────────────────────────────────────


class TestBorrar:
    @pytest.mark.parametrize("user", [MOCK_USER, LEADER], ids=["admin", "leader"])
    def test_admin_and_leader_can_delete(self, user):
        fake = FakeDB(base_tables())
        with patch("app.routers.budgets.get_data_db", return_value=fake):
            res = _client(user).delete(f"/budgets/{BUDGET}")
        assert res.status_code == 200 and res.json() == {"deleted": True}
        assert fake.tables["budgets"] == []

    def test_member_cannot(self):
        member = {**MOCK_USER, "role": "member"}
        fake = FakeDB(base_tables())
        with patch("app.routers.budgets.get_data_db", return_value=fake):
            assert _client(member).delete(f"/budgets/{BUDGET}").status_code == 403
        assert len(fake.tables["budgets"]) == 1

    def test_other_company_is_not_deleted(self):
        tables = base_tables()
        tables["budgets"][0]["org_id"] = "other-org"
        fake = FakeDB(tables)
        with patch("app.routers.budgets.get_data_db", return_value=fake):
            assert _client(LEADER).delete(f"/budgets/{BUDGET}").json() == {"deleted": False}
        assert len(fake.tables["budgets"]) == 1


# ── El modelo de Terrac y Trabajos típicos (E10) ────────────────────────────

MODELO_JSON = Path(__file__).resolve().parent.parent / "app" / "data" / "modelo_terrac.json"


class TestModeloTerrac:
    def test_the_file(self):
        data = json.loads(MODELO_JSON.read_text(encoding="utf-8"))
        rubros = data["rubros"]
        assert [r["codigo"] for r in rubros] == [str(n) for n in range(1, 10)]
        assert rubros[0]["nombre"] == "TAREAS PRELIMINARES" and rubros[8]["nombre"] == "CUBIERTA"
        assert [p["codigo"] for p in rubros[3]["pisos"]][:2] == ["4.1", "4.2"]
        assert rubros[3]["pisos"][0]["nombre"] == "FUNDACIONES"
        trabajos = [t for r in rubros for t in r["trabajos"] + [x for p in r["pisos"] for x in p["trabajos"]]]
        assert {"codigo": "4.1.3", "descripcion": "Bases aisladas", "unidad": "m3"} in trabajos
        assert all(set(t) == {"codigo", "descripcion", "unidad"} for t in trabajos)
        assert all(t["unidad"] in (None, "m2", "m3", "ml", "gl", "mes", "u") for t in trabajos)
        textos = " ".join(t["descripcion"] for t in trabajos)
        assert "Replicar" not in textos and "Dejar un item" not in textos and "REVESREVESTIMIENTO" not in textos
        assert len({t["codigo"] for t in trabajos}) == len(trabajos)

    def test_endpoint_marks_the_formulas_of_the_company(self, client):
        fake = FakeDB({"item_templates": [
            {"id": "t1", "org_id": ORG, "codigo": "4.1.3", "nombre": "BASES AISLADAS", "unidad": "m3"},
            {"id": "t2", "org_id": "other-org", "codigo": "5.1.4", "nombre": "HUECO DEL 18", "unidad": "m2"},
        ]})
        with patch("app.routers.modelo.get_data_db", return_value=fake):
            res = client.get("/modelo-terrac")
        assert res.status_code == 200, res.text
        body = res.json()
        trabajos = {t["codigo"]: t for r in body["rubros"]
                    for t in r["trabajos"] + [x for p in r["pisos"] for x in p["trabajos"]]}
        assert trabajos["4.1.3"] == {"codigo": "4.1.3", "descripcion": "Bases aisladas", "unidad": "m3",
                                     "plantilla": "4.1.3", "plantilla_nombre": "BASES AISLADAS",
                                     "plantilla_unidad": "m3"}
        assert trabajos["5.1.4"]["plantilla"] is None  # another company's fórmula does not count
        assert trabajos["1.1"]["plantilla"] is None
        assert body["resumen"]["con_formula"] == 1 and body["resumen"]["rubros"] == 9

    def test_any_user_of_the_company(self):
        fake = FakeDB({"item_templates": []})
        with patch("app.routers.modelo.get_data_db", return_value=fake):
            assert _client({**MOCK_USER, "role": "member"}).get("/modelo-terrac").status_code == 200


@pytest.fixture
def full_db():
    tables = base_tables()
    tables["item_templates"][0]["codigo"] = "4.1.1"
    fake = FakeDB(tables)
    with patch("app.routers.templates.get_data_db", return_value=fake), \
         patch("app.routers.budgets.get_data_db", return_value=fake), \
         patch("app.routers.analysis.get_data_db", return_value=fake):
        yield fake


def _full(client, items, extra_secciones=()):
    return client.post("/budgets/create-full", json={"name": "Desde el modelo", "secciones": [
        {"codigo": "4", "nombre": "ESTRUCTURA", "items": []},
        {"codigo": "4.1", "nombre": "FUNDACIONES", "items": items},
        *extra_secciones,
    ]})


def _sin_ids(rows: list[dict]) -> list[tuple]:
    keys = ("tipo", "codigo", "cantidad", "cantidad_efectiva", "precio_unitario", "subtotal", "desperdicio_pct",
            "catalog_entry_id", "formula")
    return sorted(tuple(r.get(k) for k in keys) for r in rows)


class TestCreateFullConFormula:
    def test_item_with_plantilla_is_created_with_its_formula(self, client, full_db):
        platea = {"codigo": "4.1.1", "descripcion": "PLATEA DE FUNDACIÓN", "unidad": "m2", "cantidad": 20,
                  "plantilla": "4.1.1"}
        res = _full(client, [platea], [{"codigo": "1", "nombre": "TAREAS PRELIMINARES", "items": [
            {"codigo": "1.1", "descripcion": "OBRADOR", "unidad": "mes", "cantidad": None}]}])
        assert res.status_code == 200, res.text
        body = res.json()
        assert (body["items_created"], body["con_formula"], body["sin_formula"]) == (2, 1, [])
        bid = body["budget"]["id"]
        items = {i["code"]: i for i in full_db.tables["budget_items"] if i["budget_id"] == bid}
        # 4.1 is a piso of 4
        assert items["4.1"]["parent_id"] == items["4"]["id"] and items["4"]["parent_id"] is None
        trabajo = items["4.1.1"]
        assert trabajo["template_id"] == TEMPLATE and trabajo["parent_id"] == items["4.1"]["id"]
        assert trabajo["neto_total"] > 0
        recursos = [r for r in full_db.tables["item_resources"] if r["item_id"] == trabajo["id"]]
        assert {r["codigo"]: r["precio_unitario"] for r in recursos} == {"H30": 1000, "MALLA": 50, "NYL": 5,
                                                                        "MO-OF": 200}
        # Without plantilla and without quantity: as always, at 0
        assert (items["1.1"]["cantidad"], items["1.1"]["template_id"] if "template_id" in items["1.1"] else None) \
            == (0, None)

        # The same rule as adding the work in the editor (POST /budgets/{id}/trabajos)
        res = client.post(f"/budgets/{BUDGET}/trabajos", json={"template_id": TEMPLATE, "cantidad": 20})
        assert res.status_code == 200, res.text
        otro = [r for r in full_db.tables["item_resources"] if r["item_id"] == res.json()["item"]["id"]]
        assert _sin_ids(otro) == _sin_ids(recursos)

    def test_empty_quantity(self, client, full_db):
        res = _full(client, [{"codigo": "4.1.1", "descripcion": "PLATEA", "unidad": "m2", "cantidad": None,
                              "plantilla": "4.1.1"}])
        assert res.status_code == 200, res.text
        trabajo = next(i for i in full_db.tables["budget_items"] if i.get("code") == "4.1.1")
        assert (trabajo["cantidad"], trabajo["template_id"], trabajo["neto_total"]) == (0, TEMPLATE, 0)

    def test_what_cannot_be_applied_goes_in_without_formula(self, client, full_db):
        res = _full(client, [
            {"codigo": "4.1.3", "descripcion": "Bases aisladas", "unidad": "m3", "cantidad": 5, "plantilla": "4.1.3"},
            {"codigo": "4.1.1", "descripcion": "PLATEA", "unidad": "m3", "cantidad": 5, "plantilla": "4.1.1"},
        ])
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["con_formula"] == 0 and body["items_created"] == 2
        motivos = {s["codigo"]: s["motivo"] for s in body["sin_formula"]}
        assert motivos["4.1.3"] == "La empresa no tiene una fórmula con ese código"
        assert motivos["4.1.1"].startswith("La fórmula está en m² y el trabajo en m³")
        assert all(i.get("template_id") is None for i in full_db.tables["budget_items"] if i.get("code") == "4.1.3")

    def test_a_failure_deletes_the_new_budget(self, client, full_db):
        original = full_db.table

        def table(name):
            if name == "item_resources":
                raise RuntimeError("se cortó")
            return original(name)

        full_db.table = table
        res = _full(client, [{"codigo": "4.1.1", "descripcion": "PLATEA", "unidad": "m2", "cantidad": 3,
                              "plantilla": "4.1.1"}])
        assert res.status_code == 500
        assert res.json()["detail"]["codigo"] == "NO_SE_CREO"
        assert [b["id"] for b in full_db.tables["budgets"]] == [BUDGET]


# ── Precio viejo (E13) ──────────────────────────────────────────────────────


class TestPrecioViejo:
    def test_rules(self):
        hoy = date(2026, 10, 19)
        assert precio_viejo_dias({}) == PRECIO_VIEJO_DIAS == 90
        assert precio_viejo_dias({"precio_viejo_dias": 30}) == 30
        assert precio_viejo_dias({"precio_viejo_dias": None}) == 90
        assert precio_dias({"precio_fecha": "2026-06-01"}, hoy) == 140
        assert precio_dias({"precio_fecha": None}, hoy) is None
        assert precio_dias({"precio_fecha": "no es fecha"}, hoy) is None
        assert es_precio_viejo({"precio_fecha": "2026-06-01"}, 90, hoy)
        assert not es_precio_viejo({"precio_fecha": "2026-07-21"}, 90, hoy)  # exactly 90 days: not yet
        assert not es_precio_viejo({"precio_fecha": None}, 90, hoy)

    def test_migration_013(self):
        sql = (Path(__file__).resolve().parent.parent / "migrations" / "013_precio_viejo.sql").read_text()
        code = "\n".join(line.split("--", 1)[0] for line in sql.splitlines())
        assert "ALTER TABLE indirect_config ADD COLUMN IF NOT EXISTS precio_viejo_dias integer DEFAULT 90" in code
        assert "DROP" not in code.upper() and "UPDATE " not in code.upper()

    def _indirects_db(self, rows):
        return FakeDB({"indirect_config": rows, "budgets": []})

    def test_get_and_patch_in_the_general_config(self, client):
        fake = self._indirects_db([])
        with patch("app.routers.indirects.get_data_db", return_value=fake):
            assert client.get("/indirects/general").json()["precio_viejo_dias"] == 90
            res = client.patch("/indirects/general", json={"precio_viejo_dias": 45})
            assert res.status_code == 200, res.text
            assert res.json()["precio_viejo_dias"] == 45
            assert fake.tables["indirect_config"][0]["precio_viejo_dias"] == 45
            assert client.get("/indirects/general").json()["precio_viejo_dias"] == 45
            # Together with a %: both are saved
            res = client.patch("/indirects/general", json={"precio_viejo_dias": 30, "beneficio_pct": 12})
            assert (res.json()["precio_viejo_dias"], res.json()["beneficio_pct"]) == (30, 12)
            assert client.patch("/indirects/general", json={"precio_viejo_dias": 0}).status_code == 422

    def test_without_migration_013(self, client):
        fake = self._indirects_db([{"id": "cfg", "org_id": ORG, "desperdicio_pct": 5}])
        original = fake.table

        class SinColumna:
            def __init__(self, q):
                self.q = q

            def __getattr__(self, name):
                return getattr(self.q, name)

            def update(self, data, **k):
                if "precio_viejo_dias" in data:
                    raise RuntimeError("column indirect_config.precio_viejo_dias does not exist")
                return self.q.update(data, **k)

        fake.table = lambda name: SinColumna(original(name)) if name == "indirect_config" else original(name)
        with patch("app.routers.indirects.get_data_db", return_value=fake):
            assert client.get("/indirects/general").json()["precio_viejo_dias"] == 90
            res = client.patch("/indirects/general", json={"precio_viejo_dias": 45})
        assert res.status_code == 409
        assert res.json()["detail"]["codigo"] == "FALTA_MIGRACION"
        assert "migración 013" in res.json()["detail"]["mensaje"]

    def _budget_db(self, dias=None):
        tables = base_tables()
        if dias is not None:
            tables["indirect_config"][0]["precio_viejo_dias"] = dias
        tables["item_resources"] = [
            {**RESOURCE_DEFAULTS, "id": "r1", "item_id": ITEM, "org_id": ORG, "tipo": "material", "codigo": "H30",
             "precio_unitario": 1000, "precio_fecha": "2026-06-01"},
            {**RESOURCE_DEFAULTS, "id": "r2", "item_id": ITEM, "org_id": ORG, "tipo": "material", "codigo": "MALLA",
             "precio_unitario": 50, "precio_fecha": "2026-10-01"},
            {**RESOURCE_DEFAULTS, "id": "r3", "item_id": ITEM, "org_id": ORG, "tipo": "material", "codigo": "NYL",
             "precio_unitario": 5, "precio_fecha": None},
        ]
        return FakeDB(tables)

    def test_viejos_per_item_and_days_per_resource(self, client):
        fake = self._budget_db()
        with patch("app.routers.budgets.get_data_db", return_value=fake), \
             patch("app.routers.templates.get_data_db", return_value=fake), \
             patch("app.routers.budgets.today", return_value=date(2026, 10, 19)):
            body = client.get(f"/budgets/{BUDGET}/precios-faltantes").json()
            assert (body["viejos"], body["precio_viejo_dias"]) == ({ITEM: 1}, 90)
            recursos = {r["codigo"]: r for r in client.get(f"/budgets/{BUDGET}/items/{ITEM}/resources").json()}
        assert {c: (r["precio_dias"], r["precio_viejo"]) for c, r in recursos.items()} == {
            "H30": (140, True), "MALLA": (18, False), "NYL": (None, False)}

    def test_the_company_decides(self, client):
        fake = self._budget_db(dias=10)
        with patch("app.routers.budgets.get_data_db", return_value=fake), \
             patch("app.routers.templates.get_data_db", return_value=fake), \
             patch("app.routers.budgets.today", return_value=date(2026, 10, 19)):
            body = client.get(f"/budgets/{BUDGET}/precios-faltantes").json()
        assert (body["viejos"], body["precio_viejo_dias"]) == ({ITEM: 2}, 10)


# ── Con los archivos de Emilia (si están en EXCEL_DIR) ──────────────────────

EXCEL_DIR = Path(os.environ.get("EXCEL_DIR") or Path(__file__).resolve().parent.parent / "scripts" / "excel")


def _real(name: str) -> openpyxl.Workbook:
    path = EXCEL_DIR / f"{name}.xlsx"
    if not path.exists():
        pytest.skip(f"Falta {path} (EXCEL_DIR)")
    return openpyxl.load_workbook(path, data_only=True)


def _real_bytes(name: str) -> bytes:
    path = EXCEL_DIR / f"{name}.xlsx"
    if not path.exists():
        pytest.skip(f"Falta {path} (EXCEL_DIR)")
    return path.read_bytes()


class TestArchivosDeEmilia:
    @pytest.mark.parametrize("name,hoja", [("emilia_hoja1", "Hoja1"), ("emilia_prueba3", "01_C&P")])
    def test_hoja1_and_prueba3(self, name, hoja):
        parsed = parse_any(_real(name))
        assert parsed["hoja"] == hoja
        items = [f for f in parsed["filas"] if f["nivel"] == "item"]
        assert len(items) == 244 and sum(1 for f in parsed["filas"] if f["nivel"] != "item") == 35
        codigos = {f["codigo"] for f in items}
        assert {"1.1", "1.2", "2.1", "3.1.1", "3.1.2"} <= codigos
        assert not any(plain_desc.startswith("DESCRIPCI") for plain_desc in (f["descripcion"] for f in items))
        # E12: every piso has its works
        pisos = [f for f in parsed["filas"] if f["nivel"] == "subrubro"]
        con_trabajos = {f["parent"] for f in items}
        assert all(p["orden"] in con_trabajos for p in pisos)

    def test_modelo(self):
        parsed = parse_any(_real("emilia_modelo"))
        assert parsed["hoja"] == "CYP"
        assert [f["codigo"] for f in parsed["filas"] if f["nivel"] == "rubro"] == [str(n) for n in range(1, 10)]
        assert ("4.1", "FUNDACIONES") in _filas(parsed, "subrubro")
        assert sum("es una nota" in p for p in parsed["problemas"]) == 5
        assert not any(d.startswith(("Replicar", "Rplicar", "Dejar")) for _, d in _filas(parsed, "item"))

    @pytest.mark.parametrize("name", ["emilia_hoja1", "emilia_prueba3"])
    def test_import_sends_to_cargar_obra(self, client, excel_db, name):
        res = client.post("/budgets/import-excel", files={"file": (f"{name}.xlsx", _real_bytes(name), XLSX)})
        assert res.status_code == 422 and res.json()["detail"]["codigo"] == "CARGAR_OBRA"
        assert _nothing_created(excel_db)

    def test_import_modelo(self, client, excel_db):
        res = client.post("/budgets/import-excel",
                          files={"file": ("emilia_modelo.xlsx", _real_bytes("emilia_modelo"), XLSX)})
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["hoja"] == "CYP" and body["items_inserted"] > 100 and body["resources_inserted"] > 0
        items = {i["code"]: i for i in excel_db.tables["budget_items"]}
        assert items["4.1.3"]["parent_id"] == items["4.1"]["id"]

    def test_import_ginkgo_as_always(self, client, excel_db):
        res = client.post("/budgets/import-excel", files={"file": ("ginkgo.xlsx", _real_bytes("ginkgo"), XLSX)})
        assert res.status_code == 200, res.text
        body = res.json()
        assert (body["items_inserted"], body["resources_inserted"], body["catalog_entries"]) == (279, 294, 450)
        assert body["neto_excel"] == pytest.approx(2154836047.27)

    def test_analyze_modelo_with_the_maestro(self):
        cache = EXCEL_DIR / "fake_db_cache.json"
        if not cache.exists():
            pytest.skip("Falta fake_db_cache.json (las fórmulas del Maestro)")
        from app.obra_import import excel_prices
        from app.routers import obras
        fake = FakeDB(json.loads(cache.read_text()))
        templates = obras._templates(fake, ORG)
        wb = _real("emilia_modelo")
        body = obras.analyze(parse_any(wb), templates, obras.PriceBook(fake, ORG, date(2026, 10, 19)), {}, {},
                             excel_prices(wb))
        bases = next(t for t in body["tareas"] if t["codigos"] == ["4.1.3"])
        assert (bases["receta"]["codigo"], bases["receta"]["origen"]) == ("4.1.3", "codigo")


def test_modelo_json_is_a_copy_per_request():
    """The model is read once; a request never changes it for the next one."""
    from app.routers.modelo import load_modelo, modelo_con_formulas
    antes = copy.deepcopy(load_modelo())
    modelo_con_formulas({"4.1.3": {"codigo": "4.1.3", "nombre": "BASES", "unidad": "m3"}})
    assert load_modelo() == antes
