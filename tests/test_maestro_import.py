"""Tests for the TERRAC Maestro price-list importer."""

from __future__ import annotations

from datetime import datetime

import openpyxl
import pytest

from app.maestro_import import parse_price, parse_sheet, parse_workbook, report_markdown

D = datetime(2026, 6, 3)


class TestParsePrice:
    def test_numbers(self):
        assert parse_price(1377) == 1377.0
        assert parse_price(76446.28) == 76446.28

    @pytest.mark.parametrize("text,value", [
        ("$  26.000,00", 26000.0),
        ("$  1.300.000,00", 1300000.0),
        ("$200.000,00", 200000.0),
        ("1500", 1500.0),
    ])
    def test_argentine_strings(self, text, value):
        assert parse_price(text) == value

    def test_empty(self):
        assert parse_price(None) is None
        assert parse_price("  ") is None

    def test_garbage(self):
        with pytest.raises(ValueError):
            parse_price("a consultar")


# 00_Mat-like layout: unlabeled fecha/proveedor that move between sections
MAT_ROWS = [
    (None,) * 10,
    (None,) * 10,
    ("CODIGO", "MATERIALES CORRALON", "UNIDAD", "PRECIO CON IVA", "PRECIO SIN IVA",
     "DESCUENTO", "PRECIO CON DESCUENTO", None, None, None),
    ("PIEG", "Piedra a granel", "m3", 92500.0, 76446.28, None, None, None, "Su Corralon", D),
    ("pie ", "Bolson de Piedra", "m3", 60000.0, 61700.0, None, None, None, None, None),
    ("PIEG", "Piedra otra vez", "m3", None, 80000.0, None, None, None, None, None),
    (None,) * 10,
    (None, None, "u", None, None, None, None, None, None, None),  # content but no code
    ("CODIGO", "MATERIALES AISLACION", "UNIDAD", "PRECIO CON IVA", "PRECIO SIN IVA",
     None, None, None, None, None),
    ("A-MEG", "Membrana", "u", None, "$  26.000,00", "ML", D, None, None, None),
    ("A-SIK", "Sika", "u", None, 1000.0, "Su Corralon", D, None, None, None),
    ("PI-ASF", "Pintura asfaltica", "u", None, 500.0, None, "rinde 0,5lts/m2", None, None, None),
    ("CODIGO", "MATERIALES PINTURA", "UNIDAD", None, "PRECIO CON IVA",
     None, "RENDIMIENTO", None, None, None),
    ("LAT", "Latex", "u", None, 105181.82, None, "12M2 X LITRO", "Quimtex", D, None),
    ("CODIGO", "MATERIALES PLOMERIA", "UNIDAD", "PRECIO CON IVA", "PRECIO SIN IVA",
     None, None, None, None, None),
    ("INOD", "Inodoro", "u", None, 600000.0, None, D, "https://mercadolibre.com/x", None, None),
    ("LP8", "Ladrillo", "u", None, None, None, None, None, None, None),
]

# 00_Eq / 00_Sub-like layout: labeled columns
EQ_ROWS = [
    ("EQUIPOS",),
    (),
    ("CODIGO", "MATERIALES CORRALON", "UNIDAD", "PRECIO SIN IVA", "REFERENCIA", "FECHA", None),
    ("E-MC", "Mini Cargadora", "u", 420000.0, "JP VIAL", datetime(2025, 10, 15), "CON MAQUINISTA"),
    ("E-MP", "Motopison", "u", 40000.0, None, None, None),
    ("E-X", "Raro", "u", 1.0, None, "sin dato", None),
]


class TestParseSheetMat:
    @pytest.fixture
    def parsed(self):
        return parse_sheet("00_Mat", MAT_ROWS, "material")

    def by_code(self, entries, code):
        return [e for e in entries if e["codigo"] == code]

    def test_skips_headers_blanks_and_rows_without_code(self, parsed):
        entries, issues = parsed
        assert len(entries) == 9
        assert [i["fila"] for i in issues["sin_codigo"]] == [8]

    def test_corralon_fecha_and_proveedor(self, parsed):
        e = self.by_code(parsed[0], "PIEG")[0]
        assert e["precio_sin_iva"] == 76446.28
        assert e["precio_con_iva"] == 92500.0
        assert e["fecha_precio"] == "2026-06-03"
        assert e["proveedor"] == "Su Corralon"

    def test_codes_normalized_and_reported(self, parsed):
        entries, issues = parsed
        assert self.by_code(entries, "PIE")
        assert issues["codigo_normalizado"][0]["original"] == "pie"

    def test_aislacion_columns_and_string_price(self, parsed):
        e = self.by_code(parsed[0], "A-MEG")[0]
        assert e["precio_sin_iva"] == 26000.0
        assert e["proveedor"] == "ML"
        assert e["fecha_precio"] == "2026-06-03"

    def test_note_in_date_column_is_not_proveedor(self, parsed):
        e = self.by_code(parsed[0], "PI-ASF")[0]
        assert e["proveedor"] is None
        assert e["fecha_precio"] is None

    def test_pintura_only_con_iva_is_not_guessed(self, parsed):
        entries, issues = parsed
        e = self.by_code(entries, "LAT")[0]
        assert e["precio_sin_iva"] is None
        assert e["precio_con_iva"] == 105181.82
        assert e["proveedor"] == "Quimtex"
        assert e["fecha_precio"] == "2026-06-03"
        assert [i["codigo"] for i in issues["solo_con_iva"]] == ["LAT"]

    def test_url_is_not_proveedor(self, parsed):
        e = self.by_code(parsed[0], "INOD")[0]
        assert e["proveedor"] is None
        assert e["fecha_precio"] == "2026-06-03"

    def test_missing_price_and_date_are_reported_not_invented(self, parsed):
        entries, issues = parsed
        e = self.by_code(entries, "LP8")[0]
        assert e["precio_sin_iva"] is None and e["fecha_precio"] is None
        assert "LP8" in [i["codigo"] for i in issues["sin_precio"]]
        sin_fecha = {i["codigo"] for i in issues["sin_fecha"]}
        assert {"PIE", "PI-ASF", "LP8"} <= sin_fecha
        assert "PIEG" in sin_fecha  # the second PIEG row has no date

    def test_duplicates(self, parsed):
        dups = parsed[1]["duplicados"]
        assert len(dups) == 1
        assert dups[0]["codigo"] == "PIEG"
        assert [f["fila"] for f in dups[0]["filas"]] == [4, 6]


class TestParseSheetLabeled:
    def test_labeled_columns(self):
        entries, issues = parse_sheet("00_Eq", EQ_ROWS, "equipo")
        assert len(entries) == 3
        mc = entries[0]
        assert mc["precio_sin_iva"] == 420000.0
        assert mc["proveedor"] == "JP VIAL"
        assert mc["fecha_precio"] == "2025-10-15"
        assert [i["codigo"] for i in issues["sin_fecha"]] == ["E-MP", "E-X"]
        assert issues["fecha_ilegible"][0]["valor"] == "sin dato"
        assert "duplicados" not in issues


def _workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in (("00_Mat", MAT_ROWS), ("00_Eq", EQ_ROWS)):
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(list(row))
    sub = wb.create_sheet("00_Sub")
    sub.append(["CODIGO", "PINTURA", "UNIDAD", "PRECIO SIN IVA", "REFERENCIA", "FECHA"])
    sub.append(["E-MC", "Mini Cargadora", "u", 300000, "PABLO", datetime(2026, 1, 1)])
    return wb


class TestParseWorkbook:
    def test_sheets_missing_and_cross_duplicates(self):
        entries, report = parse_workbook(_workbook())
        assert set(entries) == {"material", "equipo", "subcontrato"}
        assert report["faltantes"] == ["00_MO"]
        cross = {d["codigo"]: d["hojas"] for d in report["duplicados_entre_hojas"]}
        assert cross == {"E-MC": [("00_Eq", 420000.0), ("00_Sub", 300000.0)]}

    def test_report_markdown(self):
        _, report = parse_workbook(_workbook())
        md = report_markdown(report, "maestro.xlsx")
        assert "# Informe de importación del Maestro de precios" in md
        assert "| 00_Mat | 9 |" in md
        assert "## Códigos duplicados en la misma hoja" in md
        assert "PIEG" in md
        assert "## Precios sin fecha" in md
        assert "## Solo precio con IVA" in md
        assert "E-MC" in md and "## Códigos repetidos entre hojas" in md
        assert "$76.446,28" in md
