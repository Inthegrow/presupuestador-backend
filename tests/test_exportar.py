"""Exportar sin sorpresas (PLAN_EXPORTAR 2): the Planilla Terrac and the planilla simple.

The Planilla Terrac is Sol's Excel: 01_C&P with her 26 columns, one sheet per work with
resources and the Coeficiente de pase. It goes back through Cargar obra (parse_obra) with
the same works, rubros, pisos and costs, and its amounts are the saved ones to the cent.
"""

from __future__ import annotations

import io
import os
import re
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.calculations import (
    calc_budget_summary,
    calc_cascade_indirects,
    calc_item_from_resources,
    calc_resource_subtotal,
    cascade_factors,
    is_section,
)
from app.main import create_app
from app.obra_import import excel_prices, parse_obra
from app.terrac_export import (
    C_EQUIPOS, C_JORNALES, C_L, C_M, C_MAT, C_MO, C_MO_MAT, C_N, C_O, C_P, C_Q, C_R, C_S,
    C_SUBC, C_T, C_U, C_UNIT, C_V, C_W, C_X, C_Y, C_Z, cp_rows, sheet_name,
)
from tests.test_recipes_api import MOCK_USER, ORG, FakeDB

BUDGET = "00000000-0000-0000-0000-0000000000e1"
OTHER = "00000000-0000-0000-0000-0000000000e2"
CFG = {"beneficio_pct": 12.5}  # the budget's own %; the rest are the defaults
SIMPLE_HEADERS = [
    "Codigo", "Descripcion", "Unidad", "Cantidad", "MAT Unitario", "MO Unitario", "MAT Total",
    "MO Total", "Directo Total", "Indirecto Total", "Beneficio Total", "Impuestos Total",
    "Precio sin IVA", "IVA", "Precio con IVA", "Notas",
]


def _section(id_, code, desc, parent=None, order=0):
    return {"id": id_, "budget_id": BUDGET, "org_id": ORG, "code": code, "description": desc,
            "parent_id": parent, "sort_order": order, "notas": "Seccion", "cantidad": 0}


def _res(id_, item, tipo, codigo, cantidad=None, precio=0.0, **over):
    r = {"id": id_, "item_id": item, "org_id": ORG, "tipo": tipo, "codigo": codigo,
         "descripcion": f"Recurso {codigo}", "unidad": "u", "cantidad": cantidad,
         "desperdicio_pct": 0, "precio_unitario": precio, "lo_compra_cliente": False, **over}
    return calc_resource_subtotal(r)


def _work(id_, code, desc, parent, order, cantidad, unidad="m2", resources=(), mat=0.0, mo=0.0):
    """A work priced like the app does: from its resources (or a price by hand), then the cascade."""
    item = {"id": id_, "budget_id": BUDGET, "org_id": ORG, "code": code, "description": desc,
            "parent_id": parent, "sort_order": order, "notas": None, "unidad": unidad,
            "cantidad": cantidad}
    if resources:
        calc_item_from_resources(item, list(resources))
    else:
        item.update(mat_unitario=mat, mo_unitario=mo, mat_total=round(mat * cantidad, 2),
                    mo_total=round(mo * cantidad, 2))
        item["directo_total"] = round(item["mat_total"] + item["mo_total"], 2)
    return calc_cascade_indirects(item, CFG)


RES_BASES = [
    _res("x1", "w3", "material", "H30", 35, 120000.0, unidad="m3", desperdicio_pct=10),
    _res("x2", "w3", "material", "HADN12", 120, 9876.54, unidad="kg", desperdicio_pct=15),
    _res("x3", "w3", "material", "NYL", 50, 800.0, lo_compra_cliente=True),
    _res("x4", "w3", "mano_obra", "MO-OF", trabajadores=3, dias=7, cargas_sociales_pct=25,
         precio=55000.0, unidad="jornal"),
    _res("x5", "w3", "equipo", "MS-ER", 2, 41000.0, unidad="dia"),
    _res("x6", "w3", "mo_material", "AL-50k", 2, 33333.33, desperdicio_pct=10),
    _res("x7", "w3", "subcontrato", "BOMBA", 1, 450000.0, unidad="gl"),
]
RES_CONTRAPISO = [
    _res("y1", "w4", "material", "CEM", 37.5, 9150.0, unidad="bolsa", desperdicio_pct=5),
    _res("y2", "w4", "mano_obra", "MO-AY", trabajadores=2, dias=1.5, cargas_sociales_pct=25,
         precio=48000.0, unidad="jornal"),
]
RES_MURO = [
    _res("z1", "w5", "material", "LH18", 7.33 * 13, 1234.56, desperdicio_pct=5),
    _res("z2", "w5", "mano_obra", "MO-OF", trabajadores=2, dias=3, precio=55000.0),
]

ITEMS = [
    _section("r1", "1", "TAREAS PRELIMINARES", order=1),
    # Price by hand (no resources): E and F, no sheet of its own
    _work("w1", "1.1", "OBRADOR", "r1", 2, 1, unidad="gl", mo=2150000.0),
    _work("w2", "1.2", "CERCO  DE OBRA", "r1", 3, 1, unidad="gl", mat=3000000.0, mo=870000.0),
    _section("r2", "3", "ESTRUCTURA DE HORMIGÓN ARMADO", order=4),
    _section("p1", "3.1", "FUNDACIONES", parent="r2", order=5),
    _work("w3", "3.1.1", "BASES AISLADAS", "p1", 6, 35, unidad="u", resources=RES_BASES),
    # A piso without code (as some imports leave it): it takes 3.2 from its works
    _section("p2", "", "PLANTA BAJA", parent="r2", order=7),
    _work("w4", "3.2.1", "CONTRAPISO e=8cm", "p2", 8, 12.5, resources=RES_CONTRAPISO),
    _section("r3", "4", "ALBAÑILERÍA", order=9),
    _work("w5", "4.1", "MURO DE LADRILLO HUECO DEL 18", "r3", 10, 7.33, resources=RES_MURO),
]
RESOURCES = RES_BASES + RES_CONTRAPISO + RES_MURO
WORKS = [i for i in ITEMS if not is_section(i)]


def tables(items=ITEMS, resources=RESOURCES):
    return {
        "budgets": [
            {"id": BUDGET, "org_id": ORG, "name": "Edificio Ginkgo", "indirectos": dict(CFG)},
            {"id": OTHER, "org_id": "otra-empresa", "name": "Ajeno", "indirectos": {}},
        ],
        "budget_items": items + [{**ITEMS[1], "id": "ajeno", "budget_id": OTHER, "org_id": "otra-empresa"}],
        "item_resources": resources,
        "indirect_config": [{"id": "cfg", "org_id": ORG}],
    }


@pytest.fixture
def client():
    app = create_app()
    from app.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: MOCK_USER
    return TestClient(app)


@pytest.fixture
def db():
    fake = FakeDB(tables())
    with patch("app.routers.excel.get_data_db", return_value=fake):
        yield fake


def terrac(client, budget=BUDGET):
    r = client.get(f"/budgets/{budget}/export/excel", params={"formato": "terrac"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return openpyxl.load_workbook(io.BytesIO(r.content))


def cents(value) -> int:
    return round(float(value or 0) * 100)


def work_rows(ws) -> dict[str, int]:
    """Row of each work in 01_C&P, by code (column A)."""
    return {ws.cell(r, 1).value: r for r in range(8, ws.max_row + 1)
            if ws.cell(r, 1).value and ws.cell(r, 4).value is not None}


# ── Ida y vuelta ────────────────────────────────────────────────────────────


class TestRoundTrip:
    def test_parse_obra_reads_the_same_budget(self, client, db):
        wb = terrac(client)
        parsed = parse_obra(wb)
        assert parsed["problemas"] == []
        assert parsed["titulo"] == "Edificio Ginkgo"
        filas = parsed["filas"]

        sections = [(f["nivel"], f["codigo"], f["descripcion"]) for f in filas if f["nivel"] != "item"]
        assert sections == [
            ("rubro", "1", "TAREAS PRELIMINARES"),
            ("rubro", "3", "ESTRUCTURA DE HORMIGÓN ARMADO"),
            ("subrubro", "3.1", "FUNDACIONES"),
            ("subrubro", "3.2", "PLANTA BAJA"),
            ("rubro", "4", "ALBAÑILERÍA"),
        ]

        items = [f for f in filas if f["nivel"] == "item"]
        assert [(f["codigo"], f["descripcion"], f["unidad"], f["cantidad"]) for f in items] == [
            (w["code"], " ".join(w["description"].split()), w["unidad"], w["cantidad"]) for w in WORKS
        ]
        by_orden = {f["orden"]: f for f in filas}
        by_id = {i["id"]: i for i in ITEMS}
        for f, w in zip(items, WORKS):
            # Same parent (rubro or piso), by name
            assert by_orden[f["parent"]]["descripcion"] == by_id[w["parent_id"]]["description"]
            ex = f["excel"]
            assert cents(ex["directo"]) == cents(w["directo_total"]), w["code"]
            assert cents(ex["neto"]) == cents(w["neto_total"]), w["code"]
            assert cents(ex["mat_unit"]) == cents(w["mat_unitario"]), w["code"]
            assert cents(ex["mo_unit"]) == cents(w["mo_unitario"]), w["code"]

    def test_values_not_formulas(self, client, db):
        wb = terrac(client)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value, str) and cell.value.startswith("=")), (ws.title, cell.coordinate)

    def test_detail_sheets_give_their_prices_back(self, client, db):
        # Cargar obra reads the work sheets as Sol's detail sheets: their prices come back
        prices = excel_prices(terrac(client))
        assert prices["H30"]["precio"] == 120000.0
        assert prices["H30"]["tipo"] == "material"
        assert prices["H30"]["trabajo"] == "BASES AISLADAS"
        assert prices["MO-OF"]["tipo"] == "mano_obra"
        assert prices["MS-ER"]["tipo"] == "equipo"
        assert prices["BOMBA"]["tipo"] == "subcontrato"


# ── 01_C&P ──────────────────────────────────────────────────────────────────


class TestPlanilla:
    def test_header_like_sols(self, client, db):
        ws = terrac(client)["01_C&P"]
        assert ws["A1"].value == "Edificio Ginkgo"
        assert ws["A3"].value == "PLANILLA COMPUTO Y PRESUPUESTO"
        assert [ws[c].value for c in ("A5", "B5", "C5", "D5", "E5", "O5", "R5", "U5")] == [
            "ITEM", "DESCRIPCIÓN", "UNIDAD", "CANTIDAD", "SUBTOTAL 01: GASTOS DIRECTOS",
            "SUBTOTAL 02: GASTOS INDIRECTOS", "SUBTOTAL 03: BENEFICIO E IMPUESTOS", "TOTAL (NETO)"]
        assert [ws.cell(7, c).value for c in range(5, 12)] == [
            "MAT.", "M.O.: JORNALES", "M.O.: EQUIPOS", "M.O.: MATERIALES", "M.O.: SUBCONTRATOS", "M.O.", "GENERAL"]
        assert ws.cell(8, 2).value == "1- TAREAS PRELIMINARES"
        assert "E5:N5" in {str(r) for r in ws.merged_cells.ranges}

    def test_columns_add_up(self, client, db):
        ws = terrac(client)["01_C&P"]
        rows = work_rows(ws)
        assert set(rows) == {w["code"] for w in WORKS}
        for w in WORKS:
            v = {c: ws.cell(rows[w["code"]], c).value for c in range(5, 27)}
            assert cents(v[C_MO]) == sum(cents(v[c]) for c in (C_JORNALES, C_EQUIPOS, C_MO_MAT, C_SUBC))
            assert cents(v[C_UNIT]) == cents(v[C_MAT]) + cents(v[C_MO])
            assert cents(v[C_N]) == cents(v[C_L]) + cents(v[C_M]) == cents(w["directo_total"])
            assert cents(v[C_Q]) == cents(v[C_O]) + cents(v[C_P]) == cents(w["indirecto_total"])
            assert cents(v[C_T]) == cents(v[C_R]) + cents(v[C_S]) == cents(
                w["beneficio_total"] + w["impuestos_total"])
            assert cents(v[C_Z]) == cents(v[C_X]) + cents(v[C_Y]) == cents(w["neto_total"])
            assert cents(v[C_W]) == cents(v[C_U]) + cents(v[C_V]) == cents(w["neto_total"] / w["cantidad"])

    def test_unit_columns_by_kind_of_resource(self, client, db):
        ws = terrac(client)["01_C&P"]
        bases = {c: ws.cell(work_rows(ws)["3.1.1"], c).value for c in range(5, 12)}
        w3 = next(w for w in WORKS if w["code"] == "3.1.1")
        assert bases[C_MAT] == w3["mat_unitario"]
        sub = {t: sum(r["subtotal"] for r in RES_BASES if r["tipo"] == t)
               for t in ("mano_obra", "equipo", "mo_material", "subcontrato")}
        for col, tipo in ((C_JORNALES, "mano_obra"), (C_EQUIPOS, "equipo"), (C_MO_MAT, "mo_material"),
                          (C_SUBC, "subcontrato")):
            assert bases[col] == pytest.approx(sub[tipo] / 35, abs=0.011), tipo

    def test_price_by_hand_goes_to_e_and_f(self, client, db):
        ws = terrac(client)["01_C&P"]
        rows = work_rows(ws)
        obrador = [ws.cell(rows["1.1"], c).value for c in range(5, 11)]
        assert obrador == [0, 2150000.0, 0, 0, 0, 2150000.0]
        cerco = [ws.cell(rows["1.2"], c).value for c in range(5, 11)]
        assert cerco == [3000000.0, 870000.0, 0, 0, 0, 870000.0]

    def test_totals_row_is_the_budget_summary(self, client, db):
        ws = terrac(client)["01_C&P"]
        total = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 1).value == "TOTAL DEL PRESUPUESTO")
        summary = calc_budget_summary(ITEMS)
        assert cents(ws.cell(total, C_N).value) == cents(summary["directo_total"])
        assert cents(ws.cell(total, C_Q).value) == cents(summary["indirecto_total"])
        assert cents(ws.cell(total, C_T).value) == cents(summary["beneficio_total"] + summary["impuestos_total"])
        assert cents(ws.cell(total, C_Z).value) == cents(summary["neto_total"])
        assert cents(ws.cell(total, C_L).value) + cents(ws.cell(total, C_M).value) == cents(summary["directo_total"])
        assert cents(ws.cell(total, C_X).value) + cents(ws.cell(total, C_Y).value) == cents(summary["neto_total"])

    def test_coeficiente_de_pase(self, client, db):
        ws = terrac(client)["Coeficiente de pase"]
        values = {ws.cell(r, 1).value: (ws.cell(r, 2).value, ws.cell(r, 4).value) for r in range(1, ws.max_row + 1)}
        summary = calc_budget_summary(ITEMS)
        assert values["Beneficio"] == (0.125, pytest.approx(summary["beneficio_total"]))
        assert values["Gastos de estructura"][0] == 0.15
        assert values["Impuesto al cheque"][0] == 0.012
        assert values["IVA"][0] == 0.21
        assert values["SUBTOTAL 01 - COSTOS DIRECTOS"][1] == summary["directo_total"]
        assert values["TOTAL (NETO) SIN IVA"][1] == summary["neto_total"]
        assert values["TOTAL CON IVA"][1] == summary["total_final"]
        coef = cascade_factors({**CFG})["coeficiente"]
        assert values["Coeficiente de pase (sin IVA)"][1] == coef
        texto = next(k for k in values if isinstance(k, str) and k.startswith("Por cada $100"))
        assert f"${100 * coef:.2f}".replace(".", ",") in texto


# ── Una hoja por trabajo ────────────────────────────────────────────────────


class TestWorkSheets:
    def test_one_sheet_per_work_with_resources(self, client, db):
        wb = terrac(client)
        assert wb.sheetnames == ["01_C&P", "3.1.1", "3.2.1", "4.1", "Coeficiente de pase"]

    def test_sheet_total_is_the_direct_cost(self, client, db):
        wb = terrac(client)
        for w in WORKS:
            if w["code"] not in wb.sheetnames:
                continue
            ws = wb[w["code"]]
            assert (ws["A3"].value, ws["B3"].value, ws["H3"].value, ws["I3"].value) == (
                w["code"], w["description"], w["cantidad"], w["unidad"])
            total = next(r for r in range(1, ws.max_row + 1)
                         if ws.cell(r, 1).value == "TOTAL DEL TRABAJO (COSTO DIRECTO)")
            assert cents(ws.cell(total, 9).value) == cents(w["directo_total"])

    def test_blocks_by_kind_and_what_the_client_buys(self, client, db):
        ws = terrac(client)["3.1.1"]
        col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        for titulo in ("MATERIALES", "MANO DE OBRA - PERSONAS", "MANO DE OBRA - EQUIPOS",
                       "MANO DE OBRA - MATERIALES", "MANO DE OBRA - SUBCONTRATOS"):
            assert titulo in col_a
        nyl = col_a.index("NYL") + 1
        assert ws.cell(nyl, 10).value == "Lo compra el cliente"
        assert ws.cell(nyl, 9).value == 0
        # Materials block total: what the resources add up to (the client's costs 0)
        total_mat = col_a.index("TOTAL MATERIALES") + 1
        assert cents(ws.cell(total_mat, 9).value) == sum(
            cents(r["subtotal"]) for r in RES_BASES if r["tipo"] == "material")
        mo = col_a.index("MO-OF") + 1
        assert [ws.cell(mo, c).value for c in (4, 5, 6)] == [3, 7, 0.25]  # trabajadores, días, cargas


class TestSheetNames:
    def test_valid_and_unique(self):
        used = {"01_c&p", "coeficiente de pase"}
        codes = ["2.1", "2.1", "2.1", "1/2", "a[b]:c*d?e\\f", "X" * 40, "X" * 40, "", "'3.4'", "01_C&P"]
        names = [sheet_name(c, used, n) for n, c in enumerate(codes, start=1)]
        assert names[:3] == ["2.1", "2.1 (2)", "2.1 (3)"]
        assert names[3] == "1_2"
        assert names[7] == "Trabajo 8"
        assert names[8] == "3.4"
        assert names[9] == "01_C&P (2)"
        assert len({n.lower() for n in names}) == len(names)
        for n in names:
            assert 0 < len(n) <= 31
            assert not re.search(r"[\[\]:*?/\\]", n)
            assert not n.startswith("'") and not n.endswith("'")

    def test_repeated_long_and_slash_codes_in_the_file(self, client, db):
        long_code = "9.99.999.9999.99999.999999.9999999"
        codes = {"a": "2.1", "b": "2.1", "c": "2/3", "d": long_code, "e": ""}
        resources = [_res(f"q-{i}", i, "material", "C", 1, 10.0) for i in codes]
        items = [_section("s1", "2", "CONTRAPISOS", order=1)] + [
            _work(i, code, f"TRABAJO {i}", "s1", n, 10, resources=[r for r in resources if r["item_id"] == i])
            for n, (i, code) in enumerate(codes.items(), start=2)]
        db.tables.update(tables(items, resources))
        wb = terrac(client)
        names = wb.sheetnames[1:-1]
        assert names == ["2.1", "2.1 (2)", "2_3", long_code[:31], "Trabajo 5"]
        # Saved and opened again (Excel would refuse an invalid or repeated name)
        out = io.BytesIO()
        wb.save(out)
        assert openpyxl.load_workbook(io.BytesIO(out.getvalue())).sheetnames == wb.sheetnames


# ── Rubros y pisos ──────────────────────────────────────────────────────────


class TestTree:
    def test_loose_works_and_deep_sections_keep_their_parent(self, client, db):
        items = [
            _section("s1", "", "PRELIMINARES", order=1),
            _work("a", "1.1", "OBRADOR", "s1", 2, 1, mo=100.0),
            _section("s2", "x", "ESTRUCTURA", order=3),
            _section("s3", "", "PB", parent="s2", order=4),
            _section("s4", "", "SECTOR A", parent="s3", order=5),  # a third level: written as a piso
            _work("b", "7.1", "COLUMNAS", "s4", 6, 2, mat=10.0),
            _work("c", "7.2", "VIGAS", "s3", 7, 3, mat=20.0),  # after a deeper title: goes before it
            _work("d", "9", "LIMPIEZA FINAL", None, 8, 1, mo=50.0),  # loose, at the end
        ]
        rows = cp_rows(items)
        assert [(r["kind"], r["item"]["id"], r.get("code")) for r in rows] == [
            ("work", "d", None),
            ("rubro", "s1", "1"),
            ("work", "a", None),
            ("rubro", "s2", "7"),
            ("piso", "s3", "7.1"),
            ("work", "c", None),
            ("piso", "s4", "7.2"),
            ("work", "b", None),
        ]
        db.tables.update(tables(items, []))
        filas = parse_obra(terrac(client))["filas"]
        parent = {f["descripcion"]: (filas[f["parent"]]["descripcion"] if f["parent"] is not None else None)
                  for f in filas}
        assert parent["LIMPIEZA FINAL"] is None
        assert parent["OBRADOR"] == "PRELIMINARES"
        assert parent["VIGAS"] == "PB"
        assert parent["COLUMNAS"] == "SECTOR A"
        assert parent["SECTOR A"] == "ESTRUCTURA"


# ── Planilla simple y permisos ──────────────────────────────────────────────


class TestSimple:
    @pytest.mark.parametrize("params", [{}, {"formato": "simple"}])
    def test_same_headers_in_the_first_row(self, client, db, params):
        r = client.get(f"/budgets/{BUDGET}/export/excel", params=params)
        assert r.status_code == 200, r.text
        df = pd.read_excel(io.BytesIO(r.content))
        assert list(df.columns) == SIMPLE_HEADERS
        total = df[df["Codigo"] == "TOTAL"].iloc[0]
        assert cents(total["Precio sin IVA"]) == cents(calc_budget_summary(ITEMS)["neto_total"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["Presupuesto"]
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.b
        assert "#,##0.00" in ws["M2"].number_format

    def test_unknown_formato(self, client, db):
        assert client.get(f"/budgets/{BUDGET}/export/excel", params={"formato": "pdf"}).status_code == 422


class TestPermissions:
    @pytest.mark.parametrize("formato", ["simple", "terrac"])
    def test_other_org_is_404(self, client, db, formato):
        r = client.get(f"/budgets/{OTHER}/export/excel", params={"formato": formato})
        assert r.status_code == 404

    def test_terrac_without_items_is_404(self, client, db):
        db.tables["budget_items"] = []
        r = client.get(f"/budgets/{BUDGET}/export/excel", params={"formato": "terrac"})
        assert r.status_code == 404

    def test_file_name_with_any_obra_name(self, client, db):
        db.tables["budgets"][0]["name"] = "Torre Ñandú — etapa 2"
        r = client.get(f"/budgets/{BUDGET}/export/excel", params={"formato": "terrac"})
        assert r.status_code == 200
        assert "Planilla Terrac" in r.headers["content-disposition"]
