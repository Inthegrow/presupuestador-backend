"""API tests for "Importar Excel": the price sheets update the list of the same file."""

from __future__ import annotations

import io
import os
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.budget_prices import today
from app.main import create_app
from tests.test_recipes_api import MOCK_USER, ORG, FakeDB


def workbook(mat=None, mo=None) -> bytes:
    """Excel with 00_Mat (data from row 3) and 00_MO (data from row 4), plus one other sheet."""
    wb = openpyxl.Workbook()
    wb.active.title = "Portada"
    wb.active.append(["Obra de prueba"])
    if mat is not None:
        ws = wb.create_sheet("00_Mat")
        ws.append(["MATERIALES"])
        ws.append(["Codigo", "Descripcion", "Unidad", "Con IVA", "Sin IVA"])
        for codigo, desc, precio in mat:
            ws.append([codigo, desc, "u", precio * 1.21 if precio else precio, precio])
    if mo is not None:
        ws = wb.create_sheet("00_MO")
        ws.append(["MANO DE OBRA"])
        ws.append(["-"])
        ws.append(["Codigo", "Descripcion", "Unidad", "Precio", "Referencia"])
        for codigo, desc, precio in mo:
            ws.append([codigo, desc, "dia", precio, "UOCRA"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
def client():
    app = create_app()
    from app.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: MOCK_USER
    return TestClient(app)


@pytest.fixture
def db():
    fake = FakeDB({"indirect_config": [], "price_catalogs": [], "catalog_entries": [],
                   "catalog_price_history": [], "budgets": []})
    with patch("app.routers.excel.get_data_db", return_value=fake):
        yield fake


def upload(client, content: bytes, filename: str = "Las Heras.xlsx"):
    files = {"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post("/budgets/import-excel", files=files)
    assert r.status_code == 200, r.text
    return r.json()


def entry(db, codigo):
    rows = [e for e in db.tables["catalog_entries"] if e["codigo"] == codigo]
    assert len(rows) == 1, rows
    return rows[0]


def history(db, codigo):
    eid = entry(db, codigo)["id"]
    return [(h["precio_sin_iva"], h["fecha_precio"]) for h in db.tables["catalog_price_history"]
            if h["entry_id"] == eid]


V1_MAT = [("MAT-1", "Cemento", 100), ("MAT-2", "Arena", 50), ("MAT-4", "Cal", 80)]
V1_MO = [("MO-1", "Oficial", 200)]


class TestImportExcelPriceList:
    def test_first_import_creates_list(self, client, db):
        data = upload(client, workbook(V1_MAT, V1_MO))
        assert len(db.tables["price_catalogs"]) == 1
        catalog = db.tables["price_catalogs"][0]
        assert catalog["name"] == "Catalogo - Las Heras.xlsx" and not catalog.get("oficial")
        assert data["catalog_id"] == catalog["id"]
        assert data["catalog_entries"] == 4
        assert data["catalog_reused"] is False
        assert data["catalog_name"] == "Catalogo - Las Heras.xlsx"
        assert (data["precios_actualizados"], data["precios_nuevos"]) == (0, 4)
        # Each new price starts its history
        assert history(db, "MAT-1") == [(100, None)]
        # The response keeps what it already had
        assert data["budget_name"] == "Las Heras" and data["budget_id"]

    def test_same_file_twice_updates_the_same_list(self, client, db):
        first = upload(client, workbook(V1_MAT, V1_MO))
        # Codes match without caring about case or spaces
        mat2 = [("mat-1", "Cemento", 120), ("MAT-4", "Cal", 0), ("MAT-3", "Piedra", 30)]
        second = upload(client, workbook(mat2, V1_MO))

        assert len(db.tables["price_catalogs"]) == 1
        assert second["catalog_id"] == first["catalog_id"]
        assert second["catalog_reused"] is True
        assert second["catalog_name"] == "Catalogo - Las Heras.xlsx"
        assert (second["precios_actualizados"], second["precios_nuevos"]) == (1, 1)
        # Changed price: the new one is today's, the old one stays in the history
        assert entry(db, "MAT-1")["precio_sin_iva"] == 120
        assert entry(db, "MAT-1")["fecha_precio"] == today().isoformat()
        assert history(db, "MAT-1") == [(100, None), (120, today().isoformat())]
        # Same price: untouched
        assert entry(db, "MO-1").get("fecha_precio") is None and history(db, "MO-1") == [(200, None)]
        # A new code is added, one the Excel no longer has is kept
        assert entry(db, "MAT-3")["precio_sin_iva"] == 30
        assert entry(db, "MAT-2")["precio_sin_iva"] == 50
        # A 0 in the Excel is "sin precio": it does not replace the price
        assert entry(db, "MAT-4")["precio_sin_iva"] == 80 and history(db, "MAT-4") == [(80, None)]
        assert len(db.tables["catalog_entries"]) == 5

    def test_old_list_without_history_keeps_old_price(self, client, db):
        # Lists imported before this change have no history rows
        db.tables["price_catalogs"].append({"id": "cat-old", "org_id": ORG, "name": "Catalogo viejo",
                                            "source_file": "Las Heras.xlsx", "created_at": "2026-01-01"})
        db.tables["catalog_entries"].append({"id": "e-old", "catalog_id": "cat-old", "org_id": ORG,
                                             "tipo": "material", "codigo": "MAT-1", "precio_sin_iva": 100})
        data = upload(client, workbook([("MAT-1", "Cemento", 120)]))
        assert (data["catalog_id"], data["catalog_name"]) == ("cat-old", "Catalogo viejo")
        assert history(db, "MAT-1") == [(100, None), (120, today().isoformat())]

    def test_same_code_in_two_sheets_matches_by_tipo(self, client, db):
        upload(client, workbook([("X-1", "Material X", 10)], [("X-1", "Mano de obra X", 20)]))
        upload(client, workbook([("X-1", "Material X", 10)], [("X-1", "Mano de obra X", 25)]))
        rows = [(e["tipo"], e["precio_sin_iva"]) for e in db.tables["catalog_entries"]]
        assert sorted(rows) == [("mano_obra", 25), ("material", 10)]

    def test_code_that_moves_to_another_sheet_is_a_new_entry(self, client, db):
        # Material X-1 first; then the same file has X-1 only as mano de obra
        upload(client, workbook([("X-1", "Material X", 100)]))
        data = upload(client, workbook([], [("X-1", "Mano de obra X", 500)]))
        assert (data["precios_actualizados"], data["precios_nuevos"]) == (0, 1)
        rows = sorted((e["tipo"], e["precio_sin_iva"]) for e in db.tables["catalog_entries"])
        assert rows == [("mano_obra", 500), ("material", 100)]

    def test_code_that_moves_back_keeps_both_entries(self, client, db):
        # The inverse: mano de obra X-1 first, then X-1 only as material
        upload(client, workbook(None, [("X-1", "Mano de obra X", 500)]))
        data = upload(client, workbook([("X-1", "Material X", 100)]))
        assert (data["precios_actualizados"], data["precios_nuevos"]) == (0, 1)
        rows = sorted((e["tipo"], e["precio_sin_iva"]) for e in db.tables["catalog_entries"])
        assert rows == [("mano_obra", 500), ("material", 100)]

    def test_without_price_sheets_no_list(self, client, db):
        data = upload(client, workbook())
        assert db.tables["price_catalogs"] == [] and db.tables["catalog_entries"] == []
        assert data["catalog_id"] is None and data["catalog_name"] is None
        assert (data["catalog_reused"], data["precios_actualizados"], data["precios_nuevos"]) == (False, 0, 0)
        assert len(db.tables["budgets"]) == 1

    def test_other_file_other_list(self, client, db):
        first = upload(client, workbook(V1_MAT))
        second = upload(client, workbook(V1_MAT), filename="Lugones.xlsx")
        assert len(db.tables["price_catalogs"]) == 2
        assert second["catalog_id"] != first["catalog_id"] and second["catalog_reused"] is False
        assert len(db.tables["catalog_entries"]) == 6

    def test_list_of_another_org_is_not_reused(self, client, db):
        db.tables["price_catalogs"].append({"id": "cat-x", "org_id": "other-org", "name": "Ajena",
                                            "source_file": "Las Heras.xlsx"})
        data = upload(client, workbook(V1_MAT))
        assert data["catalog_reused"] is False and data["catalog_id"] != "cat-x"
