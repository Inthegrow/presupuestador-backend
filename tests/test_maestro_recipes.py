"""Tests for the Maestro recipes importer (solapas -> plantillas, CYP -> árbol)."""

from __future__ import annotations

import os
import re
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.maestro_recipes import (
    parse_maestro,
    report_markdown,
    slug,
    template_payload,
    template_status,
)
from app.recipes import expand_resource, param_defaults, validate_template
from import_recetas import TREE_NAME, apply_to_db
from tests.test_recipes_api import MOCK_USER, ORG, FakeDB

VLOOKUP = '=IF($A{r}=0,"",VLOOKUP($A{r},\'00_Mat\'!$A$4:$J$1175,2,FALSE))'

CATALOG = {
    "material": [
        {"codigo": c, "descripcion": d, "unidad": u}
        for c, d, u in [
            ("H30", "Hormigon H30", "m3"), ("HADN8", "Hierro 8", "u"), ("LP18", "Ladrillo 18", "u"),
            ("HADN6", "Hierro 6", "u"), ("CEM", "Cemento", "u"), ("RP-ZOC", "Zocalo", "u"),
            ("D-FIJ", "Fijacion", "u"), ("D-FIJ", "Fijacion otra", "u"), ("H-PL-M", "Bombeo", "m3"),
        ]
    ],
    "mano_obra": [{"codigo": c, "descripcion": c, "unidad": "jornal"} for c in ("MO-CA", "MO-OF", "MO-AY")],
    "equipo": [{"codigo": "E-MCM", "descripcion": "Martillo", "unidad": "dia"}],
    "subcontrato": [{"codigo": "SUB-CP", "descripcion": "Contrapiso", "unidad": "m3"}],
}


def _sections(ws, start: int) -> dict[str, int]:
    """Write the 5 section headers; return first data row of each."""
    rows = {}
    row = start
    for name in ["MATERIALES", "MANO DE OBRA - PERSONAS", "MANO DE OBRA - EQUIPOS",
                 "MANO DE OBRA - MATERIALES", "MANO DE OBRA - SUBCONTRATOS"]:
        ws[f"A{row}"] = name
        ws[f"A{row + 1}"] = "Código"
        rows[name] = row + 2
        ws[f"A{row + 9}"] = "TOTAL"
        row += 12
    return rows


def _res(ws, row, code, d=None, e=None, f=None):
    ws[f"A{row}"] = code
    ws[f"B{row}"] = VLOOKUP.format(r=row)
    if d is not None:
        ws[f"D{row}"] = d
    if e is not None:
        ws[f"E{row}"] = e
    if f is not None:
        ws[f"F{row}"] = f
    ws[f"G{row}"] = f"=ROUNDUP((D{row}+(D{row}*F{row})),0)"
    ws[f"H{row}"] = VLOOKUP.format(r=row)


def build_workbooks():
    wb = openpyxl.Workbook()
    wv = openpyxl.Workbook()  # cached values (what data_only=True returns)
    wb.remove(wb.active)
    wv.remove(wv.active)

    def sheet(name):
        return wb.create_sheet(name), wv.create_sheet(name)

    # 4.1.1: title in row 3, parameters, fixed numbers, fixed labor days
    ws, vs = sheet("4.1.1")
    ws["A1"] = "='01_C&P_E1'!A1"
    ws["A3"], ws["B3"], ws["H3"], ws["I3"] = "=CYP!A10", "PLATEA", 194, "m²"
    ws["M3"], ws["M4"] = "ML VIGA", "=(5*13)+(4*14)"
    vs["M4"] = 121
    ws["K4"], ws["K5"], ws["L5"] = "ALTURA (M)", 2.8, "=H3/K5"
    s = _sections(ws, 5)
    m = s["MATERIALES"]
    _res(ws, m, "H30", "=H3*0.2", f=0.1)
    _res(ws, m + 1, "H30", "=(0.2*0.7)*M4", f=0.1)
    _res(ws, m + 2, "HADN8", "=(107+530)*0.9", f=0.1)
    _res(ws, m + 3, "hadn6", "=((L5*4)/12)*2", f=0.1)
    _res(ws, m + 4, "H-PL-m", f"=D{m}", f=0.1)
    _res(ws, m + 5, "XX-NUEVO", "=H3", f=0.1)
    _res(ws, m + 6, "HADN8")  # no quantity: skipped
    p = s["MANO DE OBRA - PERSONAS"]
    _res(ws, p, "MO-OF", 3, 20, 0.25)
    _res(ws, p + 1, "mo-ay", 1, "=H3/15", 0.25)
    _res(ws, p + 2, "MO-CA")  # empty: skipped
    _res(ws, s["MANO DE OBRA - EQUIPOS"], "E-MCM", 1, "=H3/10", 0.15)
    _res(ws, s["MANO DE OBRA - SUBCONTRATOS"], "SUB-CP", "=H3", 1, 0)

    # 5.1.2: the example quantity typed by hand (=30/K5) and a price reference (H8)
    ws, vs = sheet("5.1.2")
    ws["A3"], ws["B3"], ws["H3"], ws["I3"] = "x", "LADRILLO", 30, "m2"
    ws["K4"], ws["K5"], ws["L5"] = "ALTURA (M)", 2.8, "=30/K5"
    s = _sections(ws, 5)
    m = s["MATERIALES"]
    _res(ws, m, "HADN6", "=((L5*4)/12)*2", f=0.1)
    _res(ws, m + 1, "CEM", f"=H{m}/16", f=0.15)
    vs[f"D{m + 1}"] = 3480.125

    # 7.2 premarcos: title in row 1, formulas copied with H$3 (empty)
    ws, vs = sheet("7.2")
    ws["B1"], ws["H1"], ws["I1"] = "PREMARCOS VENTANAS", 1, "Unid"
    s = _sections(ws, 3)
    p = s["MANO DE OBRA - PERSONAS"]
    _res(ws, p, "MO-CA", 1, "=H$3/30", 0.25)
    _res(ws, p + 1, "MO-OF", 1, "=H1/2.5", 0.25)

    # AGREGADO = zócalo (goes to 7.1.4); 7.1.4. is not loaded
    for name in ("AGREGADO", "7.1.4."):
        ws, vs = sheet(name)
        ws["B3"], ws["H3"], ws["I3"] = "ZOCALO", 1, "ml"
        s = _sections(ws, 5)
        if name == "AGREGADO":
            _res(ws, s["MATERIALES"], "rp-zoc", "=H3", f=0.1)
            _res(ws, s["MANO DE OBRA - PERSONAS"], "MO-OF", 1, "=H3/50", 0.25)

    # 5.6.1 recuadros: removed
    ws, vs = sheet("5.6.1")
    ws["B3"], ws["H3"], ws["I3"] = "RECUADROS", 1, "ml"
    _sections(ws, 5)

    # CYP
    ws, _ = sheet("CYP")
    rows = [
        ("1- TAREAS PRELIMINARES", None, None), ("1.1", "OBRADOR", "mes"),
        ("4 - ESTRUCTURA", None, None), ("4 -1 FUNDACIONES", None, None),
        ("4.1.1", "='4.1.1'!B3", "='4.1.1'!I3"),
        ("4-3 SOBRE SUBSUELO", None, None), (None, "Replicar puntos", None),
        ("5- ALBAÑILERIA", None, None), ("5-1 MAMPOSTERIA", None, None),
        ("5.1.2", "LADRILLO CERAMICO HUECO DEL 12", "='5.1.2'!I3"),
        ("5-6 RECUADROS", None, None), ("5.6.1", "RECUADROS", None),
        ("7 TERMINACIONES", None, None), ("7-1 REVESTIMIENTO ", None, None),
        ("7.1.4", "REVESTIMIENTO HORIZONTAL ZOCALO", None),
        ("7-2 VENTANAS", None, None), ("7-3 PUERTAS", None, None),
        ("7-5 OTRAS TERMINACIONES", None, None), (None, "Dejar un item libre", None),
        ("9- CUBIERTA", None, None), (None, "Dejar un item libre para cubierta", None),
    ]
    for i, (a, b, c) in enumerate(rows, start=8):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = a, b, c

    # A sheet that is not a recipe
    wb.create_sheet("00_Mat")
    return wb, wv


@pytest.fixture(scope="module")
def parsed():
    wb, wv = build_workbooks()
    return parse_maestro(wb, wv, CATALOG)


def tmpl(parsed, code):
    return next(t for t in parsed["plantillas"] if t["codigo"] == code)


def res(t, codigo, n=0):
    return [r for r in t["recursos"] if r["codigo"] == codigo][n]


class TestTranslation:
    def test_q_and_coefficients(self, parsed):
        t = tmpl(parsed, "4.1.1")
        h30 = res(t, "H30")
        assert h30["formula"] == "Q*0.2"
        assert h30["desperdicio_pct"] == 10
        assert h30["redondear"] is True and not h30["revisar"]
        assert t["unidad"] == "m2" and t["cantidad_ejemplo"] == 194

    def test_auxiliary_cell_becomes_parameter(self, parsed):
        t = tmpl(parsed, "4.1.1")
        assert res(t, "H30", 1)["formula"] == "(0.2*0.7)*ml_viga"
        params = {p["clave"]: p["valor"] for p in t["parametros"]}
        assert params["ml_viga"] == 121
        assert params["altura_m"] == 2.8

    def test_cell_that_depends_on_q_is_inlined(self, parsed):
        t = tmpl(parsed, "4.1.1")
        assert res(t, "HADN6")["formula"] == "(((Q/altura_m)*4)/12)*2"
        assert "l5" not in {p["clave"] for p in t["parametros"]}

    def test_reference_to_other_resource(self, parsed):
        assert res(tmpl(parsed, "4.1.1"), "H-PL-M")["formula"] == "(Q*0.2)"

    def test_fixed_number_is_loaded_and_marked(self, parsed):
        r = res(tmpl(parsed, "4.1.1"), "HADN8")
        assert r["formula"] == "(107+530)*0.9"
        assert r["revisar"] is True
        assert "Número fijo" in " ".join(r["notas"])

    def test_code_normalized_and_unknown_code(self, parsed):
        t = tmpl(parsed, "4.1.1")
        assert res(t, "HADN6")["descripcion"] == "Hierro 6"
        nuevo = res(t, "XX-NUEVO")
        assert nuevo["revisar"] and "no está en los catálogos" in " ".join(nuevo["notas"])

    def test_empty_rows_skipped(self, parsed):
        omitidos = {o["codigo"] for o in tmpl(parsed, "4.1.1")["omitidos"]}
        assert {"HADN8", "MO-CA"} <= omitidos

    def test_equipment_and_subcontract(self, parsed):
        t = tmpl(parsed, "4.1.1")
        assert res(t, "E-MCM")["formula"] == "Q/10"
        assert res(t, "SUB-CP")["formula"] == "Q"
        assert res(t, "SUB-CP")["tipo"] == "subcontrato"
        assert "redondear" not in res(t, "SUB-CP")


class TestLabor:
    def test_rate_from_formula(self, parsed):
        ay = res(tmpl(parsed, "4.1.1"), "MO-AY")
        assert ay["rendimiento"] == "15" and ay["trabajadores"] == 1
        assert ay["cargas_sociales_pct"] == 25
        assert not ay["revisar"]

    def test_fixed_days_become_rate(self, parsed):
        of = res(tmpl(parsed, "4.1.1"), "MO-OF")
        assert of["rendimiento"] == "9.7"  # 194 m2 in 20 days
        assert of["trabajadores"] == 3 and of["revisar"]
        row = expand_resource(of, 194, {})
        assert row["dias"] == pytest.approx(20)


class TestExcelErrors:
    def test_example_quantity_typed_by_hand(self, parsed):
        r = res(tmpl(parsed, "5.1.2"), "HADN6")
        assert r["formula"] == "(((Q/altura_m)*4)/12)*2"
        assert r["revisar"] and "cantidad de ejemplo" in " ".join(r["notas"])

    def test_price_reference_falls_back_to_value(self, parsed):
        r = res(tmpl(parsed, "5.1.2"), "CEM")
        assert r["formula"] == "3480.125"
        assert r["revisar"] and "No se pudo traducir" in " ".join(r["notas"])

    def test_cyp_name_wins(self, parsed):
        t = tmpl(parsed, "5.1.2")
        assert t["nombre"] == "LADRILLO CERAMICO HUECO DEL 12"
        assert t["nombre_solapa"] == "LADRILLO"


class TestDecisions:
    def test_premarcos_labor_always_material_optional(self, parsed):
        t = tmpl(parsed, "7.2.1")
        assert {p["clave"]: p["valor"] for p in t["parametros"]}["con_material"] == 0
        material = next(r for r in t["recursos"] if r["tipo"] == "material")
        assert material["formula"] == "Q*con_material"
        ca = res(t, "MO-CA")
        assert ca["rendimiento"] == "30" and ca["revisar"]  # H$3 was empty: taken as Q
        assert res(t, "MO-OF")["rendimiento"] == "2.5"
        # Without material by default; with con_material = 1 it is Q units
        assert expand_resource(material, 4, param_defaults(t["parametros"]))["cantidad"] == 0
        assert expand_resource(material, 4, {"con_material": 1})["cantidad"] == 4

    def test_zocalo_goes_to_714(self, parsed):
        t = tmpl(parsed, "7.1.4")
        assert t["hoja"] == "AGREGADO"
        assert res(t, "RP-ZOC")["formula"] == "Q"
        assert "7.1.4." in parsed["hojas_omitidas"]

    def test_recuadros_removed(self, parsed):
        assert "5.6.1" not in {t["codigo"] for t in parsed["plantillas"]}
        codes = {n["codigo"] for n in parsed["arbol"]}
        assert "5.6" not in codes and "5.6.1" not in codes

    def test_free_items(self, parsed):
        nodes = {n["codigo"]: n for n in parsed["arbol"]}
        assert nodes["7.5.1"]["libre"] and nodes["9.1"]["libre"]
        assert not nodes["9.1"].get("plantilla")


class TestTree:
    def test_structure_and_links(self, parsed):
        nodes = {n["codigo"]: n for n in parsed["arbol"]}
        assert nodes["4"]["nivel"] == "rubro" and nodes["4.1"]["nivel"] == "subrubro"
        assert nodes["4.1.1"]["parent"] == "4.1" and nodes["4.1.1"]["plantilla"] == "4.1.1"
        assert nodes["4.1.1"]["nombre"] == "PLATEA"  # from the solapa (CYP had a formula)
        assert nodes["1.1"]["parent"] == "1" and not nodes["1.1"].get("plantilla")
        assert nodes["7.2.1"]["parent"] == "7.2" and nodes["7.2.1"]["plantilla"] == "7.2.1"
        assert nodes["7.3.1"]["parent"] == "7.3"

    def test_per_floor_rubros_not_loaded(self, parsed):
        assert "4.3" not in {n["codigo"] for n in parsed["arbol"]}

    def test_order(self, parsed):
        orden = [n["orden"] for n in parsed["arbol"]]
        assert orden == sorted(orden)
        codes = [n["codigo"] for n in parsed["arbol"]]
        assert codes.index("4") < codes.index("4.1") < codes.index("4.1.1") < codes.index("5")


class TestOutput:
    def test_all_templates_valid_for_the_engine(self, parsed):
        for t in parsed["plantillas"]:
            payload = template_payload(t)
            assert validate_template(payload["recursos"], payload["parametros"]) == [], t["codigo"]
            for r in payload["recursos"]:
                for q in (0, 1, 50):
                    expand_resource(r, q, param_defaults(payload["parametros"]))

    def test_payload(self, parsed):
        p = template_payload(tmpl(parsed, "4.1.1"))
        assert p["codigo"] == "4.1.1" and p["origen"] == "maestro_terrac"
        assert p["categoria"] == "Estructura"
        assert all(isinstance(r["nota"], str) for r in p["recursos"])

    def test_status(self, parsed):
        assert template_status(tmpl(parsed, "4.1.1")) == "revisar"

    def test_report(self, parsed):
        md = report_markdown(parsed, "maestro.xlsx")
        assert "## Ítem por ítem" in md
        assert "4.1.1 PLATEA" in md and "⬜ **9.1" in md
        assert "Revisar material **HADN8**" in md

    def test_slug(self):
        assert slug("ALTURA (M)") == "altura_m"
        assert slug("12 cm") == "p_12_cm"
        assert slug("Ñandú año") == "nandu_ano"


class TestApply:
    def test_creates_templates_and_tree(self, parsed):
        db = FakeDB({})
        out = apply_to_db(db, ORG, parsed, "maestro.xlsx")
        assert out["plantillas_creadas"] == len(parsed["plantillas"])
        nodes = db.tables["standard_tree_nodes"]
        assert len(nodes) == len(parsed["arbol"])
        by_code = {n["codigo"]: n for n in nodes}
        tid = {t["codigo"]: t["id"] for t in db.tables["item_templates"]}
        assert by_code["4.1.1"]["template_id"] == tid["4.1.1"]
        assert by_code["4.1.1"]["parent_id"] == by_code["4.1"]["id"]
        assert by_code["9.1"]["template_id"] is None and by_code["9.1"]["libre"]
        assert db.tables["standard_trees"][0]["nombre"] == TREE_NAME

    def test_rerun_updates_in_place(self, parsed):
        db = FakeDB({})
        apply_to_db(db, ORG, parsed)
        ids = sorted(t["id"] for t in db.tables["item_templates"])
        out = apply_to_db(db, ORG, parsed)
        assert out["plantillas_creadas"] == 0
        assert sorted(t["id"] for t in db.tables["item_templates"]) == ids
        assert len(db.tables["standard_trees"]) == 1
        assert len(db.tables["standard_tree_nodes"]) == len(parsed["arbol"])


class TestStandardTreeAPI:
    @pytest.fixture
    def client_db(self, parsed):
        db = FakeDB({})
        apply_to_db(db, ORG, parsed)
        db.tables["standard_trees"].append({"id": "otro", "org_id": "otra-org", "nombre": "X"})
        app = create_app()
        from app.auth import get_current_user
        app.dependency_overrides[get_current_user] = lambda: MOCK_USER
        with patch("app.routers.standard_trees.get_data_db", return_value=db):
            yield TestClient(app), db

    def test_list_only_own_org(self, client_db):
        client, _ = client_db
        r = client.get("/standard-trees")
        assert r.status_code == 200
        assert [t["nombre"] for t in r.json()] == [TREE_NAME]

    def test_nested(self, client_db):
        client, db = client_db
        tree_id = db.tables["standard_trees"][0]["id"]
        body = client.get(f"/standard-trees/{tree_id}").json()
        rubros = {n["codigo"]: n for n in body["nodos"]}
        assert "4" in rubros and "4.1.1" not in rubros
        sub = rubros["4"]["children"][0]
        assert sub["codigo"] == "4.1" and sub["children"][0]["codigo"] == "4.1.1"

    def test_other_org_is_404(self, client_db):
        client, _ = client_db
        assert client.get("/standard-trees/otro").status_code == 404


class TestMigration006:
    SQL = Path(__file__).resolve().parent.parent / "migrations" / "006_maestro_recipes.sql"

    def _sql(self) -> str:
        return "\n".join(line.split("--", 1)[0] for line in self.SQL.read_text(encoding="utf-8").splitlines())

    def test_idempotent_and_no_data_loss(self):
        sql = self._sql()
        for line in re.findall(r"ADD COLUMN[^\n]*|CREATE (?:UNIQUE )?INDEX[^\n]*|CREATE TABLE[^\n]*", sql):
            assert "IF NOT EXISTS" in line, line
        assert not re.search(r"DROP\s+(TABLE|COLUMN)|DELETE\s+FROM|TRUNCATE", sql, re.IGNORECASE)

    def test_backend_only(self):
        sql = self._sql()
        for table in ("standard_trees", "standard_tree_nodes"):
            assert f"ALTER TABLE {table} ENABLE ROW LEVEL SECURITY" in sql
            assert f"REVOKE ALL ON {table} FROM anon, authenticated" in sql
        assert not re.search(r"CREATE POLICY|USING\s*\(\s*true\s*\)", sql, re.IGNORECASE)

    def test_columns_used_by_the_script(self):
        sql = self._sql()
        for col in ("codigo", "origen"):
            assert f"ALTER TABLE item_templates ADD COLUMN IF NOT EXISTS {col}" in sql
        for col in ("tree_id", "parent_id", "codigo", "nombre", "unidad", "nivel", "orden",
                    "template_id", "libre"):
            assert re.search(rf"^\s+{col}\s", sql, re.MULTILINE), col


class TestSqlScript:
    def test_script_for_the_sql_editor(self, parsed):
        from import_recetas import sql_script

        sql = sql_script(parsed, "maestro.xlsx")
        assert sql.splitlines()[7].strip().startswith("v_org  uuid := NULL;")  # "linea 8"
        assert "ON CONFLICT (org_id, codigo) WHERE codigo IS NOT NULL DO UPDATE" in sql
        assert "DELETE FROM standard_tree_nodes WHERE tree_id = v_tree" in sql
        assert sql.count("$maestro$") == 2
        assert f"plantillas = {len(parsed['plantillas'])}" in sql
        for t in parsed["plantillas"]:
            assert f"('{t['codigo']}', " in sql

    def test_quotes_escaped(self):
        from import_recetas import _lit

        assert _lit("d'agua") == "'d''agua'"
        assert _lit(None) == "NULL" and _lit(True) == "true" and _lit(2.5) == "2.5"
