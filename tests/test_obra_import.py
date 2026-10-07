"""Tests for the Fase 5 loader: cómputo de una obra (01_C&P) + recetas del Maestro."""

from __future__ import annotations

from datetime import datetime

import openpyxl
import pytest

from app.formulas import evaluate
from app.obra_import import (
    altura_from,
    build_plan,
    espesor_m_from,
    expand_item,
    item_notes,
    match_recipe,
    parse_obra,
    report_markdown,
    rule_for,
    suggest_recipes,
)
from import_obra import sql_script

TEMPLATES = {
    "5.1.4": {
        "codigo": "5.1.4", "unidad": "m2",
        "parametros": [{"clave": "altura_m", "valor": 2.8}],
        "recursos": [
            {"tipo": "material", "codigo": "LH18", "formula": "Q*16", "desperdicio_pct": 15.0},
            {"tipo": "material", "codigo": "HADN6", "formula": "((Q/altura_m)*4)/12"},
            {"tipo": "mano_obra", "codigo": "MO-OF", "trabajadores": 1.0, "rendimiento": "15",
             "cargas_sociales_pct": 25.0},
        ],
    },
    "4.1.7": {
        "codigo": "4.1.7", "unidad": "m3", "parametros": [],
        "recursos": [
            {"tipo": "material", "codigo": "H30", "formula": "Q", "desperdicio_pct": 10.0},
            {"tipo": "mano_obra", "codigo": "MO-OF", "trabajadores": 1.0, "rendimiento": "2"},
        ],
    },
    "8.3": {"codigo": "8.3", "unidad": "m2", "parametros": [],
            "recursos": [{"tipo": "material", "codigo": "EPS-500", "formula": "Q"}]},
    "5.2.3": {"codigo": "5.2.3", "unidad": "m3", "parametros": [],
              "recursos": [{"tipo": "material", "codigo": "ES", "formula": "Q*0.8"}]},
    "7.1.1": {"codigo": "7.1.1", "unidad": "m2", "parametros": [],
              "recursos": [{"tipo": "material", "codigo": "RP-PORC", "formula": "Q"},
                           {"tipo": "material", "codigo": "RP-KP", "formula": "(Q*8)/25", "revisar": True,
                            "nota": "número fijo"}]},
}


def _workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_C&P"
    ws["A1"] = "EDIFICIO DE PRUEBA"
    rows = [
        # A (código), B, C, D (cantidad), E (mat unit), J (MO unit), N (directo), Z (neto)
        (None, "1- TAREAS PRELIMINARES", None, None, None, None, None, None),
        (datetime(2026, 1, 1), "OBRADOR", "gl", 1, 0, 1000, 1000, 1500),
        (None, "3- ESTRUCTURA", None, None, None, None, None, None),
        (None, "3.1- FUNDACIONES", None, None, None, None, None, None),
        ("3.1-3", "TENSORES 20 cm x 40 cm.", "ml", 100, 0, 0, 0, 0),
        (None, "4. ALBAÑILERIA", None, None, None, None, None, None),
        (None, "4.2- PRIMER PISO", None, None, None, None, None, None),
        ("4.2-1", "MURO DE MAMPOSTERIA EN LADRILLO HUECO DEL 18. h 3m", "m²", 30, 100, 200, 9000, 12000),
        ("4.2-2", "ARISTAS DE YESO EN PAREDES", "m", 10, 5, 5, 100, 150),
        ("4.2-3", "TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm", "m²", 50, 1, 1, 100, 100),
        ("4.2-4", "COLOCACION DE REVESTIMIENTOS EN PISOS PORCELANATO. NO INCLUYE EL PORCELANATO", None, 20,
         1, 1, 40, 50),
        (None, "4.3- SEGUNDO PISO", None, None, None, None, None, None),
        ("4.2-5", "COLOCACION DE REVESTIMIENTOS EN PISOS PORCELANATO. NO INCLUYE EL PORCELANATO", "m²", 20,
         1, 1, 40, 50),
    ]
    for r, (a, b, c, d, e, j, n, z) in enumerate(rows, start=8):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
        ws.cell(r, 3, c)
        ws.cell(r, 4, d)
        ws.cell(r, 5, e)
        ws.cell(r, 10, j)
        ws.cell(r, 14, n)
        ws.cell(r, 26, z)
    return wb


@pytest.fixture()
def parsed() -> dict:
    return parse_obra(_workbook())


@pytest.fixture()
def plan(parsed: dict) -> dict:
    return build_plan(parsed, TEMPLATES)


def _item(plan: dict, codigo: str) -> dict:
    return next(i for i in plan["items"] if i["codigo"] == codigo)


def test_parse_tree(parsed: dict) -> None:
    filas = parsed["filas"]
    by_code = {f["codigo"]: f for f in filas}
    assert [f["nivel"] for f in filas[:2]] == ["rubro", "item"]
    # Código guardado como fecha: 01/01/2026 → 1.1
    assert by_code["1.1"]["descripcion"] == "OBRADOR"
    assert by_code["3.1.3"]["parent"] == by_code["3.1"]["orden"]
    assert by_code["3.1"]["parent"] == by_code["3"]["orden"]
    assert by_code["4.2.1"]["excel"]["mat_unit"] == 100


def test_code_outside_its_floor_is_fixed(parsed: dict) -> None:
    by_code = {f["codigo"]: f for f in parsed["filas"]}
    assert "4.2.5" not in by_code
    assert by_code["4.3.5"]["parent"] == by_code["4.3"]["orden"]
    assert any("4.2-5" in p for p in parsed["problemas"])


def test_missing_unit_taken_from_other_floor(parsed: dict) -> None:
    by_code = {f["codigo"]: f for f in parsed["filas"]}
    assert by_code["4.2.4"]["unidad"] == "m²"


def test_match_recipe_and_height() -> None:
    assert match_recipe("Muro de mampostería en ladrillo HUECO DEL 18. h 3m")["plantillas"] == [("5.1.4", 1.0)]
    assert match_recipe("MURO DE CARGA. MAMPOSTERIA EN LADRILLO HUECO DEL 18")["nota"]
    assert match_recipe("ESMALTE EN BARANDAS DE BALCONES") is None
    assert altura_from("MURO ... h 4,4m") == 4.4
    assert altura_from("MURO DE CARGA h 0.2m/1.8m") is None


# Revisión de Ginkgo: las reglas que cambian con las correcciones usan la fórmula de antes mientras la
# corrección no se aplicó (la fórmula nueva no existe en la empresa)
_HOY = {c: {"codigo": c, "unidad": "m2"} for c in ("5.5.5", "6.1", "6.5")}
_CORREGIDO = {**_HOY, "5.5.6": {"codigo": "5.5.6", "unidad": "m2"}, "6.11": {"codigo": "6.11", "unidad": "m2"}}


@pytest.mark.parametrize("descripcion, hoy, corregido", [
    ("YESO PROYECTADO EN PAREDES INTERIORES", "5.5.5", "5.5.6"),
    ("Cielorrasos de yeso suspendido en baños y cocinas", "6.5", "6.1"),
    ("CIELORRASOS APLICADOS EN YESO EN ESTAR Y DORMITORIOS", "6.1", "6.11"),
])
def test_mapeo_con_alternativa(descripcion: str, hoy: str, corregido: str) -> None:
    assert match_recipe(descripcion, _HOY)["plantillas"] == [(hoy, 1.0)]
    assert match_recipe(descripcion, _CORREGIDO)["plantillas"] == [(corregido, 1.0)]
    # Without the org's recipes: the rule as written
    assert match_recipe(descripcion)["plantillas"] == [(corregido, 1.0)]
    fila = {"descripcion": descripcion, "unidad": "m²"}
    assert rule_for(fila, _HOY)["plantillas"] == [(hoy, 1.0)]
    assert rule_for(fila, _CORREGIDO)["plantillas"] == [(corregido, 1.0)]
    assert rule_for(fila, _HOY)["falta_factor"] == [] and "alternativa" not in rule_for(fila, _HOY)


def test_mapeo_alternativa_trae_su_nota() -> None:
    texto = "CIELORRASOS DE YESO SUSPENDIDO"
    assert "placa verde" in match_recipe(texto, _HOY)["nota"]
    assert "yeso armado" in match_recipe(texto, _CORREGIDO)["nota"]
    assert "nota" not in match_recipe("YESO PROYECTADO", _HOY)


def test_mapeo_alternativa_build_plan_sin_formulas_faltantes() -> None:
    parsed = {"filas": [{"nivel": "item", "codigo": "5.1", "descripcion": "YESO PROYECTADO", "unidad": "m2",
                         "cantidad": 10, "excel": {}}]}
    plan = build_plan(parsed, _HOY)
    assert plan["plantillas_faltantes"] == [] and plan["items"][0]["plantillas"] == ["5.5.5"]


# Revisión de Ginkgo, grupo B: revoque con silleta (5.5.7) y contrapiso de hormigón celular (5.2.4)
_P524 = {"codigo": "5.2.4", "unidad": "m2", "parametros": [{"clave": "espesor_m", "valor": 0.1}],
         "recursos": [{"tipo": "material", "codigo": "CEM", "formula": "Q*espesor_m*15"},
                      {"tipo": "subcontrato", "codigo": "SUB-CONT", "formula": "Q"}]}
_SIN_B = {**TEMPLATES, **{c: {"codigo": c, "unidad": "m2", "parametros": [], "recursos": []}
                         for c in ("5.5.3", "5.4.1")}}
_CON_B = {**_SIN_B, "5.5.7": {"codigo": "5.5.7", "unidad": "m2", "parametros": [], "recursos": []}, "5.2.4": _P524}
_CELULAR = [{"texto": "BOMBEO HORMIGON CELULAR. INCLUYE MANO DE OBRA Y BOMBA.", "hoja": "4.1-6"}]
_CASCOTE = [{"texto": "CASCOTE", "hoja": "4.2-7"}]
_GINKGO_CONTRAPISOS = [
    ("CONTRAPISO e=10cm", ["5.2.4"], 0.10),
    ("CONTRAPISO EN HALL + RAMPAS. ESP.=10cm.", ["5.2.4"], 0.10),
    ("CONTRAPISO EN INTERIOR e: 10cm", ["5.2.4"], 0.10),
    ("CONTRAPISO/ CARPETA e=10cm", ["5.2.4"], 0.10),
    ("TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm", ["8.3", "5.2.4"], 0.08),
    ("TELGOPOR 50 mm + CONTRAPISO/CARPETA EN AZOTEA INACCESIBLE (BALCONES) e=4cm", ["8.3", "5.2.4"], 0.04),
    ("TELGOPOR 50 mm + CONTRAPISO/ CARPETA EN AZOTEA ACCESIBLE e: 8cm", ["8.3", "5.2.4"], 0.08),
]


@pytest.mark.parametrize("descripcion", ["REVOQUE EXTERIOR CON HIDROFUGO CON SILLETA",
                                         "Revoque exterior a la cal con silleta"])
def test_mapeo_silleta(descripcion: str) -> None:
    assert match_recipe(descripcion, _SIN_B)["plantillas"] == [("5.5.3", 1.0)]
    assert "silleta no está en la fórmula" in match_recipe(descripcion, _SIN_B)["nota"]
    assert match_recipe(descripcion, _CON_B)["plantillas"] == [("5.5.7", 1.0)]
    # Without silleta it is the common one, with or without the new recipe
    assert match_recipe("REVOQUE EXTERIOR CON HIDROFUGO", _CON_B)["plantillas"] == [("5.5.3", 1.0)]


@pytest.mark.parametrize("descripcion, codigos, espesor", _GINKGO_CONTRAPISOS)
def test_mapeo_contrapiso_celular_si_la_obra_lo_usa(descripcion: str, codigos: list, espesor: float) -> None:
    fila = {"codigo": "4.2.7", "descripcion": descripcion, "unidad": "m²", "cantidad": 100, "nivel": "item",
            "excel": {}}
    rule = rule_for(fila, _CON_B, obra=_CELULAR)
    assert [c for c, _ in rule["plantillas"]] == codigos
    assert rule["porque"] == "La obra usa hormigón celular (hoja 4.1-6)"
    assert rule["nota"].startswith("La obra usa hormigón celular (hoja 4.1-6).")
    rows, params, _ = expand_item(fila, rule, _CON_B)
    assert params["espesor_m"] == espesor
    cem = next(r for r in rows if r["codigo"] == "CEM")
    assert cem["cantidad"] == pytest.approx(100 * espesor * 15)
    # Today's rule when the obra does not use it, when there is no detail sheet, or without 5.2.4
    for templates, obra in ((_CON_B, _CASCOTE), (_CON_B, None), (_SIN_B, _CELULAR)):
        hoy = rule_for(fila, templates, obra=obra)
        assert [c for c, _ in hoy["plantillas"]] == [c if c != "5.2.4" else "5.2.3" for c in codigos]
        assert "porque" not in hoy


@pytest.mark.parametrize("descripcion", ["CONTRAPISO ALIVIANADO e=12cm", "Contrapiso de hormigón celular e: 12 cm",
                                         "RELLENO CON CONTRAPISO BOMBEADO ESP. 12CM"])
def test_mapeo_contrapiso_celular_por_palabra(descripcion: str) -> None:
    fila = {"codigo": "1", "descripcion": descripcion, "unidad": "m2", "cantidad": 10, "nivel": "item", "excel": {}}
    rule = rule_for(fila, _CON_B)  # obra without detail sheets
    assert rule["plantillas"] == [("5.2.4", 1.0)] and "dice" in rule["porque"]
    assert expand_item(fila, rule, _CON_B)[1]["espesor_m"] == 0.12
    sin = rule_for(fila, _SIN_B)
    assert sin["plantillas"] == [("5.2.3", 0.12)] and "porque" not in sin


def test_mapeo_telgopor_celular_por_palabra() -> None:
    rule = match_recipe("TELGOPOR 50 mm + CONTRAPISO CELULAR e: 8cm", _CON_B)
    assert rule["plantillas"] == [("8.3", 1.0), ("5.2.4", 1.0)]


def test_obra_usa_solo_las_hojas_de_detalle() -> None:
    from app.obra_import import obra_usa
    excel = {"SUB-CONT": {"descripcion": "BOMBEO HORMIGON CELULAR", "origen": "detalle", "hoja": "4.1-6"},
             "SUB-HC": {"descripcion": "Hormigón celular", "origen": "lista", "hoja": "00_Sub"}}
    assert obra_usa(excel) == [{"texto": "BOMBEO HORMIGON CELULAR", "hoja": "4.1-6"}]
    assert obra_usa(None) == []


def _ginkgo_workbook() -> openpyxl.Workbook:
    """The contrapisos of Ginkgo's 01_C&P and its sheet 4.1-6 (bombeo de hormigón celular)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_C&P"
    ws["A1"] = "EDIFICIO GINKGO"
    filas = [(None, "4. ALBAÑILERIA", None, None)] + [
        (f"4.{i}", desc, "m²", 100) for i, (desc, _, _) in enumerate(_GINKGO_CONTRAPISOS, start=1)
    ] + [("4.9", "CONTRAPISO/CARPETA EN BALCONES e=4cm", "m²", 20)]
    for r, (a, b, c, d) in enumerate(filas, start=8):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
        ws.cell(r, 3, c)
        ws.cell(r, 4, d)
    _detail_sheet(wb, "4.1-6", "4.1-6", "CONTRAPISO EN HALL + RAMPAS. ESP.=10cm.", [
        ("MATERIALES", [("CEM", "Bolsa Cemento Loma Negra 25 k", "u", 9000)]),
        ("SUBCONTRATOS", [("SUB-CONT", "BOMBEO HORMIGON CELULAR. Incluye mano de obra y bomba.", "m2", 13000)]),
    ])
    return wb


def test_ginkgo_elige_hormigon_celular() -> None:
    from app.obra_import import excel_prices, obra_usa
    wb = _ginkgo_workbook()
    obra = obra_usa(excel_prices(wb))
    plan = build_plan(parse_obra(wb), _CON_B, obra=obra)
    elegidas = {i["descripcion"]: i["plantillas"] for i in plan["items"] if i["nivel"] == "item"}
    for descripcion, codigos, _ in _GINKGO_CONTRAPISOS:
        assert elegidas[descripcion] == codigos, descripcion
    assert elegidas["CONTRAPISO/CARPETA EN BALCONES e=4cm"] == ["5.4.1"]
    # Without the new recipe (correction not applied): as today
    plan = build_plan(parse_obra(wb), _SIN_B, obra=obra)
    assert {tuple(i["plantillas"]) for i in plan["items"] if i["nivel"] == "item"} == {
        ("5.2.3",), ("8.3", "5.2.3"), ("5.4.1",)}


@pytest.mark.parametrize("descripcion, espesor", [
    ("TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm", 0.08),
    ("CONTRAPISO EN HALL + RAMPAS. ESP.=10cm.", 0.10),
    ("CONTRAPISO e=10cm", 0.10),
    ("CONTRAPISO/ CARPETA e=10cm", 0.10),
    ("TELGOPOR 50 mm + CONTRAPISO/CARPETA EN AZOTEA INACCESIBLE (BALCONES) e=4cm", 0.04),
    ("LOSA E 12 CM", 0.12),
    ("Contrapiso espesor 7,5 cm", 0.075),
    ("TELGOPOR 50 mm + CONTRAPISO", None),   # 50 mm is the board, not the contrapiso
    ("CONTRAPISO DE 10 CM", None),           # no E / ESP / ESPESOR before the number
    ("CONTRAPISO", None),
])
def test_espesor_from_the_name(descripcion: str, espesor: float | None) -> None:
    assert espesor_m_from(descripcion) == espesor


@pytest.mark.parametrize("descripcion, factores, origenes", [
    ("TELGOPOR 50 mm + CONTRAPISO EN AZOTEA ACCESIBLE e: 8cm", [1.0, 0.08], ["regla", "nombre"]),
    ("TELGOPOR 50 mm + CONTRAPISO/CARPETA EN AZOTEA INACCESIBLE (BALCONES) e=4cm", [1.0, 0.04],
     ["regla", "nombre"]),
    ("TELGOPOR 50 mm + CONTRAPISO EN AZOTEA", [1.0, 0.08], ["regla", "supuesto"]),
    ("CONTRAPISO EN HALL + RAMPAS. ESP.=10cm.", [0.10], ["nombre"]),
    ("CONTRAPISO/ CARPETA e=12cm", [0.12], ["nombre"]),
    ("CONTRAPISO", [0.10], ["supuesto"]),
])
def test_contrapiso_factor_is_its_thickness(descripcion: str, factores: list, origenes: list) -> None:
    rule = rule_for({"descripcion": descripcion, "unidad": "m²"}, TEMPLATES)
    assert [f for _, f in rule["plantillas"]] == factores
    assert rule["origen_factor"] == origenes
    assert rule["falta_factor"] == []
    # In another unit than the rule's, the thickness does not convert: it is asked
    rule = rule_for({"descripcion": descripcion, "unidad": "m3"}, TEMPLATES)
    assert "5.2.3" in rule["falta_factor"]


def test_recipe_uses_obra_height(plan: dict) -> None:
    item = _item(plan, "4.2.1")
    assert item["plantilla"] == "5.1.4"
    assert item["parametros"] == {"altura_m": 3.0}
    hierro = next(r for r in item["recursos"] if r["codigo"] == "HADN6")
    assert hierro["cantidad"] == pytest.approx((30 / 3) * 4 / 12, abs=1e-4)
    mo = next(r for r in item["recursos"] if r["tipo"] == "mano_obra")
    assert mo["dias"] == 2.0
    # Sin desperdicio propio: se hereda en el SQL (plantilla > organización)
    assert hierro["desperdicio_pct"] is None
    assert next(r for r in item["recursos"] if r["codigo"] == "LH18")["desperdicio_pct"] == 15.0


def test_unit_factor_is_kept_in_formula(plan: dict) -> None:
    item = _item(plan, "3.1.3")  # 100 ml de tensor = 8 m³
    h30 = next(r for r in item["recursos"] if r["codigo"] == "H30")
    assert h30["cantidad"] == pytest.approx(8.0)
    # La fórmula guardada sigue funcionando al recalcular con Q en ml
    assert evaluate(h30["formula"], {"Q": 100}) == pytest.approx(8.0)
    mo = next(r for r in item["recursos"] if r["tipo"] == "mano_obra")
    assert mo["dias"] == pytest.approx(4.0)
    assert evaluate(f"Q / ({mo['rendimiento']})", {"Q": 100}) == pytest.approx(4.0)


def test_composite_and_client_material(plan: dict) -> None:
    item = _item(plan, "4.2.3")
    assert item["plantillas"] == ["8.3", "5.2.3"]
    es = next(r for r in item["recursos"] if r["codigo"] == "ES")
    assert es["cantidad"] == pytest.approx(50 * 0.08 * 0.8)

    porc = next(r for r in _item(plan, "4.2.4")["recursos"] if r["codigo"] == "RP-PORC")
    assert porc["lo_compra_cliente"] is True
    assert _item(plan, "4.2.4")["revisar"]


def test_items_without_recipe(plan: dict) -> None:
    codes = [i["codigo"] for i in plan["sin_receta"]]
    assert codes == ["1.1", "4.2.2"]
    assert "Sin fórmula" in item_notes(_item(plan, "1.1"))
    assert plan["plantillas_faltantes"] == []


def test_missing_template_goes_without_recipe(parsed: dict) -> None:
    templates = {k: v for k, v in TEMPLATES.items() if k != "4.1.7"}
    plan = build_plan(parsed, templates)
    assert plan["plantillas_faltantes"] == ["4.1.7"]
    assert "3.1.3" in [i["codigo"] for i in plan["sin_receta"]]


def test_report(parsed: dict, plan: dict) -> None:
    md = report_markdown(parsed, plan, "obra.xlsx")
    assert "## Ítems sin fórmula en el Maestro" in md
    assert "ARISTAS DE YESO" in md
    assert "3.1.3" in md  # con receta pero en $0 en el Excel


def test_sql_script(plan: dict) -> None:
    sql = sql_script(plan, "OBRA (Fase 5)", "obra.xlsx")
    assert "RAISE EXCEPTION 'Faltan plantillas del Maestro" in sql
    assert "ya existe" in sql
    assert "ARRAY['4.1.7', '5.1.4', '5.2.3', '7.1.1', '8.3']::text[]" in sql
    # Fase 4: indirectos de la obra, fecha de precios y fecha del precio de cada recurso
    assert "Falta correr migrations/008" in sql
    assert "status, indirectos, precios_al)" in sql
    assert "'estructura_pct', coalesce((x.c->>'estructura_pct')::numeric, 15)" in sql
    assert "cantidad_redondeo, precio_fecha)" in sql
    # Auditoría: precio vigente a la fecha del presupuesto (regla de la Fase 4), el mismo
    # para el control y para la carga; sin precio válido se frena, salvo que se fuerce.
    # El comportamiento se prueba contra Postgres en tests/test_obra_import_pg.py.
    assert "x.fecha IS NULL OR x.fecha <= v_fecha" in sql
    assert "JOIN obra_precios op ON op.n = r.n" in sql
    assert "v_permitir_sin_precio boolean := false" in sql
    assert "RAISE EXCEPTION 'Recursos sin precio valido" in sql
    assert "Recalcular obra" in sql
    assert sql.count("$obra$") == 2
    pglast = pytest.importorskip("pglast")
    assert len(pglast.parse_sql(sql)) == 2


def test_price_problems_in_report(parsed: dict, plan: dict) -> None:
    from datetime import date

    from app.obra_import import price_problems

    entries = [
        {"tipo": "material", "codigo": "LH18", "precio_sin_iva": 500, "fecha_precio": None},
        {"tipo": "material", "codigo": "HADN6", "precio_sin_iva": 0, "fecha_precio": None},       # sin precio
        {"tipo": "material", "codigo": "H30", "precio_sin_iva": 100, "fecha_precio": "2099-01-01"},  # futuro
        {"tipo": "material", "codigo": "ES", "precio_sin_iva": 10, "fecha_precio": None},
        {"tipo": "material", "codigo": "es", "precio_sin_iva": 12, "fecha_precio": None},          # duplicado
    ]
    probs = {p["codigo"]: p for p in price_problems(plan, entries, date(2026, 9, 30))}
    assert probs["HADN6"]["motivo"] == "sin precio"
    assert probs["H30"]["motivo"] == "sin precio"
    assert probs["ES"]["motivo"] == "código duplicado en el catálogo"
    assert probs["MO-OF"]["motivo"] == "no está en el catálogo"
    assert "LH18" not in probs
    assert "RP-PORC" not in probs  # lo compra el cliente: no hace falta precio
    md = report_markdown(parsed, plan, "obra.xlsx", list(probs.values()))
    assert "## Recursos sin precio válido (la carga se frena)" in md
    assert "`HADN6`" in md


def test_dated_zero_is_a_price(plan: dict) -> None:
    """'Va en $0': a 0 with a date is a confirmed price; an undated 0 is still no price."""
    from datetime import date

    from app.obra_import import price_problems

    entries = [
        {"tipo": "material", "codigo": "HADN6", "precio_sin_iva": 0, "fecha_precio": "2026-09-01"},
        {"tipo": "material", "codigo": "LH18", "precio_sin_iva": 0, "fecha_precio": None},
    ]
    probs = {p["codigo"]: p for p in price_problems(plan, entries, date(2026, 9, 30))}
    assert "HADN6" not in probs
    assert probs["LH18"]["motivo"] == "sin precio"


SUGGEST = {
    "5.1.4": {"nombre": "LADRILLO CERAMICO HUECO DEL 18", "unidad": "m2"},
    "5.1.5": {"nombre": "LADRILLO CERAMICO HUECO DEL 12", "unidad": "m2"},
    "5.1.7": {"nombre": "LADRILLO COMUN", "unidad": "m2"},
    "4.2.6": {"nombre": "escaleras", "unidad": "gl",
              "descripcion": "Importado del Maestro TERRAC (solapa 4.2.6). Cantidad de ejemplo del Excel: 18 gl."},
    "6.1": {"nombre": "CIELORRASO", "unidad": "m2", "descripcion": "Aplicado de yeso"},
}


def test_suggest_recipes() -> None:
    found = suggest_recipes("Muro de mampostería en ladrillo hueco del 18", SUGGEST)
    assert [c for c, _, _ in found] == ["5.1.4", "5.1.5", "5.1.7"]
    assert found[0][1] >= 0.6
    assert found[0][2] == "Se parece por 'ladrillo', 'hueco', '18'"
    assert found[1][1] < found[0][1]  # the number counts


def test_suggest_recipes_needs_words() -> None:
    assert suggest_recipes("Obrador", SUGGEST) == []
    assert suggest_recipes("Bolsas de 18", SUGGEST) == []  # a number alone is not enough
    # One shared word is too generic to propose the recipe by itself
    [(codigo, score, _)] = suggest_recipes("ESMALTE EN ESCALERA COMPLETA", SUGGEST)
    assert (codigo, score) == ("4.2.6", 0.5)
    # Typos and plurals, and the recipe's own description (not the import boilerplate)
    assert suggest_recipes("CIELORASOS APLICADOS EN YESO", SUGGEST)[0][:2] == ("6.1", 1.0)


# ── Precios que el Excel de la obra ya trae ──────────────────────────────────


def _detail_sheet(wb: openpyxl.Workbook, title: str, codigo: str, trabajo: str, sections: list) -> None:
    """A detail sheet like the TERRAC ones: row 3 = task, then sections with a 'Código' header."""
    ws = wb.create_sheet(title)
    ws["A1"] = "EDIFICIO DE PRUEBA"
    ws["A3"], ws["B3"] = codigo, trabajo
    r = 5
    for name, rows in sections:
        ws.cell(r, 1, name)
        header = ["Código", "Descripción", "Unidad", "Cantidad", "Dias", "Desperdicio",
                  "Cantidad + Desperdicio", "Precio Unitario", "Subtotal"]
        for c, label in enumerate(header, start=1):
            ws.cell(r + 1, c, label)
        r += 2
        for cod, desc, unidad, precio in rows:
            ws.cell(r, 1, cod)
            ws.cell(r, 2, desc)
            ws.cell(r, 3, unidad)
            ws.cell(r, 4, 1)
            ws.cell(r, 8, precio)
            r += 1
        ws.cell(r, 1, "TOTAL")
        ws.cell(r + 1, 1, "m²")
        ws.cell(r + 1, 8, 999)  # below TOTAL: not a resource
        r += 4


def _prices_workbook(with_lists: bool = True) -> openpyxl.Workbook:
    wb = _workbook()
    if with_lists:
        mat = wb.create_sheet("00_Mat")
        mat.append(["CODIGO", "MATERIALES CORRALON", "UNIDAD", "PRECIO CON IVA", "PRECIO SIN IVA"])
        mat.append(["RE-PLI20", "Latex interior 20 l", "u", 121000, 100000, "ML", datetime(2026, 9, 17)])
        mat.append(["CEM", "Cemento 25 k", "u", 9000, 7438.02, "Corralón", datetime(2026, 9, 17)])
        mat.append(["LP8", "Ladrillo portante 8", "u", None, 0])  # price 0: not a price
        mat.append(["RE-END15", "Enduido 15 l", "u", 30661, None])  # only "con IVA": proposed with a warning
    _detail_sheet(wb, "5.2-6", "5.2-6", "EJECUCION DE PINTURA EN PAREDES.  INCLUYE ENDUIDO", [
        ("MATERIALES", [
            ("RE-PLI20", "Albalatex Extra Mate 20 l (u)", "u", 200000),
            ("re-end", "Enduido  plástico\n", "u", 60000),
            ("RE-CIN", "Cinta de papel", "u", 0),            # price 0: ignored
            ("RE-ROD", "Rodillo", "u", "#REF!"),             # error: ignored
            ("RE-LIJ", "Lija", "u", "a confirmar"),          # text: ignored
            (None, None, None, 500),                         # no code: ignored
        ]),
        ("MANO DE OBRA - PERSONAS", [("MO-OF", "Oficial", "u", 100000)]),
        ("MANO DE OBRA - SUBCONTRATOS", [("SUB-PI", "Pintura interior", "m2", 15000)]),
    ])
    _detail_sheet(wb, "5.1-3", "5.1-3", "EJECUCION DE PINTURA EN CIELORRASOS", [
        ("MATERIALES", [("RE-PLI20", "Albalatex", "u", 140000), ("RE-END", "Enduido", "u", 60000)]),
        ("MANO DE OBRA - EQUIPOS", [("E-AND", "Andamio", "u", 5000)]),
        ("SUBCONTRATOS", [("SUB-PET", "Pintura exterior", "m2", 18000)]),
    ])
    # Odd sheets: a name with spaces (a copy of 5.2-6) and one that is a date: skipped
    _detail_sheet(wb, "REV PROY", "5.2-6", "COPIA", [("MATERIALES", [("XX-1", "Otro", "u", 5)])])
    _detail_sheet(wb, "2026-02-02", "2.2", "FECHA", [("MATERIALES", [("XX-2", "Otro", "u", 5)])])
    # A detail sheet broken by #REF! in the title and the headers: nothing to read, no failure
    broken = wb.create_sheet("2.3")
    broken["A3"], broken["B3"], broken["A5"], broken["A6"] = "#REF!", "#REF!", "MATERIALES", "#REF!"
    broken["A7"], broken["H7"] = "#REF!", "#REF!"
    return wb


def test_excel_prices_detail_wins_over_the_list() -> None:
    from app.obra_import import excel_prices

    precios = excel_prices(_prices_workbook())
    pli = precios["RE-PLI20"]
    assert pli == {
        "codigo": "RE-PLI20", "descripcion": "Albalatex Extra Mate 20 l (u)", "unidad": "u", "tipo": "material",
        "precio": 200000.0, "fecha": None, "proveedor": None, "nota": None, "origen": "detalle", "hoja": "5.2-6",
        "trabajo": "EJECUCION DE PINTURA EN PAREDES. INCLUYE ENDUIDO",
        # the other detail sheet and the list, once per different price
        "otros": [{"precio": 140000.0, "hoja": "5.1-3"}, {"precio": 100000.0, "hoja": "00_Mat"}],
    }
    # Same price in another sheet is not "otro"; codes are normalized
    assert precios["RE-END"]["otros"] == []
    assert precios["RE-END"]["descripcion"] == "Enduido plástico"
    # Sections → tipo
    assert precios["MO-OF"]["tipo"] == "mano_obra"
    assert precios["SUB-PI"]["tipo"] == "subcontrato"
    assert precios["E-AND"]["tipo"] == "equipo"
    assert precios["SUB-PET"]["tipo"] == "subcontrato"


def test_excel_prices_from_the_list_and_what_is_ignored() -> None:
    from app.obra_import import excel_prices

    precios = excel_prices(_prices_workbook())
    assert precios["CEM"] == {
        "codigo": "CEM", "descripcion": "Cemento 25 k", "unidad": "u", "tipo": "material", "precio": 7438.02,
        "fecha": "2026-09-17", "proveedor": "Corralón", "nota": None, "origen": "lista", "hoja": "00_Mat",
        "trabajo": None, "otros": [],
    }
    # A row with only "precio con IVA" is still proposed, with the warning
    assert (precios["RE-END15"]["precio"], precios["RE-END15"]["nota"]) == (
        30661.0, "En la lista del Excel figura como precio con IVA: fijate si va sin IVA.")
    for codigo in ("RE-CIN", "RE-ROD", "RE-LIJ", "LP8", "XX-1", "XX-2", "#REF!"):
        assert codigo not in precios, codigo
    # Rows below TOTAL are not resources
    assert all(p["precio"] != 999 for p in precios.values())


def test_excel_prices_without_lists_or_details() -> None:
    from app.obra_import import excel_prices

    precios = excel_prices(_prices_workbook(with_lists=False))
    assert precios["RE-PLI20"]["otros"] == [{"precio": 140000.0, "hoja": "5.1-3"}]
    assert "CEM" not in precios
    assert excel_prices(_workbook()) == {}  # only the cómputo: nothing, and no failure
