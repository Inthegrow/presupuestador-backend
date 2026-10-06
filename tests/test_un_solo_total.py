"""Un solo total (PLAN_UN_SOLO_TOTAL 2): the price of each work is its direct cost through the
cascade of its budget, written by the server on every path that changes the direct cost, and
every reader (editor summary, client PDF, exported Excel, versions) shows that saved number.

After each operation, for every work: neto == cascade(directo, config) to the cent,
total_final == neto × 1,21, and summary neto == Σ neto == client PDF total sin IVA ==
total row of the exported Excel.
"""

from __future__ import annotations

import copy
import io
import os
import re
from unittest.mock import patch

os.environ.setdefault("SUPABASE_URL", "https://test.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

import app.routers.excel as excel_router
from app.budget_prices import effective_indirects, general_indirects
from app.calculations import CASCADE_FIELDS, calc_budget_summary, calc_cascade_indirects, is_section
from app.main import create_app
from tests.test_budget_prices import resource
from tests.test_recipes_api import MOCK_USER, ORG, FakeDB

B1 = "00000000-0000-0000-0000-0000000000b1"
B2 = "00000000-0000-0000-0000-0000000000b2"
SEC = "00000000-0000-0000-0000-0000000000c1"
W1 = "00000000-0000-0000-0000-0000000000a1"
W2 = "00000000-0000-0000-0000-0000000000a2"
W3 = "00000000-0000-0000-0000-0000000000a3"
CATALOG = "00000000-0000-0000-0000-0000000000d1"
R11 = "00000000-0000-0000-0000-000000000111"
R12 = "00000000-0000-0000-0000-000000000112"
R13 = "00000000-0000-0000-0000-000000000113"
R14 = "00000000-0000-0000-0000-000000000114"
R21 = "00000000-0000-0000-0000-000000000121"
R22 = "00000000-0000-0000-0000-000000000122"
R23 = "00000000-0000-0000-0000-000000000123"
R31 = "00000000-0000-0000-0000-000000000131"
R41 = "00000000-0000-0000-0000-000000000141"
NUMBERS = (
    "cantidad", "mat_unitario", "mo_unitario", "mat_total", "mo_total", "directo_total", *CASCADE_FIELDS,
)


def _work(iid, code, cantidad, budget=B1, parent=SEC, **over):
    # Stale totals: the fixture prices everything with the reference path first
    return {"id": iid, "budget_id": budget, "org_id": ORG, "parent_id": parent, "code": code,
            "description": f"Trabajo {code}", "unidad": "m2", "cantidad": cantidad, "notas": None,
            "sort_order": int(code.replace(".", "")), "template_id": None, "parametros": {},
            "mat_unitario": 0, "mo_unitario": 0, **dict.fromkeys(NUMBERS[3:], 0), **over}


def tables():
    return {
        "budgets": [
            {"id": B1, "org_id": ORG, "name": "Ginkgo", "desperdicio_pct": None, "indirectos": {},
             "precios_al": "2026-09-01", "created_at": "2026-10-01T10:00:00", "description": None},
            # All its own values: a change of the general ones never touches it
            {"id": B2, "org_id": ORG, "name": "Propio", "desperdicio_pct": None,
             "indirectos": {**general_indirects({}), "beneficio_pct": 15}, "precios_al": "2026-09-01"},
        ],
        "budget_items": [
            {"id": SEC, "budget_id": B1, "org_id": ORG, "parent_id": None, "code": "1",
             "description": "ALBAÑILERÍA", "notas": "Seccion", "sort_order": 0, "cantidad": None,
             **dict.fromkeys(NUMBERS[1:], 0)},
            _work(W1, "1.1", 10),
            _work(W2, "1.2", 20, parametros={"espesor": 0.1}),
            _work(W3, "1.3", 3.5),
            _work("00000000-0000-0000-0000-0000000000e1", "1.1", 7, budget=B2, parent=None),
        ],
        "item_resources": [
            # W1: every resource tipo, plus one the client buys
            resource(id=R11, item_id=W1, codigo="LAD", cantidad=100, precio_unitario=52.37),
            resource(id=R12, item_id=W1, tipo="mano_obra", codigo="OF", trabajadores=1, dias=2,
                     cargas_sociales_pct=0, precio_unitario=1234.57),
            resource(id=R13, item_id=W1, tipo="equipo", codigo="EQ", cantidad=1, precio_unitario=300),
            resource(id=R14, item_id=W1, codigo="CLI", cantidad=3, precio_unitario=999,
                     lo_compra_cliente=True),
            # W2: a formula with a parameter, a mo_material and a subcontrato
            resource(id=R21, item_id=W2, codigo="H30", formula="Q * espesor", cantidad=2,
                     precio_unitario=1000.01),
            resource(id=R22, item_id=W2, tipo="mo_material", codigo="CLAV", cantidad=1, precio_unitario=55.5),
            resource(id=R23, item_id=W2, tipo="subcontrato", codigo="SUB", cantidad=1, precio_unitario=777.77),
            # W3: one material
            resource(id=R31, item_id=W3, codigo="PINT", cantidad=7, precio_unitario=333.33),
            resource(id=R41, item_id="00000000-0000-0000-0000-0000000000e1", codigo="LAD", cantidad=10,
                     precio_unitario=52.37),
        ],
        "indirect_config": [{"id": "cfg", "org_id": ORG, "desperdicio_pct": None}],
        "item_templates": [],
        "budget_versions": [],
        "price_catalogs": [{"id": CATALOG, "org_id": ORG, "name": "Lista octubre"}],
        "catalog_entries": [
            {"id": "e1", "catalog_id": CATALOG, "org_id": ORG, "tipo": "material", "codigo": "LAD",
             "precio_sin_iva": 61.11, "fecha_precio": "2026-10-01"},
            {"id": "e2", "catalog_id": CATALOG, "org_id": ORG, "tipo": "equipo", "codigo": "EQ",
             "precio_sin_iva": 450, "fecha_precio": "2026-10-01"},
            {"id": "e3", "catalog_id": CATALOG, "org_id": ORG, "tipo": "material", "codigo": "CLI",
             "precio_sin_iva": 5000, "fecha_precio": "2026-10-01"},
        ],
        "catalog_price_history": [],
        "item_audits": [],
        "audit_logs": [],
    }


@pytest.fixture
def client():
    app = create_app()
    from app.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: MOCK_USER
    return TestClient(app)


@pytest.fixture
def db(client):
    fake = FakeDB(tables())
    targets = ("budgets", "analysis", "indirects", "excel", "catalogs", "templates", "ai")
    patches = [patch(f"app.routers.{name}.get_data_db", return_value=fake) for name in targets]
    for p in patches:
        p.start()
    try:
        # Start from a consistent budget: the reference path (Coeficiente de pase)
        for bid in (B1, B2):
            assert client.post(f"/budgets/{bid}/cascade-recalculate").status_code == 200
        yield fake
    finally:
        for p in patches:
            p.stop()


# ── Helpers ─────────────────────────────────────────────────────────────────


def cents(value) -> int:
    return round(float(value or 0) * 100)


def item(db, iid) -> dict:
    return next(i for i in db.tables["budget_items"] if i["id"] == iid)


def works(db, bid=B1) -> list[dict]:
    return [i for i in db.tables["budget_items"] if i["budget_id"] == bid and not is_section(i)]


def config(db, bid=B1) -> dict:
    budget = next(b for b in db.tables["budgets"] if b["id"] == bid)
    org = db.tables["indirect_config"][0] if db.tables["indirect_config"] else {}
    return effective_indirects(org, budget)


def numbers(db, bid=B1) -> dict:
    return {i["id"]: {k: i.get(k) for k in NUMBERS} for i in works(db, bid)}


def client_pdf(client, bid) -> dict:
    """What the client PDF shows (client_pdf_data as the endpoint calls it) + its text check."""
    seen: dict = {}
    real = excel_router.client_pdf_data

    def spy(all_items, cfg):
        seen["data"] = real(all_items, cfg)
        return seen["data"]

    with patch("app.routers.excel.client_pdf_data", side_effect=spy):
        r = client.get(f"/budgets/{bid}/export/pdf?vista=cliente")
    assert r.status_code == 200, r.text
    fitz = pytest.importorskip("fitz")
    text = "\n".join(page.get_text() for page in fitz.open(stream=r.content, filetype="pdf"))
    total = "$ " + f"{seen['data']['total_sin_iva']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    assert re.search(r"Total sin IVA\s+" + re.escape(total), text), text
    return seen["data"]


def excel_totals(client, bid) -> dict:
    r = client.get(f"/budgets/{bid}/export/excel")
    assert r.status_code == 200, r.text
    df = pd.read_excel(io.BytesIO(r.content))
    total = df[df["Codigo"] == "TOTAL"].iloc[0]
    body = df[df["Codigo"] != "TOTAL"]
    # The totals row adds up its column, and the ladder adds up to the price without IVA
    for col in ("Directo Total", "Indirecto Total", "Beneficio Total", "Impuestos Total", "Precio sin IVA"):
        assert cents(total[col]) == cents(pd.to_numeric(body[col], errors="coerce").sum()), col
    assert abs(total["Directo Total"] + total["Indirecto Total"] + total["Beneficio Total"]
               + total["Impuestos Total"] - total["Precio sin IVA"]) <= 0.011 * len(body)
    return total


def check(client, db, bid=B1) -> dict:
    """Every work priced by the cascade of its budget, and every reader shows the same total."""
    cfg = config(db, bid)
    iva = 1 + cfg["iva_pct"] / 100
    rows = works(db, bid)
    assert rows
    for w in rows:
        expected = calc_cascade_indirects({"directo_total": w["directo_total"]}, cfg)
        for key in CASCADE_FIELDS:
            assert cents(w[key]) == cents(expected[key]), (w["code"], key, w[key], expected[key])
        assert abs(w["total_final"] - w["neto_total"] * iva) <= 0.01, w["code"]

    neto = sum(cents(w["neto_total"]) for w in rows)
    summary = client.get(f"/budgets/{bid}/full").json()["analysis"]
    assert summary["items_count"] == len(rows)  # T11: rubros are not works
    assert cents(summary["neto_total"]) == neto
    assert cents(summary["total_final"]) == sum(cents(w["total_final"]) for w in rows)
    pdf = client_pdf(client, bid)
    assert cents(pdf["total_sin_iva"]) == neto
    assert cents(pdf["total_con_iva"]) == cents(summary["total_final"])
    assert cents(excel_totals(client, bid)["Precio sin IVA"]) == neto
    return summary


def url(*parts) -> str:
    return "/budgets/" + "/".join(parts)


# ── Starting point ──────────────────────────────────────────────────────────


class TestReference:
    def test_default_cascade_of_100(self):
        it = calc_cascade_indirects({"directo_total": 100}, {})
        assert (it["indirecto_total"], it["beneficio_total"], it["impuestos_total"]) == (34, 13.4, 12.09)
        assert (it["neto_total"], it["iva_total"], it["total_final"]) == (159.49, 33.49, 192.98)

    def test_fixture_is_consistent(self, client, db):
        check(client, db)
        check(client, db, B2)
        # W1: material + mano de obra + equipo; what the client buys costs 0
        # (unit prices are rounded to cents, like the editor shows them)
        assert item(db, W1)["directo_total"] == pytest.approx(5237 + 2469.14 + 300, abs=0.1)


# ── Writes: every path that changes the direct cost prices the work ─────────


class TestItemPatch:
    def test_cantidad(self, client, db):
        r = client.patch(url(B1, "items", W3), json={"cantidad": 7})
        assert r.status_code == 200, r.text
        assert item(db, W3)["directo_total"] == pytest.approx(7 * item(db, W3)["mat_unitario"])
        check(client, db)
        assert r.json()["item"]["neto_total"] == item(db, W3)["neto_total"]

    def test_cantidad_with_formula_resources(self, client, db):
        r = client.patch(url(B1, "items", W2), json={"cantidad": 40})
        assert r.status_code == 200, r.text
        h30 = next(x for x in db.tables["item_resources"] if x["id"] == R21)
        assert h30["cantidad"] == pytest.approx(4)
        check(client, db)

    def test_unit_price(self, client, db):
        assert client.patch(url(B1, "items", W3), json={"mo_unitario": 100.5}).status_code == 200
        assert item(db, W3)["mo_total"] == pytest.approx(351.75)
        check(client, db)

    @pytest.mark.parametrize("body", [
        {"notas_calculo": "Medido en obra el 05/10"},
        {"description": "Muro de ladrillo hueco"},
        {"unidad": "m3"},
        {"code": "1.9"},
    ])
    def test_text_only_changes_no_number(self, client, db, body):
        # T1: saving a note used to drop Ingresos Brutos and cheque from the price
        antes = numbers(db)
        r = client.patch(url(B1, "items", W1), json=body)
        assert r.status_code == 200, r.text
        assert numbers(db) == antes
        check(client, db)

    def test_totals_from_the_client_are_ignored(self, client, db):
        antes = numbers(db)
        r = client.patch(url(B1, "items", W1), json={"neto_total": 1, "indirecto_total": 2})
        assert r.status_code == 200
        assert numbers(db) == antes


class TestResources:
    def test_create(self, client, db):
        r = client.post(url(B1, "items", W3, "resources"), json={
            "tipo": "equipo", "codigo": "AND", "descripcion": "Andamio", "cantidad": 2, "precio_unitario": 150.25})
        assert r.status_code == 200, r.text
        assert item(db, W3)["directo_total"] == pytest.approx(7 * 333.33 + 300.5, abs=0.05)
        check(client, db)

    def test_update(self, client, db):
        r = client.patch(url(B1, "items", W1, "resources", R13), json={"precio_unitario": 512.34})
        assert r.status_code == 200, r.text
        assert item(db, W1)["directo_total"] == pytest.approx(5237 + 2469.14 + 512.34, abs=0.1)
        check(client, db)

    def test_delete(self, client, db):
        assert client.delete(url(B1, "items", W2, "resources", R23)).status_code == 204
        check(client, db)

    def test_bulk(self, client, db):
        r = client.post(url(B1, "items", W3, "resources", "bulk"), json={"resources": [
            {"tipo": "material", "codigo": "A", "cantidad": 1, "precio_unitario": 10.01},
            {"tipo": "subcontrato", "codigo": "B", "cantidad": 1, "precio_unitario": 20.02},
        ]})
        assert r.status_code == 200, r.text
        check(client, db)

    def test_parameter(self, client, db):
        r = client.patch(url(B1, "items", W2, "parametros"), json={"parametros": {"espesor": 0.25}})
        assert r.status_code == 200, r.text
        assert next(x for x in db.tables["item_resources"] if x["id"] == R21)["cantidad"] == pytest.approx(5)
        check(client, db)


class TestNewWork:
    def test_by_hand_gets_the_full_price(self, client, db):
        # T3: a work entered by hand used to have no indirects nor profit
        r = client.post(url(B1, "items"), json=[{
            "parent_id": SEC, "code": "1.4", "description": "Revoque", "unidad": "m2", "cantidad": 12,
            "mat_unitario": 100.1, "mo_unitario": 55.55, "indirecto_total": 0, "beneficio_total": 0}])
        assert r.status_code == 200, r.text
        nuevo = next(i for i in works(db) if i["code"] == "1.4")
        assert nuevo["directo_total"] == pytest.approx(1867.8)
        assert nuevo["neto_total"] > nuevo["directo_total"] * 1.5
        check(client, db)

    def test_section_by_hand_has_no_price(self, client, db):
        r = client.post(url(B1, "items"), json=[{"description": "PINTURA", "notas": "Seccion"}])
        assert r.status_code == 200
        seccion = next(i for i in db.tables["budget_items"] if i.get("description") == "PINTURA")
        assert not seccion.get("neto_total") and not seccion.get("total_final")
        check(client, db)


class TestPriceList:
    """T9: applying a price list is the same recalculation as "Actualizar precios"."""

    def _assert_w1(self, db):
        w1 = item(db, W1)
        # Every resource tipo (equipo too) and what the client buys stays at $0
        assert w1["directo_total"] == pytest.approx(round(100 * 61.11, 2) + 2469.14 + 450, abs=0.1)
        cli = next(x for x in db.tables["item_resources"] if x["id"] == R14)
        assert cli["subtotal"] == 0 and cli["precio_unitario"] == 5000

    def test_assign_catalog(self, client, db):
        r = client.post(url(B1, "assign-catalog", CATALOG))
        assert r.status_code == 200, r.text
        assert r.json()["updated_count"] == 3
        self._assert_w1(db)
        summary = check(client, db)
        assert r.json()["summary"]["neto_total"] == summary["neto_total"]

    def test_catalogs_apply(self, client, db):
        r = client.post(f"/catalogs/apply/{B1}/{CATALOG}")
        assert r.status_code == 200, r.text
        assert (r.json()["items_matched"], r.json()["items_unmatched"]) == (3, 5)
        assert r.json()["total_updated"] == pytest.approx(round(100 * 61.11, 2) + 450)
        self._assert_w1(db)
        check(client, db)


class TestRecalculate:
    def test_same_as_cascade_recalculate(self, client, db):
        # T4: the editor's "Recálculo completo" used to drop the taxes
        for w in works(db):
            w.update(neto_total=1, impuestos_total=None, iva_total=None, total_final=None)
        r = client.post(url(B1, "recalculate"))
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["items_total"] == 4 and body["items_updated"] == 3
        summary = check(client, db)
        assert body["summary"]["neto_total"] == summary["neto_total"]


class TestPercentages:
    def test_get_returns_the_percentage_and_the_coefficient(self, client, db):
        data = client.get(url(B1, "indirects")).json()
        assert data["indirecto_pct"] == 34  # the 5 concepts, imprevistos included
        assert data["coeficiente"] == pytest.approx(1.34 * 1.1 * 1.082, abs=1e-4)
        assert client.get(url(B2, "indirects")).json()["coeficiente"] == pytest.approx(1.34 * 1.15 * 1.082, abs=1e-4)
        general = client.get("/indirects/general").json()
        assert general["indirecto_pct"] == 34 and general["coeficiente"] == data["coeficiente"]

    def test_budget_change_reprices_it(self, client, db):
        otro = numbers(db, B2)
        r = client.patch(url(B1, "indirects"), json={"beneficio_pct": 20, "imprevistos_pct": 5})
        assert r.status_code == 200, r.text
        assert r.json()["actualizados"] == 1
        assert r.json()["indirecto_pct"] == 36
        assert config(db)["beneficio_pct"] == 20
        check(client, db)
        assert numbers(db, B2) == otro

    def test_only_waste_reprices_nothing(self, client, db):
        r = client.patch(url(B1, "indirects"), json={"desperdicio_pct": 5})
        assert r.status_code == 200 and r.json()["actualizados"] == 0

    def test_affected_budgets(self, client, db):
        antes = copy.deepcopy(db.tables)
        r = client.post("/indirects/general/afectados", json={"beneficio_pct": 20})
        assert r.status_code == 200, r.text
        assert r.json() == {"presupuestos": [{"id": B1, "nombre": "Ginkgo"}]}
        assert db.tables == antes  # nothing saved
        # A value equal to the current one changes nobody
        assert client.post("/indirects/general/afectados", json={"beneficio_pct": 10}).json() == {"presupuestos": []}

    def test_general_change_with_aplicar(self, client, db):
        otro = numbers(db, B2)
        r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 200, r.text
        assert r.json()["actualizados"] == 1 and r.json()["beneficio_pct"] == 20
        assert "aplicar" not in db.tables["indirect_config"][0]
        assert config(db)["beneficio_pct"] == 20
        check(client, db)
        # The budget with its own values does not change
        assert numbers(db, B2) == otro
        check(client, db, B2)

    def test_general_change_without_aplicar_saves_only(self, client, db):
        antes = numbers(db)
        r = client.patch("/indirects/general", json={"beneficio_pct": 20})
        assert r.status_code == 200, r.text
        assert "actualizados" not in r.json()
        assert db.tables["indirect_config"][0]["beneficio_pct"] == 20
        assert numbers(db) == antes


class TestGeneralSaveIsAtomic:
    """Codex PR #42 [P1]: saving the general values reprices first and saves last, so a failure
    leaves nothing half done and a retry finishes the job."""

    B3 = "00000000-0000-0000-0000-0000000000b3"

    def _second_follower(self, client, db):
        db.tables["budgets"].append({"id": self.B3, "org_id": ORG, "name": "Otra obra", "desperdicio_pct": None,
                                     "indirectos": {}, "precios_al": "2026-09-01"})
        db.tables["budget_items"].append(_work("00000000-0000-0000-0000-0000000000f1", "1.1", 4,
                                               budget=self.B3, parent=None, mat_unitario=25, mo_unitario=0))
        assert client.post(f"/budgets/{self.B3}/cascade-recalculate").status_code == 200

    def _failing_on(self, call: int):
        import app.routers.indirects as indirects
        real = indirects.reprice_budget
        calls = {"n": 0}

        def flaky(*args, **kwargs):
            calls["n"] += 1
            result = real(*args, **kwargs)  # the budget gets written...
            if calls["n"] == call:
                raise RuntimeError("se cortó la base")  # ...and then the request dies
            return result
        return patch("app.routers.indirects.reprice_budget", side_effect=flaky)

    def test_failure_restores_and_the_retry_finishes(self, client, db):
        self._second_follower(client, db)
        antes = {bid: numbers(db, bid) for bid in (B1, self.B3)}
        with self._failing_on(2):
            r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 500
        assert r.json()["detail"]["codigo"] == "NO_SE_APLICO"
        # Nothing saved, nothing half priced: both budgets are as they were
        assert db.tables["indirect_config"][0].get("beneficio_pct") is None
        for bid in (B1, self.B3):
            assert numbers(db, bid) == antes[bid]
        # The retry finds the same two budgets and prices them with the new values
        assert len(client.post("/indirects/general/afectados", json={"beneficio_pct": 20}).json()["presupuestos"]) == 2
        r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 200, r.text
        assert r.json()["actualizados"] == 2
        assert config(db)["beneficio_pct"] == 20
        for bid in (B1, self.B3):
            check(client, db, bid)
        w1 = item(db, W1)
        assert cents(w1["neto_total"]) == cents(
            calc_cascade_indirects({"directo_total": w1["directo_total"]}, {"beneficio_pct": 20})["neto_total"])

    def test_failure_while_saving_restores_the_prices(self, client, db):
        antes = numbers(db)
        with patch("app.routers.indirects.save_org_config", side_effect=RuntimeError("sin base")):
            r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 500 and r.json()["detail"]["codigo"] == "NO_SE_APLICO"
        assert numbers(db) == antes
        check(client, db)

    def test_restore_also_fails_then_the_retry_heals(self, client, db):
        self._second_follower(client, db)
        with self._failing_on(2), patch("app.routers.indirects._restore_prices",
                                        side_effect=RuntimeError("tampoco")):
            r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 500 and r.json()["detail"]["codigo"] == "A_MEDIAS"
        assert "Volvé a guardar" in r.json()["detail"]["mensaje"]
        # The values were not saved, so the retry still sees both budgets and fixes them
        r = client.patch("/indirects/general", json={"beneficio_pct": 20, "aplicar": True})
        assert r.status_code == 200 and r.json()["actualizados"] == 2
        for bid in (B1, self.B3):
            check(client, db, bid)


class TestOldItemsUseTheirBudgetIva:
    """Codex PR #42 [P2]: items saved before iva_total existed take their budget's IVA everywhere."""

    def test_analysis_full_and_exports_agree(self, client, db):
        budget = next(b for b in db.tables["budgets"] if b["id"] == B2)
        budget["indirectos"] = {**budget["indirectos"], "iva_pct": 10}
        for w in works(db, B2):
            w["iva_total"] = None
            w["total_final"] = None
        neto = sum(cents(w["neto_total"]) for w in works(db, B2))
        esperado = round(neto * 0.10)
        analysis = client.get(url(B2, "analysis")).json()
        full = client.get(url(B2, "full")).json()["analysis"]
        assert abs(cents(analysis["iva_total"]) - esperado) <= 1
        assert cents(analysis["iva_total"]) == cents(full["iva_total"])
        assert cents(analysis["total_final"]) == cents(full["total_final"]) == cents(analysis["neto_total"]) + cents(analysis["iva_total"])
        assert cents(client_pdf(client, B2)["total_con_iva"]) == cents(full["total_final"])
        # And the recalculation summary too
        recalc = client.post(url(B2, "cascade-recalculate")).json()["summary"]
        assert abs(cents(recalc["iva_total"]) - round(cents(recalc["neto_total"]) * 0.10)) <= 1


class TestCopy:
    def test_copy_keeps_the_whole_price(self, client, db):
        r = client.post(url(B1, "copy"), json={})
        assert r.status_code == 200, r.text
        nuevo = r.json()["budget_id"]
        check(client, db, nuevo)
        assert sorted(cents(w["total_final"]) for w in works(db, nuevo)) == sorted(
            cents(w["total_final"]) for w in works(db))


# ── Import Excel (T12) ──────────────────────────────────────────────────────


def computation_sheet(rows) -> bytes:
    """01_C&P with a section and works: (code, desc, cantidad, mat_unit, mo_unit, directo, neto)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_C&P"
    for _ in range(7):
        ws.append(["-"])
    ws.append(["1", "ALBAÑILERÍA"])
    for code, desc, cantidad, mat_u, mo_u, directo, neto in rows:
        line = [None] * 26
        line[0:4] = [code, desc, "m2", cantidad]
        line[4], line[9] = mat_u, mo_u
        line[11], line[12], line[13] = mat_u * cantidad, mo_u * cantidad, directo
        line[16], line[19], line[25] = directo * 0.2, directo * 0.1, neto
        ws.append(line)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class TestImportExcel:
    def test_direct_cost_from_the_excel_price_from_the_cascade(self, client, db):
        rows = [("1.1", "Muro", 10, 100, 50, 1500, 1980), ("1.2", "Revoque", 4, 20.5, 10.25, 123, 162.36)]
        files = {"file": ("Ginkgo.xlsx", computation_sheet(rows), "application/octet-stream")}
        r = client.post("/budgets/import-excel", files=files)
        assert r.status_code == 200, r.text
        body = r.json()
        bid = body["budget_id"]
        assert body["neto_excel"] == pytest.approx(1980 + 162.36)
        imported = works(db, bid)
        assert sorted(w["directo_total"] for w in imported) == [123, 1500]
        assert cents(body["neto_app"]) == sum(cents(w["neto_total"]) for w in imported)
        assert cents(body["neto_app"]) == sum(
            cents(calc_cascade_indirects({"directo_total": d}, {})["neto_total"]) for d in (1500, 123))
        check(client, db, bid)


# ── Readers ─────────────────────────────────────────────────────────────────


class TestReaders:
    def test_internal_pdf_ladder_adds_up(self, client, db):
        fitz = pytest.importorskip("fitz")
        r = client.get(url(B1, "export", "pdf"))
        assert r.status_code == 200
        text = "\n".join(page.get_text() for page in fitz.open(stream=r.content, filetype="pdf"))
        summary = client.get(url(B1, "full")).json()["analysis"]

        def ars(value):
            return "$ " + f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        for label, key in (("Impuestos (IIBB + cheque)", "impuestos_total"), ("PRECIO SIN IVA", "neto_total"),
                           ("PRECIO CON IVA", "total_final")):
            assert re.search(re.escape(label) + r"\s+" + re.escape(ars(summary[key])), text), label
        assert cents(summary["directo_total"]) + cents(summary["indirecto_total"]) + cents(
            summary["beneficio_total"]) + cents(summary["impuestos_total"]) == cents(summary["neto_total"])
        # The cascade page shows the saved sums, not a recalculation from the config
        assert re.search(r"= PRECIO SIN IVA\s+" + re.escape(ars(summary["neto_total"])), text)
        assert "Costos Indirectos (34%)" in text

    def test_analysis_has_the_whole_ladder(self, client, db):
        data = client.get(url(B1, "analysis")).json()
        for key in ("impuestos_total", "iva_total", "total_final"):
            assert data[key] > 0
        assert data["items_count"] == 3


class TestVersions:
    def test_each_version_has_its_price(self, client, db):
        assert client.post(url(B1, "versions"), json={}).status_code == 200
        primero = client.get(url(B1, "full")).json()["analysis"]["neto_total"]
        assert client.patch(url(B1, "items", W3), json={"cantidad": 70}).status_code == 200
        assert client.post(url(B1, "versions"), json={}).status_code == 200
        segundo = check(client, db)["neto_total"]
        assert segundo > primero

        r = client.get(url(B1, "versions"))
        assert r.status_code == 200, r.text
        netos = {v["version"]: v["neto_total"] for v in r.json()}
        assert netos == {1: primero, 2: segundo}
        assert all("data" not in v for v in r.json())


# ── Sections of every kind are not works ────────────────────────────────────


class TestSections:
    def test_ai_section_is_not_a_work(self, client, db):
        r = client.post(url(B1, "items", "from-ai"), json={"items": [{
            "seccion_nombre": "Pintura", "seccion_codigo": "2", "codigo": "2.1", "descripcion": "Pintura látex",
            "unidad": "m2", "cantidad": 10,
            "recursos": {"materiales": [{"codigo": "LAD", "descripcion": "Ladrillo", "cantidad_por_unidad": 2}]},
        }]})
        assert r.status_code == 200, r.text
        seccion = next(i for i in db.tables["budget_items"] if i.get("notas") == "Sección generada por IA")
        assert seccion["cantidad"] == 1 and is_section(seccion)
        nuevo = next(i for i in works(db) if i["code"] == "2.1")
        assert nuevo["directo_total"] == pytest.approx(20 * 61.11)  # priced by the cascade at insertion
        assert check(client, db)["items_count"] == 4  # the AI section is not counted
        # Neither the full recalculation nor a change of % prices the AI section
        assert client.post(url(B1, "cascade-recalculate")).json()["items_skipped"] == 2
        assert client.patch(url(B1, "indirects"), json={"beneficio_pct": 12}).status_code == 200
        seccion = next(i for i in db.tables["budget_items"] if i["id"] == seccion["id"])
        assert not seccion.get("neto_total") and not seccion.get("total_final")
        check(client, db)
        rows = excel_router.excel_rows(db.tables["budget_items"], 21)
        assert next(x for x in rows if x["Codigo"] == "2")["Precio sin IVA"] == ""

    def test_summary_ignores_every_section_marker(self):
        items = [{"notas": "Seccion", "neto_total": 5}, {"notas": "SECCIÓN GENERADA POR IA", "cantidad": 1},
                 {"notas": "Sugerido por IA", "neto_total": 10}, {"notas": None, "neto_total": 1}]
        summary = calc_budget_summary(items)
        assert summary["items_count"] == 2 and summary["neto_total"] == 11


class TestInternalPdfDepth:
    """Works under a piso (rubro → piso → trabajo) are listed, so the rows add up to the total."""

    def _two_levels(self, db):
        piso = "00000000-0000-0000-0000-0000000000f1"
        db.tables["budget_items"].append({
            "id": piso, "budget_id": B1, "org_id": ORG, "parent_id": SEC, "code": "", "description": "Planta alta",
            "notas": "Seccion", "sort_order": 20, "cantidad": None})
        deep = "00000000-0000-0000-0000-0000000000f2"
        db.tables["budget_items"].append(_work(deep, "1.5", 4, parent=piso, description="Revoque planta alta",
                                               sort_order=21, mat_unitario=150.15))
        return deep

    def test_groups_reach_every_work(self, client, db):
        self._two_levels(db)
        assert client.post(url(B1, "cascade-recalculate")).status_code == 200
        items = sorted((i for i in db.tables["budget_items"] if i["budget_id"] == B1),
                       key=lambda i: i["sort_order"])
        groups = excel_router.pdf_detail_groups(items)
        listed = [r for _, rows in groups for r in rows if not is_section(r)]
        assert sorted(r["id"] for r in listed) == sorted(w["id"] for w in works(db))
        assert sum(cents(r["neto_total"]) for r in listed) == cents(check(client, db)["neto_total"])

    def test_pdf_lists_the_deep_work_and_its_piso(self, client, db):
        fitz = pytest.importorskip("fitz")
        self._two_levels(db)
        assert client.post(url(B1, "cascade-recalculate")).status_code == 200
        r = client.get(url(B1, "export", "pdf"))
        assert r.status_code == 200
        text = "\n".join(page.get_text() for page in fitz.open(stream=r.content, filetype="pdf"))
        assert "Planta alta" in text and "Revoque planta alta" in text
        deep = next(w for w in works(db) if w["code"] == "1.5")
        ars = "$ " + f"{deep['neto_total']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        assert ars in text
